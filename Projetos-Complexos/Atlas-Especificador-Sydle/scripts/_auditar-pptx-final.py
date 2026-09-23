# -*- coding: utf-8 -*-
"""Auditoria detalhada: Requisitos Atlas - Final.pptx × protótipo × fontes."""
import json
import os
import re
import sys
from collections import defaultdict

from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PPTX = (
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas - Final.pptx"
)
FORMS_PATH = os.path.join(ROOT, "data/subprojects/atlas-prototipo/epics/prototipo/forms.json")
FONTES = os.path.join(ROOT, "docs/_fontes_txt")

sys.path.insert(0, os.path.join(ROOT, "scripts"))
import atlas_classes_config as acc

# espelho exclusões PPTX
EXCLUDE_FIELD_IDS = {
    "patlasv4proto-uo-convites",
    "patlasv4proto-uo-trib-validacao-status",
    "patlasv4proto-uo-trib-validado-por",
    "patlasv4proto-uo-trib-validacao-data",
    "patlasv4proto-uo-meth-aprovar-cadastro",
    "patlasv4proto-uo-meth-solicitar-ajuste",
    "patlasv4proto-uo-meth-gerar-convite",
}
EXCLUDE_METHOD_FORMS = {
    "form-patlasv4-proto-metodo-aprovar-cadastro-org",
    "form-patlasv4-proto-metodo-solicitar-ajuste-org",
}
EXCLUDE_STANDALONE = {
    "form-patlasv4-proto-convite-cadastro",
    "form-patlasv4-proto-metodo-gerar-convite",
}

forms = json.load(open(FORMS_PATH, encoding="utf-8"))
fbi = {f["id"]: f for f in forms}
lbl = {fl["id"]: fl.get("label", fl["id"]) for f in forms for fl in f.get("fields", [])}

# --- extrair PPTX ---
prs = Presentation(PPTX)
slides_data = []
all_table_text = []
tabs_on_slide = []
field_rows = []  # (screen, tab, campo, tipo, obrig, obs)

for si, slide in enumerate(prs.slides):
    slide_title = ""
    tabs = []
    for sh in slide.shapes:
        if sh.has_text_frame and not sh.has_table:
            t = sh.text_frame.text.strip()
            if t and len(t) < 120:
                slide_title = slide_title or t
        if sh.has_table:
            t = sh.table
            for ri in range(len(t.rows)):
                c0 = t.cell(ri, 0).text.strip()
                if c0.startswith("Tela:") and "Aba:" in c0:
                    tabs.append(c0)
                if ri > 0 and c0 and c0 not in ("Campo",) and not c0.startswith("Tela:"):
                    if len(t.rows[ri].cells) >= 4:
                        try:
                            row = [t.cell(ri, ci).text.strip() for ci in range(4)]
                        except Exception:
                            continue
                        if row[0] and row[0] != "Campo":
                            field_rows.append(row)
                            all_table_text.append(" ".join(row))
    if len(tabs) > 1:
        tabs_on_slide.append((si + 1, tabs))
    slides_data.append({"n": si + 1, "title": slide_title, "tabs": tabs})

pptx_full = "\n".join(all_table_text)

# --- campos esperados no PPTX (lógica do gerador) ---
def collect_expected_fields(form_id, screen_name, out):
    form = fbi.get(form_id)
    if not form:
        return
    secs = form.get("sections", [])
    flds = form.get("fields", [])

    def emit(field_list, prefix=""):
        for fld in field_list:
            if fld.get("id") in EXCLUDE_FIELD_IDS:
                continue
            if fld.get("type") == "embeddedReference":
                group = fld.get("label", "Tabela")
                linked = fbi.get(fld.get("linkedFormId"))
                if linked:
                    emit(linked.get("fields", []), prefix=group)
                continue
            label = fld.get("label", "")
            if prefix:
                label = f"{prefix} › {label}"
            out.add(label)

    if not secs:
        emit(flds)
    else:
        tops = [s for s in secs if not s.get("parentSectionId")]
        for sec in tops:
            sec_ids = {sec["id"]}
            for sub in [s for s in secs if s.get("parentSectionId") == sec["id"]]:
                sec_ids.add(sub["id"])
            emit([f for f in flds if f.get("sectionId") in sec_ids])

expected_fields = set()
MODULE_FORMS = []
for key, cfg in acc.CONFIG.items():
    if cfg.get("form_id") and cfg["form_id"] not in EXCLUDE_STANDALONE:
        MODULE_FORMS.append((key, cfg["form_id"], cfg.get("titulo", key)))
        collect_expected_fields(cfg["form_id"], "", expected_fields)
    for ax in cfg.get("anexos_forms", []):
        if ax["form_id"] not in EXCLUDE_STANDALONE:
            MODULE_FORMS.append((key, ax["form_id"], ax.get("titulo", ax["form_id"])))
            collect_expected_fields(ax["form_id"], "", expected_fields)

EXTRA = [
    ("catalogo", "form-patlasv4-proto-cat-produto"),
    ("catalogo", "form-patlasv4-proto-cat-catalogo"),
    ("catalogo", "form-patlasv4-proto-cat-universal"),
    ("catalogo", "form-patlasv4-proto-cat-dados-parceria"),
    ("fluxo", "form-patlasv4-proto-processos"),
    ("fluxo", "form-patlasv4-proto-proposta"),
    ("fluxo", "form-patlasv4-proto-documentos"),
    ("fluxo", "form-patlasv4-proto-contrato"),
    ("fluxo", "form-patlasv4-proto-emissao-ordem-servico"),
    ("fluxo", "form-patlasv4-proto-projeto"),
]
for _, fid in EXTRA:
    collect_expected_fields(fid, "", expected_fields)

METHOD_FORMS = [
    "form-patlasv4-proto-metodo-cadastrar-cargo",
    "form-patlasv4-proto-metodo-migrar-cargos-uo",
    "form-patlasv4-proto-metodo-adicionar-grupos-mipp",
    "form-patlasv4-proto-metodo-substituir-unidade",
    "form-patlasv4-proto-metodo-atribuir-cargo-pessoa",
    "form-patlasv4-proto-metodo-opcoes-pessoa",
    "form-patlasv4-proto-metodo-precadastro-acesso",
    "form-patlasv4-proto-metodo-criar-pessoa",
    "form-patlasv4-proto-metodo-preparar-proposta",
    "form-patlasv4-proto-metodo-aprovar-bloco-produto",
    "form-patlasv4-proto-metodo-criar-os",
    "form-patlasv4-proto-metodo-zerar-pedido-venda",
]

pptx_fields = {r[0] for r in field_rows}
missing_in_pptx = sorted(expected_fields - pptx_fields)
extra_in_pptx = []  # hard to detect orphans

# métodos
pptx_methods = set()
for s in slides_data:
    for sh in prs.slides[s["n"] - 1].shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            if "Método:" in t:
                pptx_methods.add(t)

method_slides = []
for si, slide in enumerate(prs.slides):
    for sh in slide.shapes:
        if sh.has_text_frame and "› Método:" in sh.text_frame.text:
            method_slides.append(sh.text_frame.text.strip())

# fontes keywords
SOURCE_CHECKS = {
    "Pré-cadastro CPF (F1 primário)": (r"pr[eé].?cadastr.*CPF|CPF.*pr[eé].?cadastr", ["Discovery 15 06.txt", "Checkpoint Atlas.txt"]),
    "Convite código (backlog/secundário)": (r"convite.*(c[oó]digo|link)", ["Discovery 15 06.txt"]),
    "Habilitação MIPP 4 grupos": (r"jur[ií]dica|t[eé]cnica|financeira|[Cc]ompliance", ["Projeto Atlas - Modelagem Fase 1 (5).txt"]),
    "Tributos + comprovante F1": (r"tributo|isen[cç][aã]o.*ICMS|DAFI", ["Discovery 15 06.txt", "Projeto - Atlas"]),
    "Assinatura paralela + recusa envelope": (r"paralel|recusa|envelope", ["[Projeto Atlas MT] Assinatura.txt", "Checkpoint Atlas.txt"]),
    "Senha + MFA ou certificado": (r"senha|MFA|certificado|Gov\.?br|MT Login", ["Projeto Atlas - Modelagem Fase 1 (5).txt", "Projeto - Atlas"]),
    "Pro-Rata cobrança": (r"[Pp]ro.?[Rr]ata", ["Gestão Centralizada - DIRC.txt", "Overview - Atlas.txt"]),
    "Catálogo pré-requisito F1": (r"cat[aá]logo", ["Overview - Atlas.txt"]),
    "Cargo ocupante titular/substituto": (r"titular|substituto|suplente", ["Perfil de Permissões (Cargo).txt"]),
    "Pessoa específica workflow opcional": (r"pessoa espec[ií]fica", ["Atlas _ Requisitos.txt"]),
    "Aprovar cadastro parceiro MIPP": (r"aprovar.*cadastro|habilita[cç][aã]o", ["Checkpoint Atlas.txt"]),
    "Versão processo vigente": (r"vers[aã]o.*vigente|inst[aâ]ncia", ["Reunião 1 junho.txt"]),
}

def source_mentions(pattern):
    hits = []
    if not os.path.isdir(FONTES):
        return hits
    for fn in os.listdir(FONTES):
        if not fn.endswith(".txt"):
            continue
        text = open(os.path.join(FONTES, fn), encoding="utf-8", errors="ignore").read()
        if re.search(pattern, text, re.I):
            hits.append(fn)
    return hits

print("=" * 70)
print("AUDITORIA: Requisitos Atlas - Final.pptx")
print("=" * 70)
print(f"Arquivo: {PPTX}")
print(f"Slides: {len(prs.slides)}")
print(f"Linhas de campo na tabela: {len(field_rows)}")
print(f"Campos esperados (gerador): {len(expected_fields)}")
print(f"Campos únicos no PPTX: {len(pptx_fields)}")
print(f"Slides com >1 aba: {len(tabs_on_slide)}")

print("\n--- CAMPOS ESPERADOS AUSENTES NO PPTX (exclusões + gaps) ---")
excluded_labels = []
for fid in EXCLUDE_FIELD_IDS:
    if fid in lbl:
        excluded_labels.append(lbl[fid])
for m in sorted(missing_in_pptx)[:60]:
    flag = " [EXCLUÍDO]" if any(x in m for x in ["Convite", "Validação tributária", "Validado por", "Data da validação"]) or "Aprovar" in m else ""
    print(f"  - {m}{flag}")
if len(missing_in_pptx) > 60:
    print(f"  ... +{len(missing_in_pptx)-60}")

print("\n--- MÉTODOS: esperados vs slides no PPTX ---")
for mf in METHOD_FORMS:
    name = fbi.get(mf, {}).get("name", mf)
    in_pptx = any(mf.replace("form-patlasv4-proto-metodo-", "") in m.lower() or name.lower() in m.lower() for m in method_slides)
    excl = mf in EXCLUDE_METHOD_FORMS
    status = "EXCLUÍDO" if excl else ("OK" if in_pptx else "FALTA")
    print(f"  [{status}] {name}")

missing_methods_proto = [
    "form-patlasv4-proto-metodo-enviar-assinatura",
    "form-patlasv4-proto-metodo-recusar-assinatura",
    "form-patlasv4-proto-metodo-selecionar-metodo-assinatura",
    "form-patlasv4-proto-metodo-aprovar-cadastro-org",
    "form-patlasv4-proto-metodo-solicitar-ajuste-org",
]
print("\n--- MÉTODOS NO PROTÓTIPO SEM SLIDE NO PPTX ---")
for mf in missing_methods_proto:
    print(f"  - {fbi.get(mf, {}).get('name', mf)}")

print("\n--- FONTES × PRESENÇA NO PPTX ---")
for label, (pat, _) in SOURCE_CHECKS.items():
    src_files = source_mentions(pat)
    in_pptx = bool(re.search(pat, pptx_full, re.I))
  # special cases
    if label == "Pré-cadastro CPF (F1 primário)":
        in_pptx = "Pré-cadastrar" in pptx_full or "pré-cadastr" in pptx_full.lower()
    if label == "Aprovar cadastro parceiro MIPP":
        in_pptx = "Aprovar cadastro" in pptx_full or "Aprovada pela MTI" in pptx_full
    if label == "Tributos + comprovante F1":
        in_pptx = "Isento de ICMS" in pptx_full and ("DAFI" in pptx_full or "Regime de Tributação" in pptx_full)
    if label == "Convite código (backlog/secundário)":
        in_pptx = "convite" in pptx_full.lower() and "backlog" in pptx_full.lower()
    ok = "OK" if in_pptx else "PARCIAL/FALTA"
    print(f"  [{ok}] {label}")
    print(f"         fontes: {', '.join(src_files[:3]) or '—'}")

print("\n--- NARRATIVA CONFIG vs PPTX (inconsistências conhecidas) ---")
checks = [
    ("Senha + MFA no protótipo", "Senha + MFA" in pptx_full),
    ("DAFI validação tributária", "DAFI" in pptx_full),
    ("Compliance MIPP", "Compliance" in pptx_full),
    ("Pro-Rata", re.search(r"Pro.?Rata", pptx_full, re.I) is not None),
    ("Portal do Parceiro", "Portal" in pptx_full),
    ("Solicitação de vínculo", "Solicitação de vínculo" in pptx_full or "vínculo" in pptx_full.lower()),
    ("Slide método Aprovar cadastro", any("Aprovar cadastro" in m for m in method_slides)),
    ("Slide método Solicitar ajuste", any("Solicitar ajuste" in m for m in method_slides)),
    ("Slide Enviar assinatura", any("Enviar" in m and "assinatura" in m.lower() for m in method_slides)),
]
for name, ok in checks:
    print(f"  [{'OK' if ok else 'FALTA'}] {name}")

print("\n--- MÓDULOS NO PPTX (forms cobertos) ---")
for key, fid, title in MODULE_FORMS[:20]:
    fm = fbi.get(fid)
    n = len([r for r in field_rows if fm and fm.get("name", "")[:15] in " ".join(r)])
    print(f"  {title[:50]} ({fid})")
print(f"  ... total módulos/forms configurados: {len(MODULE_FORMS)}")
