"""Extract red/green (and other highlight) text runs from the Catalog PPTX."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX_CANDIDATES = [
    Path(r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Requisitos_Finais_Catalogo_Produtos_Atlas.pptx"),
    Path(
        r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle"
        r"\exports\Requisitos_Finais_Catalogo_Produtos_Atlas.pptx"
    ),
]
OUT = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle"
    r"\exports\_pptx_color_audit.json"
)
OUT_TXT = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle"
    r"\exports\_pptx_color_audit.txt"
)


def find_pptx() -> Path:
    for p in PPTX_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError("PPTX not found")


def rgb_of(run):
    try:
        color = run.font.color
        if color is None or color.type is None:
            return None
        try:
            rgb = color.rgb
            if rgb is not None:
                s = str(rgb)
                return (int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
        except Exception:
            pass
        try:
            return f"theme:{color.theme_color}"
        except Exception:
            pass
    except Exception as e:
        return f"err:{e}"
    return None


def classify_color(rgb):
    if rgb is None:
        return "default"
    if isinstance(rgb, str):
        return rgb
    r, g, b = rgb
    if r >= 160 and g <= 110 and b <= 110 and r > g + 30 and r > b + 30:
        return "red"
    if g >= 100 and g > r + 15 and g >= b:
        return "green"
    if r >= 180 and 70 <= g <= 190 and b <= 90:
        return "orange"
    if b >= 130 and b > r and b >= g:
        return "blue"
    return f"other:{r:02X}{g:02X}{b:02X}"


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def collect_runs_from_tf(tf, slide_i, bag, where="shape"):
    for para in tf.paragraphs:
        for run in para.runs:
            t = (run.text or "").strip()
            if not t:
                continue
            rgb = rgb_of(run)
            kind = classify_color(rgb)
            bag.append(
                {
                    "slide": slide_i,
                    "kind": kind,
                    "rgb": None
                    if rgb is None
                    else (rgb if isinstance(rgb, str) else f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"),
                    "text": t[:800],
                    "where": where,
                }
            )


def keep_highlight(item) -> bool:
    kind = item["kind"]
    if kind in ("red", "green", "orange"):
        return True
    if not isinstance(kind, str) or not kind.startswith("other:"):
        return False
    hexc = kind[6:]
    boring = {
        "000000",
        "FFFFFF",
        "333333",
        "666666",
        "595959",
        "7F7F7F",
        "404040",
        "1A1A1A",
        "1F4E79",
        "2F5496",
        "5B9BD5",
        "4472C4",
        "203864",
        "0D1B2A",
        "F2F2F2",
        "D9D9D9",
        "BFBFBF",
        "A6A6A6",
    }
    if hexc in boring:
        return False
    try:
        r = int(hexc[0:2], 16)
        g = int(hexc[2:4], 16)
        b = int(hexc[4:6], 16)
    except Exception:
        return False
    # saturated enough to be a highlight
    return max(r, g, b) - min(r, g, b) > 70 and max(r, g, b) > 120


def main():
    pptx_path = find_pptx()
    prs = Presentation(str(pptx_path))
    print(f"Path: {pptx_path}")
    print(f"Slides: {len(prs.slides)}")

    all_colored = []
    summary = []
    full_slide_text = []

    for si, slide in enumerate(prs.slides, 1):
        bag = []
        plain = []
        for shape in iter_shapes(slide.shapes):
            if shape.has_text_frame:
                collect_runs_from_tf(shape.text_frame, si, bag, "shape")
                plain.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        collect_runs_from_tf(cell.text_frame, si, bag, "table")
                        plain.append(cell.text_frame.text.strip())

        title = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()[:140]
        elif plain:
            title = plain[0][:140]

        colored = [x for x in bag if keep_highlight(x)]
        # dedupe within slide
        seen = set()
        uniq = []
        for x in colored:
            key = (x["kind"], x["text"])
            if key in seen:
                continue
            seen.add(key)
            uniq.append(x)

        all_colored.extend(uniq)
        summary.append({"slide": si, "title": title, "n_highlight": len(uniq)})
        full_slide_text.append(
            {
                "slide": si,
                "title": title,
                "text": "\n".join(t for t in plain if t)[:6000],
            }
        )

    # Prefer red/green first in report
    order = {"red": 0, "green": 1, "orange": 2}
    all_colored.sort(key=lambda x: (x["slide"], order.get(x["kind"], 9), x["text"]))

    payload = {
        "path": str(pptx_path),
        "slides": len(prs.slides),
        "summary": summary,
        "colored": all_colored,
        "slides_text": full_slide_text,
        "counts": dict(Counter(x["kind"] for x in all_colored)),
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = []
    lines.append(f"PPTX: {pptx_path}")
    lines.append(f"Slides: {len(prs.slides)}")
    lines.append(f"Counts: {payload['counts']}")
    lines.append("")
    lines.append("=== HIGHLIGHTS (red/green/orange/other saturated) ===")
    for x in all_colored:
        lines.append(f"S{x['slide']:02d}\t[{x['kind']}]\t{x.get('rgb','')}\t{x['text']}")
    lines.append("")
    lines.append("=== SLIDE TITLES ===")
    for s in summary:
        lines.append(f"S{s['slide']:02d}\t({s['n_highlight']} hl)\t{s['title']}")

    OUT_TXT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Highlights: {len(all_colored)}")
    print(f"Counts: {payload['counts']}")
    print(f"Wrote {OUT}")
    print(f"Wrote {OUT_TXT}")
    for x in all_colored[:100]:
        print(f"S{x['slide']:02d} [{x['kind']}] {x['text'][:160]}")


if __name__ == "__main__":
    main()
