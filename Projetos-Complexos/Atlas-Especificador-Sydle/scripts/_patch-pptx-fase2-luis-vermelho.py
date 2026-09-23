# -*- coding: utf-8 -*-
"""Aplica correções da validação Luís no PPT Fase 2 — texto alterado em vermelho.
Não altera layout (posições/tamanhos/estrutura de shapes); só conteúdo de células/parágrafos.
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt, Emu
from lxml import etree

SRC = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle"
    r"\exports\Fase2_Catalogo_Copia_REVISADO_LUIS.pptx"
)
# also copy to Videos for user convenience
OUT_VIDEOS = Path(r"c:\Users\LucasSantos\Videos\Fase 2 - Catalogo - REVISADO_LUIS.pptx")

RED = RGBColor(0xC0, 0x00, 0x00)


def set_cell_text_red(cell, text: str, font_size_pt: float | None = None) -> None:
    """Replace cell text; all runs red. Keep cell geometry."""
    tf = cell.text_frame
    # clear paragraphs after first
    p0 = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    # clear runs in first paragraph
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    run = p0.add_run()
    run.text = text
    run.font.color.rgb = RED
    run.font.bold = False
    if font_size_pt:
        run.font.size = Pt(font_size_pt)
    else:
        # try keep reasonable size
        run.font.size = Pt(10)


def set_shape_text_red(shape, text: str) -> None:
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    # support multilines via \n → new paragraphs
    lines = text.split("\n")
    p0_run = p0.add_run()
    p0_run.text = lines[0]
    p0_run.font.color.rgb = RED
    p0_run.font.size = Pt(12)
    for line in lines[1:]:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.color.rgb = RED
        r.font.size = Pt(12)


def find_table(slide, min_rows=1):
    tables = [sh for sh in slide.shapes if sh.has_table]
    return tables


def cell(table, r, c):
    return table.cell(r, c)


def patch_slide5(slide) -> None:
    # MTI list shape[2], Parceiro shape[4]
    shapes = [s for s in slide.shapes if s.has_text_frame]
    # by content
    for sh in shapes:
        t = sh.text_frame.text
        if t.startswith("1.") and "Listas básicas" in t:
            set_shape_text_red(
                sh,
                "1.  Listas básicas\n"
                "2.  Parceria e solução\n"
                "3.  Habilita catálogo da parceria (focais/DTIC + toggles)\n"
                "4.  Catálogo universal (só redireciona créditos)\n"
                "5.  Análise / aprovar produtos e catálogo\n"
                "6.  Publicar / versionar\n"
                "7.  Apostilar no contrato",
            )
        elif t.startswith("1.") and "Rascunho" in t:
            set_shape_text_red(
                sh,
                "1.  Cadastro de produtos (no catálogo da parceria)\n"
                "2.  Produtos manual/CSV\n"
                "3.  Visão planilha\n"
                "4.  Enviar à MTI\n"
                "5.  Corrigir após ajuste\n"
                "6.  Status espelhado\n"
                "7.  Não publica / não apostila",
            )


def patch_slide7(slide) -> None:
    table = find_table(slide)[0].table
    # row 5 Modelo de Venda - note pending
    set_cell_text_red(cell(table, 5, 1), "Modelo de Venda (a confirmar)")
    set_cell_text_red(cell(table, 5, 3), "Pendente Luís — possível enxugar")
    # row 8 Catálogo → Catálogo da parceria + hierarquia
    set_cell_text_red(cell(table, 8, 1), "Catálogo da parceria")
    set_cell_text_red(cell(table, 8, 3), "Listas + Parceria · focais/DTIC · toggles")
    # row 9 Produto depends
    set_cell_text_red(cell(table, 9, 3), "Catálogo parceria + listas (sem focais)")
    # row 10 Dados - not validated
    set_cell_text_red(cell(table, 10, 1), "Dados de Parceria (não validado 05/08)")
    # row 11 Universal
    set_cell_text_red(cell(table, 11, 1), "Catálogo Universal (redirecionador)")
    set_cell_text_red(cell(table, 11, 3), "Catálogos elegíveis · SEM focais")


def patch_slide10(slide) -> None:
    tables = [sh.table for sh in find_table(slide)]
    story, campos = tables[0], tables[1]
    set_cell_text_red(
        cell(story, 0, 1),
        "Definir QUANDO o produto é cobrado. Valores fechados: Mensal · Anual · Conforme homologação. "
        "Sob demanda/OS = forma de CONSUMO (proposta/Fase 3), não tipo de cobrança.",
    )
    set_cell_text_red(
        cell(story, 1, 1),
        "A MTI mantém a lista. O Produto seleciona o tipo. "
        "Cobrança no Catálogo da parceria = a confirmar (Luís). "
        "Não exibir no card da listagem.",
    )
    set_cell_text_red(
        cell(story, 2, 1),
        "RF-COB-01: CRUD tipos de cobrança.\n"
        "RF-COB-02: Obrigatório no PRODUTO (catálogo = a confirmar).\n"
        "RN-COB: Distinguir cobrança × forma de consumo (Fase 3).",
    )
    set_cell_text_red(cell(campos, 1, 3), "Mensal · Anual · Conforme homologação")
    set_cell_text_red(
        cell(campos, 2, 0),
        "Recorrência (ocultar)",
    )
    set_cell_text_red(
        cell(campos, 2, 3),
        "NÃO usar no cadastro Fase 2 — consumo/OS na proposta (Luís 05/08).",
    )


def patch_slide14(slide) -> None:
    tables = [sh.table for sh in find_table(slide)]
    story, campos = tables[0], tables[1]
    set_cell_text_red(
        cell(story, 2, 1),
        "Validação Luís 05/08: campos mínimos = Identificador + Parceiro (Organização) + Soluções. "
        "Status DIREX / datas / responsável = fora do mínimo validado.",
    )
    # mark extra fields in red as fora do escopo
    set_cell_text_red(cell(campos, 3, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 4, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 5, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 6, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 8, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 1, 3), "Nome da parceria — MANTIDO (Luís)")
    set_cell_text_red(cell(campos, 2, 3), "Organização (CNPJ) — MANTIDO (Luís)")
    set_cell_text_red(cell(campos, 7, 3), "Soluções da parceria — MANTIDO (Luís)")


def patch_slide15(slide) -> None:
    tables = [sh.table for sh in find_table(slide)]
    story, campos = tables[0], tables[1]
    set_cell_text_red(
        cell(story, 0, 1),
        "Descrever a oferta da parceria. Luís 05/08: só Nome + Parceria + Descrição.",
    )
    set_cell_text_red(
        cell(story, 1, 1),
        "Ligada à Parceria. Fabricante e documentos de apoio = FORA do mínimo validado.",
    )
    set_cell_text_red(cell(campos, 1, 3), "Nome da solução — MANTIDO (Luís)")
    set_cell_text_red(cell(campos, 2, 3), "Parceria — MANTIDO (Luís)")
    set_cell_text_red(cell(campos, 3, 3), "Descrição (se necessário) — MANTIDO (Luís)")
    set_cell_text_red(cell(campos, 4, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 5, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 6, 3), "FORA do mínimo validado (Luís 05/08)")
    set_cell_text_red(cell(campos, 7, 3), "FORA do mínimo validado (Luís 05/08)")


def patch_slide18(slide) -> None:
    tables = [sh.table for sh in find_table(slide)]
    story, campos = tables[0], tables[1]
    set_cell_text_red(
        cell(story, 1, 1),
        "MTI cadastra modelos. Valores PROVISÓRIOS (Luís): Por Licença · Por Serviço · Por Pacote "
        "(no lugar de Homologação). Luís vai pensar se mantém a classe ou descreve no produto.",
    )
    set_cell_text_red(
        cell(campos, 1, 3),
        "Por Licença · Por Serviço · Por Pacote — A CONFIRMAR com Luís",
    )


def patch_slide21(slide) -> None:
    table = find_table(slide)[0].table
    set_cell_text_red(
        cell(table, 0, 1),
        "O parceiro cadastra PRODUTOS no catálogo da parceria e envia à MTI, "
        "sem acessar o backoffice. Não publica nem apostila.",
    )
    set_cell_text_red(
        cell(table, 1, 1),
        "Parceiro: alimenta produtos (tela/CSV) no catálogo da parceria → envia. "
        "MTI: analisa produtos/catálogo, habilita toggles (Universal? / Serviços?), publica e apostila. "
        "Cliente: só consome versão apostilada.",
    )
    set_cell_text_red(
        cell(table, 2, 1),
        "Fluxo (Luís): produtos aprovados → formam catálogo da parceria → "
        "se elegível, aponta ao universal. CSV pode variar por parceria. "
        "RF-PARC mantidos; RN-PARC-06: parceiro não publica/apostila.",
    )


def patch_slide24(slide) -> None:
    table = find_table(slide)[0].table
    set_cell_text_red(
        cell(table, 1, 1),
        "MTI: listas, parcerias, soluções; analisa envios; habilita Universal?/Catálogo de serviços? "
        "no catálogo da parceria; mantém universal (redirecionador); publica; apostila.",
    )
    set_cell_text_red(
        cell(table, 2, 1),
        "RF-MTI: analisar produtos e catálogo; publicar só após Homologado; "
        "apostilar; Catálogo Universal SEM focais; conferir fator Tipo 3.",
    )


def patch_slide25(slide) -> None:
    # BPMN with many text shapes — find and recolor key labels
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        replacements = {
            "2. Criar rascunho\ndo catálogo": "2. Formar catálogo\nda parceria (após produtos)",
            "2. Criar rascunho\rdo catálogo": "2. Formar catálogo\nda parceria (após produtos)",
            "3. Incluir produtos\nmanual ou CSV": "3. Incluir/aprovar produtos\nmanual ou CSV",
            "3. Incluir produtos\rmanual ou CSV": "3. Incluir/aprovar produtos\nmanual ou CSV",
            "1. Configurar base\nlistas, parceria, solução": "1. Configurar base\nlistas, parceria, solução\n(+ toggles no catálogo)",
        }
        # normalize newlines
        key = t.replace("\r", "\n")
        for old, new in list(replacements.items()):
            if key == old.replace("\r", "\n") or t.replace("\r\n", "\n") == old.replace("\r", "\n"):
                set_shape_text_red(sh, new)
                break
        # rules footer
        if t.startswith("Regras-chave:"):
            set_shape_text_red(
                sh,
                "Regras-chave (Luís 05/08): produtos aprovados → catálogo da parceria → "
                "universal (redirecionador, sem focais) · focais/DTIC no catálogo da parceria · "
                "enviar ≥1 produto · publicar ≠ cliente ver (só apostila) · parceiro não publica/apostila · "
                "MTI habilita Universal?/Catálogo de serviços?.",
            )


def patch_slide28(slide) -> None:
    # Classe Catálogo — texto only, no layout change
    tables = [sh.table for sh in find_table(slide)]
    if not tables:
        return
    story = tables[0]
    set_cell_text_red(
        cell(story, 1, 1),
        "Catálogo DA PARCERIA (cardápio). Luís: produtos aprovados formam o catálogo; "
        "focais/DTIC ficam AQUI; toggles Universal? e Catálogo de serviços?; "
        "CSV pode variar por parceria. Cliente só vê após apostila.",
    )
    set_cell_text_red(
        cell(story, 2, 1),
        "RF-CTL mantidos + RF: toggles Universal?/Serviços?. "
        "RN: hierarquia Produto → Catálogo parceria → Universal. "
        "Publicar só MTI. CSV round-trip (layout pode variar por parceria).",
    )


def patch_slide29(slide) -> None:
    """Campos do catálogo — NÃO muda layout/linhas; só textos em células existentes."""
    table = find_table(slide)[0].table
    # Cobrança obs
    set_cell_text_red(
        cell(table, 6, 3),
        "Mensal · Anual · Conforme homologação — obrigatoriedade no catálogo A CONFIRMAR (Luís)",
    )
    # Focais — reinforce correct place
    set_cell_text_red(cell(table, 8, 3), "Contato comercial — FICA NO CATÁLOGO DA PARCERIA (Luís)")
    set_cell_text_red(cell(table, 9, 3), "Pós-vendas — FICA NO CATÁLOGO DA PARCERIA (Luís)")
    set_cell_text_red(cell(table, 10, 3), "UO MTI — FICA NO CATÁLOGO DA PARCERIA (Luís)")
    # Parceria obs — inject toggles note without new row
    set_cell_text_red(
        cell(table, 12, 3),
        "Parceria vinculada. INCLUIR toggles (mesmo layout): Universal? · Catálogo de serviços? (ambos podem Ser Sim)",
    )
    # Observações — reinforce
    set_cell_text_red(
        cell(table, 15, 3),
        "Notas. CSV: estrutura de colunas PODE VARIAR por parceria (Cloud ≠ Simplifica).",
    )


def patch_slide30(slide) -> None:
    tables = [sh.table for sh in find_table(slide)]
    story = tables[0]
    set_cell_text_red(
        cell(story, 1, 1),
        "Página única. Card: nome + tipo (Licença|Serviço) + status. "
        "Luís: só Licença ou Serviço. Complexidade/peso/qtde métrica SÓ se Serviço. "
        "Focais/DTIC NÃO ficam no produto.",
    )
    set_cell_text_red(
        cell(story, 2, 1),
        "RF-PROD: catálogo, tipo, métrica, cobrança (Mensal/Anual/Conforme homologação), "
        "modelo venda, valor, Universal/Individualizado, QUANTIDADE DA MÉTRICA (serviço). "
        "Sem DTIC/focais no produto. Fator se Universal=Sim.",
    )


def patch_slide31(slide) -> None:
    table = find_table(slide)[0].table
    # map rows by first column
    rows_by_name = {}
    for r in range(1, len(table.rows)):
        name = table.cell(r, 0).text.strip()
        rows_by_name[name] = r

    def touch(name, col, text):
        r = rows_by_name.get(name)
        if r is not None:
            set_cell_text_red(cell(table, r, col), text)

    touch("Cobrança", 3, "Mensal · Anual · Conforme homologação (Luís)")
    touch("Modelo de venda", 3, "Por Licença · Por Serviço · Por Pacote — a confirmar")
    touch("Complexidade", 3, "SÓ Serviço; níveis muito baixa→muito alta + coeficiente")
    touch("Peso", 3, "SÓ Serviço — quando o cálculo exigir")
    touch(
        "Grupo de atendimento",
        0,
        "Quantidade da métrica",
    )
    touch(
        "Quantidade da métrica",
        1,
        "Número",
    )
    touch(
        "Quantidade da métrica",
        2,
        "Cond. Serviço",
    )
    touch(
        "Quantidade da métrica",
        3,
        "Qtde da MÉTRICA selecionada (não fixar HST/UST) — Luís 05/08",
    )
    # If rename didn't update map (we set col0 after map built), find by old or new
    # Re-scan
    for r in range(1, len(table.rows)):
        name = table.cell(r, 0).text.strip()
        if name == "Grupo de atendimento":
            set_cell_text_red(cell(table, r, 0), "Quantidade da métrica")
            set_cell_text_red(cell(table, r, 1), "Número")
            set_cell_text_red(cell(table, r, 2), "Cond. Serviço")
            set_cell_text_red(
                cell(table, r, 3),
                "Qtde da MÉTRICA selecionada (não fixar HST/UST) — Luís 05/08",
            )
        if name == "Unidade DTIC":
            set_cell_text_red(cell(table, r, 0), "Unidade DTIC")
            set_cell_text_red(cell(table, r, 2), "NÃO")
            set_cell_text_red(
                cell(table, r, 3),
                "REMOVIDO do Produto → fica no Catálogo da parceria (Luís 05/08)",
            )
        if name.startswith("Moeda universal"):
            set_cell_text_red(
                cell(table, r, 3),
                "Só se Universal=Sim; base do fator automático",
            )
        if name == "Fator":
            set_cell_text_red(
                cell(table, r, 3),
                "Calculado (unitário ÷ moeda); até 6 casas — Luís",
            )


def patch_slide32(slide) -> None:
    tables = [sh.table for sh in find_table(slide)]
    story = tables[0]
    campos = tables[1] if len(tables) > 1 else None
    set_cell_text_red(
        cell(story, 0, 1),
        "Redirecionador de créditos Tipo 3. Agrega catálogos de parceria elegíveis. "
        "SEM focais/responsáveis (Luís).",
    )
    set_cell_text_red(
        cell(story, 1, 1),
        "MTI mantém. Universal aponta para catálogos da parceria com Universal?=Sim. "
        "Não substitui o cardápio da parceria. Nem toda parceria entra no crédito universal.",
    )
    if campos:
        set_cell_text_red(
            cell(campos, 1, 3),
            "Ativo · Homologado · Paralisado (sem Concluído — Luís)",
        )
        set_cell_text_red(
            cell(campos, 2, 3),
            "Só catálogos elegíveis (toggle Universal?). SEM responsáveis neste formulário.",
        )


def patch_slide2(slide) -> None:
    table = find_table(slide)[0].table
    set_cell_text_red(
        cell(table, 1, 1),
        "Flags Universal/Individualizado + métrica definem elegibilidade. "
        "Tipo 1/2/3 = objeto do CONTRATO (não escolha isolada no cadastro). "
        "Tipo 2 amarra Catálogo da parceria + moeda; Grupo NÃO é objeto do contrato.",
    )


def patch_slide20(slide) -> None:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Classe - Dados"):
            set_shape_text_red(sh, "Classe - Dados de Parceria por Produto (NÃO validado Luís 05/08)")
            break
    tables = find_table(slide)
    if tables:
        story = tables[0].table
        set_cell_text_red(
            cell(story, 2, 1),
            "ATENÇÃO: esta classe NÃO passou no passe final de validação com Luís (05/08). "
            "Manter como protótipo até próxima agenda.",
        )


def main() -> None:
    prs = Presentation(str(SRC))
    patches = {
        2: patch_slide2,
        5: patch_slide5,
        7: patch_slide7,
        10: patch_slide10,
        14: patch_slide14,
        15: patch_slide15,
        18: patch_slide18,
        20: patch_slide20,
        21: patch_slide21,
        24: patch_slide24,
        25: patch_slide25,
        28: patch_slide28,
        29: patch_slide29,
        30: patch_slide30,
        31: patch_slide31,
        32: patch_slide32,
    }
    for num, fn in patches.items():
        fn(prs.slides[num - 1])
        print("patched slide", num)

    prs.save(str(SRC))
    print("saved", SRC)
    try:
        import shutil

        shutil.copy2(SRC, OUT_VIDEOS)
        print("copied", OUT_VIDEOS)
    except Exception as e:
        print("videos copy failed", e)
        # try shared write via bytes
        try:
            OUT_VIDEOS.write_bytes(SRC.read_bytes())
            print("copied bytes", OUT_VIDEOS)
        except Exception as e2:
            print("videos copy2 failed", e2)


if __name__ == "__main__":
    main()
