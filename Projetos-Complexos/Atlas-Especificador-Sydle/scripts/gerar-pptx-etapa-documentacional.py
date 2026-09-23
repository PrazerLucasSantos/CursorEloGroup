# -*- coding: utf-8 -*-
"""PPTX — Classe Etapa documentacional (catálogo) e uso na Organização. Layout Pessoa."""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "docs/entregaveis/pptx/Atlas requisitos.pptx")
OUT_PATH = os.path.join(ROOT, "exports/Classe_Etapa_Documentacional_Organizacao.pptx")

BLUE = RGBColor(0x18, 0x22, 0xDC)
DARK = RGBColor(0x27, 0x27, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xBF, 0xBF)
GREY_HDR = RGBColor(0xD9, 0xD9, 0xD9)

F_TITLE = "PP Telegraf"
F_HEAD = "PP Telegraf"
F_LABEL = "PP Telegraf SemiBold"
F_BODY = "Arial"

COL_W = [Inches(2.4), Inches(2.6), Inches(2.4), Inches(4.9)]
HEADERS = ["Grupo", "Campo", "Opções", "Regra / Validação"]


def clear_slides(prs: Presentation) -> None:
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        r_id = sld.get(qn("r:id"))
        try:
            prs.part.drop_rel(r_id)
        except Exception:
            pass
        xml_slides.remove(sld)


def set_fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def bottom_border(cell, rgb=LINE, w=6350):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("a:lnB"))
    if existing is not None:
        tc_pr.remove(existing)
    ln = tc_pr.makeelement(qn("a:lnB"), {"w": str(w), "cap": "flat"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])})
    fill.append(clr)
    ln.append(fill)
    tc_pr.append(ln)


def style_cell(
    cell,
    text,
    *,
    bold=False,
    size=10,
    color=DARK,
    font=F_BODY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text or "").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(1)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        clear_slides(self.prs)
        self.layout = next(
            (l for l in self.prs.slide_layouts if l.name == "Todo branco"),
            self.prs.slide_layouts[0],
        )

    def add_slide(self):
        return self.prs.slides.add_slide(self.layout)

    def add_title(self, slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.0), Inches(0.45))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = DARK

    def intro_slide(self, title, historia, contexto, requisitos, regras):
        slide = self.add_slide()
        self.add_title(slide, title)
        tbl = slide.shapes.add_table(3, 2, Inches(0.45), Inches(0.72), Inches(12.3), Inches(5.95)).table
        tbl.columns[0].width = Inches(2.1)
        tbl.columns[1].width = Inches(10.2)
        for i, lab in enumerate(["História", "Contexto da História", "Requisitos e Regras"]):
            set_fill(tbl.cell(i, 0), BLUE)
            style_cell(
                tbl.cell(i, 0),
                lab,
                bold=True,
                size=11,
                color=WHITE,
                font=F_LABEL,
                anchor=MSO_ANCHOR.TOP,
            )
        style_cell(tbl.cell(0, 1), historia, size=10, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(1, 1), contexto, size=10, anchor=MSO_ANCHOR.TOP)

        c2 = tbl.cell(2, 1)
        c2.vertical_anchor = MSO_ANCHOR.TOP
        tf = c2.text_frame
        tf.word_wrap = True

        def par(text, *, bold=False, bullet=False, first=False, color=DARK):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.space_before = Pt(2 if not first else 0)
            p.space_after = Pt(1)
            r = p.add_run()
            r.text = ("•  " if bullet else "") + text
            r.font.size = Pt(10)
            r.font.bold = bold
            r.font.name = F_BODY
            r.font.color.rgb = color

        par("Requisitos Funcionais", bold=True, color=BLUE, first=True)
        for req in requisitos:
            par(req, bullet=True)
        par("Regras de Negócio", bold=True, color=BLUE)
        for rg in regras:
            par(rg, bullet=True)

        for i in range(3):
            bottom_border(tbl.cell(i, 0))
            bottom_border(tbl.cell(i, 1))

    def fields_slide(self, title: str, rows: list[tuple[str, str, str, str]]):
        """Telas e Campos — Grupo | Campo | Opções | Regra."""
        slide = self.add_slide()
        self.add_title(slide, title)
        # bar
        bar = slide.shapes.add_table(1, 1, Inches(0.45), Inches(0.68), Inches(12.3), Inches(0.32)).table
        set_fill(bar.cell(0, 0), BLUE)
        style_cell(bar.cell(0, 0), "Telas e Campos", bold=True, size=11, color=WHITE, font=F_HEAD)

        n = len(rows) + 1
        height = min(Inches(5.5), Inches(0.38 * n + 0.2))
        tbl = slide.shapes.add_table(n, 4, Inches(0.45), Inches(1.08), Inches(12.3), height).table
        for i, w in enumerate(COL_W):
            tbl.columns[i].width = w
        for ci, h in enumerate(HEADERS):
            set_fill(tbl.cell(0, ci), GREY_HDR)
            style_cell(tbl.cell(0, ci), h, bold=True, size=10, color=DARK, font=F_HEAD)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                style_cell(tbl.cell(ri, ci), val, size=9, anchor=MSO_ANCHOR.TOP)
            for ci in range(4):
                bottom_border(tbl.cell(ri, ci))

    def build(self):
        # --- Classe catálogo ---
        self.intro_slide(
            "Etapa documentacional",
            "Como gestor de cadastro, quero definir etapas documentacionais com grupos e tipos de documento, "
            "para que a Organização inclua só a etapa e receba automaticamente a lista de documentos exigidos.",
            "Classe de catálogo (form-patlasv4-proto-cat-etapa-documentacional). "
            "Modelo canônico: Etapa → Grupo → Tipo de Documento. "
            "Exemplos de etapa: Habilitação documental, Pré-rito parceria, Execução da parceria. "
            "Em cada etapa montam-se os Grupos de Documento e, dentro de cada grupo, os Tipos (obrigatórios ou não). "
            "Na Organização, ao selecionar e Adicionar a etapa, grupos e documentos são carregados do catálogo — "
            "não se escolhe grupo na tela da Organização. "
            "Classes relacionadas: Grupo de Documento, Tipo de Documento; embutidos Grupo na etapa e Tipo na etapa.",
            [
                "Cadastrar etapas com Nome, Ordem e Ativo.",
                "Em cada etapa, adicionar Grupos de Documento do catálogo.",
                "Em cada grupo da etapa, adicionar/remover Tipos de Documento e marcar obrigatoriedade.",
                "Disponibilizar etapas ativas na seleção “Etapas a incluir” da Organização.",
                "Ao incluir a etapa na Organização, copiar grupos e tipos configurados para a grade documental.",
            ],
            [
                "Etapas inativas não aparecem em Etapas a incluir.",
                "A composição documental é definida na Etapa (não na Organização).",
                "Grupo de Documento referencia a Etapa documentacional do catálogo.",
                "Tipo de Documento permanece no catálogo MIPP; na etapa indica-se quais tipos entram e se são obrigatórios.",
                "Parceiro/cliente no portal só anexam arquivos nos tipos já incluídos — não criam etapa/grupo/tipo.",
            ],
        )

        self.fields_slide(
            "Etapa documentacional — Catálogo",
            [
                ("Etapa", "Nome", "Texto", "Obrigatório. Ex.: Habilitação documental, Pré-rito, Execução."),
                ("Etapa", "Ordem", "Número", "Ordem de exibição ao incluir etapas na Organização."),
                ("Etapa", "Ativo", "Sim / Não", "Inativa: não aparece em Etapas a incluir."),
                ("Etapa", "Grupos de documentos", "Embutido (vários)", "Lista de grupos da etapa (classe Grupo na etapa)."),
                ("Grupo na etapa", "Grupo de documentos", "Referência", "Seleciona Grupo de Documento do catálogo."),
                ("Grupo na etapa", "Documentos (tipos)", "Tabela embutida", "Tipos desta etapa/grupo; adicionar/remover."),
                ("Tipo na etapa", "Tipo de documento", "Referência / opções", "Tipo do catálogo MIPP."),
                ("Tipo na etapa", "Obrigatório na etapa", "Sim / Não", "Se obrigatório ao incluir a etapa na Organização."),
            ],
        )

        # --- Uso na Organização ---
        self.intro_slide(
            "Organização — Documentos (Etapas)",
            "Como analista MTI ou responsável no portal, quero incluir etapas documentacionais na Organização "
            "e acompanhar o progresso dos anexos, para habilitar a parceria sem montar a lista documento a documento.",
            "Na aba Documentos da Organização: campos Etapas a incluir, botão Adicionar e Etapas incluídas. "
            "Etapas incluídas usa a classe embutida form-patlasv4-proto-uo-etapa-documentacional "
            "(espelho da etapa do catálogo na instância da Organização). "
            "Cada etapa embute Grupos (form-patlasv4-proto-uo-grupo-doc-etapa) e cada grupo embute "
            "Documentos (form-patlasv4-proto-uo-documento) para anexo e validação. "
            "No portal, a mesma hierarquia aparece em Documentos (Etapas a incluir / Progresso / Etapas incluídas). "
            "Seleção de “Grupos de documento” multipla antiga é legado e permanece oculta.",
            [
                "Selecionar uma ou mais etapas ativas e clicar em Adicionar.",
                "Exibir Etapas incluídas em acordeão: etapa → grupos → documentos.",
                "Mostrar progresso da etapa e do grupo (ex.: 2/9).",
                "Permitir anexar arquivo e acompanhar status de validação em cada documento.",
                "No portal, refletir as etapas já incluídas pela MTI (somente leitura da estrutura).",
            ],
            [
                "Na Organização só se seleciona a etapa — grupo e tipos vêm do catálogo.",
                "Não é permitido criar etapa/grupo/tipo na Organização ou no portal.",
                "Campos de identidade da etapa (nome, ordem, ativo, progresso) são somente leitura após inclusão.",
                "Documentos de execução pertencem à etapa Execução (quando incluída); tabela legada de execução fica oculta se coberta pela etapa.",
                "Salvar no portal envia o conjunto documental para análise da MTI.",
            ],
        )

        self.fields_slide(
            "Organização — Aba Documentos (uso da Etapa)",
            [
                ("Documentos", "Etapas a incluir", "Habilitação documental · Pré-rito · Execução", "Multi-seleção do catálogo (apenas ativas)."),
                ("Documentos", "Adicionar", "Ação (+)", "Inclui etapas selecionadas e carrega grupos/tipos."),
                ("Documentos", "Etapas incluídas", "Embutido Etapa documentacional (UO)", "Instâncias da etapa na Organização."),
                ("Etapa (UO)", "Etapa documentacional", "Texto / opções (leitura)", "Identidade no título do acordeão."),
                ("Etapa (UO)", "Ordem / Ativo", "Número / Sim-Não", "Somente leitura; ocultos no corpo."),
                ("Etapa (UO)", "Progresso da etapa", "Texto (leitura)", "Destaque à direita (ex.: 22/40)."),
                ("Etapa (UO)", "Grupos e documentos", "Embutido Grupo na etapa", "Vem do catálogo ao Adicionar."),
                ("Grupo (UO)", "Grupo de documentos", "Referência (leitura)", "Identidade do acordeão do grupo."),
                ("Grupo (UO)", "Progresso", "Texto (leitura)", "Ex.: 12/19 no cabeçalho."),
                ("Grupo (UO)", "Documentos", "Tabela Documento UO", "Anexar arquivo e status de validação."),
                ("Documento (UO)", "Tipo / Arquivo / Status", "Conforme formulário UO", "Anexo do parceiro/cliente; validação MTI."),
            ],
        )

        # --- Relacionamento resumido ---
        self.intro_slide(
            "Etapa documentacional — Relacionamento",
            "Como equipe Atlas, quero deixar explícito o vínculo entre catálogo e Organização, "
            "para evitar o modelo antigo de multi-seleção de grupos na Organização.",
            "Cadeia: Tipo de Documento e Grupo de Documento (catálogo) → Etapa documentacional (catálogo) "
            "compõe Grupos + Tipos → Organização.Documentos inclui Etapa → gera Etapa (UO) → Grupo (UO) → Documento (UO). "
            "IDs: catálogo form-patlasv4-proto-cat-etapa-documentacional; "
            "embutido na Organização form-patlasv4-proto-uo-etapa-documentacional "
            "(campo patlasv4proto-uo-etapas-documentacionais).",
            [
                "Manter uma única fonte de verdade no catálogo Etapa documentacional.",
                "Propagar estrutura para a Organização apenas via Adicionar etapas.",
                "Usar a mesma hierarquia no portal do cliente/parceiro.",
            ],
            [
                "Não reintroduzir seleção de grupos na Organização (legado oculto).",
                "Alteração de catálogo afeta novas inclusões de etapa; instâncias já incluídas preservam a cópia carregada até reprocessamento definido pela MTI.",
                "Portal: estrutura read-only; edição = anexos e dados permitidos; Salvar envia à MTI.",
            ],
        )

        os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
        self.prs.save(OUT_PATH)
        print(OUT_PATH)


if __name__ == "__main__":
    Deck().build()
