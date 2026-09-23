# -*- coding: utf-8 -*-
"""PowerPoint — Portal Cliente/Parceiro no layout do anexo (História · Contexto · Requisitos e Regras)."""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "docs/entregaveis/pptx/Atlas requisitos.pptx")
OUT_PATH = os.path.join(ROOT, "exports/Portal_Cliente_Parceiro_Requisitos_Fluxo.pptx")

BLUE = RGBColor(0x18, 0x22, 0xDC)
DARK = RGBColor(0x27, 0x27, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xBF, 0xBF)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
SOFT = RGBColor(0xEA, 0xF4, 0xFB)
NAVY = RGBColor(0x00, 0x33, 0x66)
GREEN = RGBColor(0x16, 0x65, 0x34)
ORANGE = RGBColor(0xC2, 0x41, 0x0C)

F_TITLE = "PP Telegraf"
F_HEAD = "PP Telegraf"
F_LABEL = "PP Telegraf SemiBold"
F_BODY = "Arial"


def clear_slides(prs: Presentation) -> None:
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        r_id = sld.get(qn("r:id"))
        try:
            prs.part.drop_rel(r_id)
        except Exception:
            pass
        xml_slides.remove(sld)


def set_fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def bottom_border(cell, rgb=LINE, w=6350):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("a:lnB"))
    if existing is not None:
        tc_pr.remove(existing)
    ln = tc_pr.makeelement(qn("a:lnB"), {"w": str(w), "cap": "flat"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])})
    fill.append(clr)
    ln.append(fill)
    tc_pr.append(ln)


def style_cell(
    cell,
    text,
    *,
    bold=False,
    size=10,
    color=DARK,
    font=F_BODY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.1)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text or "").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        p.space_before = Pt(0)
        for r in list(p.runs):
            r.text = ""
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        clear_slides(self.prs)
        self.layout = next(
            (l for l in self.prs.slide_layouts if l.name == "Todo branco"),
            self.prs.slide_layouts[0],
        )
        self.sw = self.prs.slide_width

    def add_slide(self):
        return self.prs.slides.add_slide(self.layout)

    def add_title(self, slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.0), Inches(0.45))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = DARK

    def section_divider(self, title, subtitle):
        slide = self.add_slide()
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
        for y, txt, sz, bold in (
            (2.5, title, 34, True),
            (3.55, subtitle, 14, False),
        ):
            box = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(11.5), Inches(0.9))
            p = box.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.name = F_TITLE if bold else F_BODY
            r.font.color.rgb = WHITE

    def intro_slide(self, title, historia, contexto, requisitos, regras):
        """Layout do anexo: tabela 3×2 com faixas azuis à esquerda."""
        slide = self.add_slide()
        self.add_title(slide, title)
        tbl = slide.shapes.add_table(3, 2, Inches(0.45), Inches(0.72), Inches(12.3), Inches(5.95)).table
        tbl.columns[0].width = Inches(2.1)
        tbl.columns[1].width = Inches(10.2)
        for i, lab in enumerate(["História", "Contexto da História", "Requisitos e Regras"]):
            set_fill(tbl.cell(i, 0), BLUE)
            style_cell(
                tbl.cell(i, 0),
                lab,
                bold=True,
                size=11,
                color=WHITE,
                font=F_LABEL,
                anchor=MSO_ANCHOR.TOP,
            )
        style_cell(tbl.cell(0, 1), historia, size=10, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(1, 1), contexto, size=10, anchor=MSO_ANCHOR.TOP)

        c2 = tbl.cell(2, 1)
        c2.vertical_anchor = MSO_ANCHOR.TOP
        tf = c2.text_frame
        tf.word_wrap = True

        def par(text, *, bold=False, bullet=False, first=False, color=DARK, size=10):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.space_before = Pt(2 if not first else 0)
            p.space_after = Pt(1)
            r = p.add_run()
            r.text = ("•  " if bullet else "") + text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = F_BODY
            r.font.color.rgb = color

        par("Requisitos Funcionais", bold=True, color=BLUE, first=True)
        for req in requisitos:
            par(req, bullet=True)
        par("Regras de Negócio", bold=True, color=BLUE)
        for rg in regras:
            par(rg, bullet=True)

        for i in range(3):
            bottom_border(tbl.cell(i, 0))
            bottom_border(tbl.cell(i, 1))

    def fluxo_slide(self):
        slide = self.add_slide()
        self.add_title(slide, "Portal — Fluxo (diagrama / modelagem)")

        steps = [
            "MTI inicia\ncadastro",
            "Disponibiliza\nao responsável",
            "Acesso portal\n(CPF)",
            "Completa\ndados/docs",
            "Salvar =\nenviar MTI",
            "Meu painel\nacompanha",
            "Parecer\nMTI",
        ]
        box_w = Inches(1.45)
        box_h = Inches(0.95)
        gap = Inches(0.18)
        arrow_w = Inches(0.22)
        n = len(steps)
        total = n * box_w + (n - 1) * (gap + arrow_w)
        x = (self.sw - total) // 2
        y = Inches(1.15)

        for i, label in enumerate(steps):
            sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
            sh.fill.solid()
            sh.fill.fore_color.rgb = SOFT if i < n - 1 else NAVY
            sh.line.color.rgb = LINE
            try:
                sh.adjustments[0] = 0.15
            except Exception:
                pass
            tf = sh.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            r.text = label
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.name = F_BODY
            r.font.color.rgb = WHITE if i == n - 1 else DARK
            x += box_w
            if i < n - 1:
                ar = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    x + Inches(0.04),
                    y + Inches(0.38),
                    arrow_w,
                    Inches(0.18),
                )
                ar.fill.solid()
                ar.fill.fore_color.rgb = BLUE
                ar.line.fill.background()
                x += gap + arrow_w

        # Decisions
        box = slide.shapes.add_textbox(Inches(0.45), Inches(2.4), Inches(12.3), Inches(0.35))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = "Decisão da MTI"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.name = F_HEAD
        r.font.color.rgb = BLUE
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        outcomes = [
            (GREEN, "Aprovar — cadastro liberado; acesso comercial conforme regra"),
            (ORANGE, "Solicitar ajuste — mensagem no Meu painel com botão Atender"),
            (RGBColor(0x99, 0x1B, 0x1B), "Reprovar — cadastro encerrado"),
        ]
        ow = Inches(3.8)
        og = Inches(0.3)
        ox = (self.sw - (ow * 3 + og * 2)) // 2
        for i, (color, text) in enumerate(outcomes):
            sh = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                ox + i * (ow + og),
                Inches(2.85),
                ow,
                Inches(0.85),
            )
            sh.fill.solid()
            sh.fill.fore_color.rgb = WHITE
            sh.line.color.rgb = color
            tf = sh.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            r.text = text
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.name = F_BODY
            r.font.color.rgb = color

        # Numbered modeling notes
        notes = [
            "1–9  Back-office: cadastrar Organização, etapas documentais, validar Responsável/CPF/Cargo e disponibilizar.",
            "10–11  Portal: login MT Login/Gov.br; match do CPF com a Pessoa responsável vinculada.",
            "12–16  Completar Meus dados e Cadastro da organização (campos visíveis do perfil). Dados do Cadastro = somente leitura.",
            "17–18  Salvar no portal envia à MTI (status Completa – aguardando MTI). Sem botão separado de envio.",
            "19–20  MTI analisa e decide. Em ajuste: usuário abre a solicitação no Meu painel → mensagem → Atender → mesmo formulário.",
        ]
        box = slide.shapes.add_textbox(Inches(0.55), Inches(3.95), Inches(12.1), Inches(2.5))
        tf = box.text_frame
        tf.word_wrap = True
        for i, note in enumerate(notes):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(4)
            r = p.add_run()
            r.text = note
            r.font.size = Pt(11)
            r.font.name = F_BODY
            r.font.color.rgb = DARK

    def build(self):
        self.section_divider(
            "Portal do Cliente e do Parceiro",
            "Atlas · MTI — História, requisitos, regras e fluxo",
        )

        # 1. Visão geral
        self.intro_slide(
            "Portal — Visão geral",
            "Como cliente ou parceiro da MTI, quero acessar um portal para completar o cadastro da organização, "
            "atualizar Meus dados, acompanhar solicitações no Meu painel e responder a ajustes, "
            "para que a MTI analise e aprove sem eu precisar operar o back-office.",
            "Na Fase 1 o portal cobre documentação (etapas), dados da organização/parceria e pessoas. "
            "A MTI inicia o cadastro no back-office (nome, CNPJ, sigla, etapas/documentos exigidos, responsável) "
            "e disponibiliza ao parceiro/cliente. O responsável acessa o portal (match por CPF), completa os dados "
            "visíveis e salva — o salvamento envia para análise da MTI. Colaboradores pré-cadastrados veem Meus dados. "
            "Solicitações e pareceres aparecem no Meu painel; ajuste da MTI exibe Atender na mensagem.",
            [
                "Disponibilizar portal exclusivo a CPFs pré-cadastrados na classe Pessoa.",
                "Permitir ao responsável completar Organização (campos editáveis) e documentos das etapas incluídas.",
                "Exibir Meus dados ao colaborador logado e Cadastro da organização ao responsável habilitado.",
                "Mostrar progresso documental por etapa/grupo (ex.: 2/9 anexados).",
                "Ao Salvar, enviar o cadastro à MTI (sem etapa separada de “Enviar para aprovação”).",
                "No Meu painel, listar solicitações; no detalhe, histórico; em ajuste, botão Atender na mensagem.",
                "Atender abre o mesmo layout do Cadastro da organização com os dados enviados para correção.",
            ],
            [
                "Somente CPF pré-cadastrado acessa o portal (sem auto cadastro aberto).",
                "Somente o Responsável da Organização envia o cadastro inicial.",
                "Dados do Cadastro (Nome, Sigla, CNPJ, Código, Ativo) são somente leitura no portal.",
                "Não há Validar CNPJ no portal; preenchimento fiscal editável fica em Dados da Organização.",
                "Parceiro/cliente não cria etapas, grupos ou tipos de documento — só anexa o que a MTI configurou.",
                "MTI pode Aprovar, Solicitar ajuste ou Reprovar, sempre com notificação/histórico.",
                "Pós-aprovação, alterações geram nova solicitação sem bloquear o acesso ao portal.",
                "Formulários do portal reutilizam regras do back-office; seções sem campos visíveis não aparecem.",
            ],
        )

        # 2. Fluxo
        self.fluxo_slide()

        # 3. Cadastro da organização
        self.intro_slide(
            "Portal — Cadastro da organização",
            "Como responsável, quero completar e enviar o cadastro da organização no portal, "
            "para habilitar a parceria/cliente perante a MTI.",
            "Tela em acordeão alinhada ao Sydle: Dados do Cadastro (leitura), Dados da Organização "
            "(identificação, inscrições/CNAE, contas), Documentos (Etapas a incluir / Progresso / Etapas incluídas), "
            "Habilitação documental, Tributos e Encargos, Contato e Usuários. "
            "Salvar persiste e encaminha à MTI. Cancelar descarta o rascunho local.",
            [
                "Exibir formulário completo em seções colapsáveis, no padrão visual do portal Sydle/MTI.",
                "Permitir edição dos campos liberados (organização, contas, tributos, contato, anexos).",
                "Em Documentos, espelhar o back-office: Etapas a incluir, Progresso e Etapas incluídas (Etapa → Grupo → Documento).",
                "Permitir upload de anexos nos tipos exigidos e consulta de status do documento.",
                "Ao clicar em Salvar, enviar para análise da MTI e registrar solicitação no Meu painel.",
            ],
            [
                "Dados do Cadastro não podem ser alterados no portal.",
                "Removido Validar CNPJ no portal (sem consulta automática obrigatória na tela).",
                "Não existe aba/seção “Envio / aprovação”: o envio ocorre no Salvar.",
                "Documentos seguem catálogo Etapa → Grupo → Tipo definido pela MTI.",
                "Status típicos: Incompleta → Em apresentação → Completa – aguardando MTI → Aprovada/Reprovada.",
            ],
        )

        # 4. Meus dados
        self.intro_slide(
            "Portal — Meus dados",
            "Como colaborador logado, quero manter meus dados pessoais, profissionais e de contato atualizados no portal.",
            "Tela em acordeão: Pessoa (dados pessoais, documentos, informações complementares), "
            "Profissional (organização, cargo, condição, região) e Contato. "
            "Salvar atualiza o cadastro da Pessoa vinculada ao login.",
            [
                "Exibir e editar dados da Pessoa autenticada.",
                "Permitir documentos pessoais, filiação, demografia, telefones, e-mails, endereços e redes.",
                "Exibir vínculo profissional (organização, cargo, condição, região) conforme regras do perfil.",
            ],
            [
                "Acesso a Meus dados para qualquer Pessoa pré-cadastrada com login válido.",
                "Campos somente leitura (ex.: status de vínculo calculado) não são editáveis no portal.",
                "Alterações relevantes podem gerar fluxo de validação conforme política da MTI.",
            ],
        )

        # 5. Meu painel
        self.intro_slide(
            "Portal — Meu painel",
            "Como responsável ou colaborador, quero acompanhar solicitações do cadastro e responder a pedidos de ajuste da MTI.",
            "Lista no padrão Sydle (protocolo, status, solicitado em, situação, última atualização). "
            "Ao clicar na solicitação, abre o detalhe com stepper, resumo e Histórico. "
            "Quando a MTI solicita ajuste, a mensagem do histórico exibe o botão Atender, "
            "que abre o Cadastro da organização com os dados enviados.",
            [
                "Listar solicitações relacionadas ao cadastro da organização (e demais do portal).",
                "Filtrar por nome/status/situação e ordenar por atualização.",
                "Abrir detalhe com histórico de mensagens do sistema/MTI.",
                "Exibir Atender apenas na mensagem de solicitação de ajuste.",
                "Após Atender e Salvar, reenviar à MTI e atualizar status da solicitação.",
            ],
            [
                "Fluxo obrigatório: lista → detalhe/mensagens → Atender (não atalho direto na lista).",
                "Atender usa o mesmo layout do Cadastro da organização.",
                "Status de painel (ex.: Em andamento, Aguardando, Concluído) e situação textual são independentes do status documental da Organização, mas permanecem alinhados ao parecer da MTI.",
                "Enquanto aguarda MTI, novas edições ficam bloqueadas até decisão ou reabertura por ajuste.",
            ],
        )

        # 6. Cliente x Parceiro
        self.intro_slide(
            "Portal — Cliente × Parceiro",
            "Como MTI, quero o mesmo portal com escopos distintos por perfil, para Cliente e Parceiro completarem apenas o que lhes cabe.",
            "O direcionamento ocorre após disponibilização: perfil Cliente ou Parceiro. "
            "Ambos usam Meus dados e Meu painel. O Parceiro concentra documentos MIPP/etapas e dados fiscais; "
            "o Cliente completa os campos visíveis da Organização conforme a hierarquia (raiz / não raiz).",
            [
                "Direcionar o responsável ao Portal Cliente ou Portal Parceiro conforme perfil.",
                "Cliente: completar Pessoa e campos visíveis da Organização; revisar e Salvar (envia).",
                "Parceiro: completar Pessoa e Organização; anexar documentos das etapas; Salvar (envia).",
                "Aplicar visibilidade por Organização raiz / não raiz (ocultar seções sem campos).",
            ],
            [
                "Parceiro e Cliente não alteram a estrutura de etapas/grupos/tipos.",
                "Campos ocultos no back-office permanecem ocultos no portal.",
                "Validação no envio considera apenas campos visíveis e obrigatórios do perfil.",
            ],
        )

        os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
        self.prs.save(OUT_PATH)
        print(OUT_PATH)


if __name__ == "__main__":
    Deck().build()
