# -*- coding: utf-8 -*-
"""Gera documentação Fase 2 — ajustes validação 17/08 (texto original + correções em vermelho)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "exports" / "Fase2-Ajustes-Requisitos-Validacao-17ago.docx"

BLACK = RGBColor(0, 0, 0)
RED = RGBColor(192, 0, 0)


def add_run(paragraph, text: str, *, red: bool = False, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RED if red else BLACK
    run.bold = bold
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_mixed(paragraph, parts: list[tuple[str, bool]]) -> None:
    for text, is_red in parts:
        add_run(paragraph, text, red=is_red)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = BLACK


def bullet(doc: Document, parts: list[tuple[str, bool]], *, bold_label: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_label:
        add_run(p, bold_label, bold=True)
    add_mixed(p, parts)


def slide_ref(doc: Document, slides: str, action: str) -> None:
    p = doc.add_paragraph()
    add_run(p, "Slides PPTX: ", bold=True)
    add_run(p, slides)
    add_run(p, "  |  Ação no deck: ", bold=True)
    add_run(p, action, red=True)


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, "Projeto Atlas — Fase 2\n", bold=True)
    add_run(title, "Documentação de Ajustes de Requisitos\n", bold=True)
    add_run(title, "(Validação 17/08/2026 + auditoria Fase 2.pptx)\n\n")
    add_run(title, "Legenda: ", bold=True)
    add_run(title, "texto em ")
    add_run(title, "vermelho", red=True)
    add_run(title, " = alteração/correção em relação ao texto enviado ou ao PPTX atual.\n")
    add_run(title, "Texto em preto = mantido conforme enviado ou já alinhado à fonte.\n")

    doc.add_paragraph()
    heading(doc, "Introdução", 2)
    p = doc.add_paragraph()
    add_mixed(
        p,
        [
            (
                "Este documento reproduz a lista de alterações enviada (9 classes), "
                "confrontada com a validação de 17/08/2026, o protótipo atlas-prototipo "
                "e a auditoria slide a slide do arquivo ",
                False,
            ),
            ("Fase 2.pptx", False),
            (" (34 slides). ", False),
            (
                "Onde o texto enviado diverge da fonte ou do PPTX, a redação corrigida "
                "aparece em vermelho.",
                True,
            ),
        ],
    )

    # --- 1 Métrica ---
    heading(doc, "1. Classe: Métrica", 2)
    slide_ref(doc, "Slide 1", "Incluir painel Onde é usada (Parcerias, Catálogos, Produtos)")
    bullet(
        doc,
        [
            ("Nível de complexidade de implementação: ", False),
            ("Médio", True),
            (".", False),
        ],
        bold_label="",
    )
    bullet(
        doc,
        [
            (
                "Inclusão: Painel de Relacionamento (Grid): A tela de Métrica deve listar "
                "de forma dinâmica em quais parcerias, catálogos e produtos a métrica "
                "(UST, HST, USN) está vinculada.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Justificativa: Impedir a inativação ou exclusão de moedas ativas, "
                "evitando registros órfãos no banco de dados.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Complemento (PPTX/protótipo): incluir grid explícito de Catálogos "
                "(hoje só Parcerias e Produtos no protótipo). Grids somente leitura "
                "(consulta reversa). Cadastro mínimo: Nome + Ativo; descrição opcional.",
                True,
            ),
        ],
    )

    # --- 2 Tipo Cobrança ---
    heading(doc, "2. Classe: Tipo de Cobrança", 2)
    slide_ref(doc, "Slide 2", "Incluir valor Única; remover Recorrência/pro rata")
    bullet(doc, [("Excluído: Campo \"Recorrência\"", False)])
    bullet(
        doc,
        [
            (
                "Incluído: valor ",
                True,
            ),
            ("Única", True),
            (
                " na lista fechada (Mensal · Anual · Conforme homologação · Única). "
                "Usar Única quando Modelo de Venda = Perpétuo. Recorrência e pro rata "
                "permanecem no contrato (Fase 3), não nesta classe.",
                True,
            ),
        ],
    )

    # --- 3 Vertical ---
    heading(doc, "3. Classe: Categoria de Serviços", 2)
    slide_ref(doc, "Slide 3 (+ refs. 6, 15, 16, 22, 24, 25)", "Renomear; ajustar vínculo Parceria")
    bullet(
        doc,
        [
            (
                "Alterado (Nomenclatura): Renomeada para \"Vertical de Serviço de TI\" "
                "para evitar ambiguidade técnica com o \"Catálogo de Serviços\".",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Incluído: Relacionamento com Parceria: Vinculação obrigatória da Vertical "
                "diretamente à Parceria (e não ao CNPJ do parceiro), permitindo que um "
                "mesmo parceiro atue em verticais de TI distintas por meio de parcerias separadas.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção (fonte 17/08): a Vertical NÃO exige cadastro obrigatório de Parceria. "
                "O vínculo ocorre no Produto (Serviço) ao selecionar a Vertical; na tela da "
                "Vertical, exibir grid somente leitura de Parcerias que a utilizam (consulta reversa).",
                True,
            ),
        ],
    )

    # --- 4 Grupo ---
    heading(doc, "4. Classe: Grupo", 2)
    slide_ref(doc, "Slide 6 (+ 15, 16)", "Parceria (não Solução); filtro; Grupo obrigatório")
    bullet(
        doc,
        [
            (
                "Incluído: Filtro por Parceria: A listagem de grupos em tela deve ser filtrada "
                "obrigatoriamente de acordo com a Parceria selecionada para evitar sobrecarga "
                "de dados de outros parceiros.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Alterado (Obrigatoriedade): O campo Grupo passa a ser de preenchimento "
                "obrigatório no cadastro do produto.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "O campo SKU (Part Number) passa a ser de preenchimento opcional, pois "
                "fabricantes (como a Sidel) não utilizam esse padrão.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção (PPTX slide 6): Grupo deve referenciar Parceria (não Solução). "
                "Remover campo Categoria do Grupo. No backoffice MTI, filtro Parceria "
                "com opção \"Todos\" para visão consolidada.",
                True,
            ),
        ],
    )

    # --- 5 Modelo Venda ---
    heading(doc, "5. Classe: Modelo de Venda", 2)
    slide_ref(doc, "Slide 7", "Confirmar classe; incluir Perpétuo e Por UST")
    bullet(
        doc,
        [
            (
                "Incluído (item estava vazio na lista enviada): Classe mantida como cadastro "
                "separado da MTI (lista será reduzida antes da carga). Valores candidatos: "
                "Por Licença · Por Serviço · Por Pacote · Perpétuo · Por UST. "
                "Regra: Perpétuo → Tipo de Cobrança do Produto = Única. "
                "Modelo de venda (forma comercial) ≠ Tipo de Cobrança (ritmo de faturamento).",
                True,
            ),
        ],
    )

    # --- 6 Parceria ---
    heading(doc, "6. Classe: Parceria", 2)
    slide_ref(doc, "Slide 4 (+ 22, 24, 28)", "Incluir geração automática de catálogo rascunho")
    bullet(
        doc,
        [
            (
                "Incluído (Automação): Geração Automática de Catálogo: Ao salvar um novo "
                "cadastro de Parceria pela MTI, o sistema deve gerar automaticamente um "
                "registro de Catálogo em branco (status Rascunho) para posterior preenchimento "
                "do parceiro.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Complemento (fluxo): MTI cria o casulo → parceiro monta catálogo, grupos e "
                "produtos no Portal do Parceiro → envia à MTI para análise.",
                True,
            ),
        ],
    )

    # --- 7 Solução ---
    heading(doc, "7. Classe: Solução", 2)
    slide_ref(doc, "Slide 5", "Autofill Catálogo; remover Documentos de Apoio da tabela")
    bullet(
        doc,
        [
            (
                "Incluído (Automação): Autofill do Catálogo: O campo de cadastro de Catálogo "
                "no formulário de Solução deve ser calculado e preenchido de forma automática "
                "com base na parceria selecionada.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Remover da tabela de campos: Documentos de Apoio e Observação (fora do "
                "escopo mínimo Fase 2). Parceria da Solução: obrigatória.",
                True,
            ),
        ],
    )

    # --- 8 Catálogo ---
    heading(doc, "8. Classe: Catálogo", 2)
    slide_ref(doc, "Slides 9–12 (+ 21, 26)", "Remover complexidade; incluir toggle Licenciamento")
    bullet(
        doc,
        [
            (
                "Excluído: Tabela de Complexidade e Coeficientes no Catálogo: Removida por "
                "completo do nível do Catálogo.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Justificativa: Os serviços de um catálogo possuem esforços individuais. "
                "Essa parametrização deve ser tratada diretamente no nível do Produto/Serviço.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Incluído: Toggles de Propriedades: Inclusão dos toggles de Sim/Não "
                "independentes: \"É Universal?\", \"É Catálogo de Serviços?\" e "
                "\"É Catálogo de Licença?\".",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção nomenclatura: usar \"É catálogo de licenciamento?\" "
                "(não \"Licença\"). Os três toggles são independentes (podem ser Sim simultaneamente).",
                True,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Valor Unitário da Métrica: Obrigatoriedade de preenchimento do valor "
                "comercial da métrica na tabela embedded do catálogo (ex: UST de parceria = R$ 161,02).",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Remover do slide de campos (PPTX): Início/Fim de vigência, Estrutura de variação, "
                "seção Complexidades e coeficientes. Título padronizado: Catálogo da Parceria.",
                True,
            ),
        ],
    )

    # --- 9 Produto ---
    heading(doc, "9. Classe: Produto", 2)
    slide_ref(doc, "Slides 13–16 (+ 8 Dados comerciais)", "Corrigir período mínimo, coeficiente, custo")
    bullet(
        doc,
        [
            (
                "Incluído (Obrigatoriedade): Campo \"Período Mínimo\" (Vigência): "
                "Preenchimento obrigatório através de opções fechadas "
                "(12, 24, 36, 48, 60 meses ou Perpétuo).",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção (fonte 17/08): Período mínimo fica em Dados de Parceria por Produto "
                "(classe/bloco comercial separado), opções 12/24/36/48/60 meses — SEM Perpétuo. "
                "Perpétuo = Modelo de Venda + cobrança Única (não é opção de período mínimo). "
                "Remover campo Vigência dos dados comerciais.",
                True,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Excluído (Obrigatoriedade): Códigos SIAG / Proteus: O preenchimento dos códigos "
                "de integração do ERP não deve ser obrigatório no cadastro do catálogo. "
                "Passa a ser obrigatório estritamente na geração do Contrato (Fase 3).",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção: SIAG/Protheus opcionais no cadastro do Produto e Catálogo; "
                "obrigatórios no Contrato. Código imediato só N2 + USN/UST/HST; N1 na criação do contrato.",
                True,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Incluído (Cálculo Financeiro e Comportamento de Tela): Cálculo automático de "
                "Distribuição: O parceiro preenche o Custo Parceiro e a MTI aplica o Markup. "
                "O sistema calcula e exibe em tempo real a Distribuição do Parceiro e a "
                "Distribuição da MTI (ambos em %).",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção (fonte 17/08): Custo do parceiro é informado à MTI à parte — "
                "NÃO entra no CSV do parceiro nem é preenchido livremente pelo parceiro no Produto. "
                "Markup aplicado pela MTI. Percentuais parceiro/MTI em somente leitura (cálculo automático).",
                True,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Se Tipo de Produto = \"Licença\": O sistema oculta por completo os campos: "
                "Complexidade, Coeficiente da Complexidade, Peso e Quantidade da Métrica por Execução.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Se Tipo de Produto = \"Serviço\": O sistema exibe obrigatoriamente a seleção de "
                "Vertical, Complexidade, Coeficiente (modo leitura herdado da faixa do catálogo), "
                "Peso e Quantidade da Métrica por Execução.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Correção (17/08): Coeficiente NÃO é herdado de tabela/faixa do Catálogo "
                "(complexidade saiu do catálogo). Coeficiente pertence ao Produto/Serviço; "
                "pode ser somente leitura calculado pela faixa de complexidade escolhida NO ITEM.",
                True,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Herança de Oferta: O campo \"Tipo de Oferta\" (Universal vs. Individualizado) "
                "herda o valor e fica travado em modo leitura com base na configuração do "
                "Catálogo selecionado no início do formulário.",
                False,
            ),
        ],
    )
    bullet(
        doc,
        [
            (
                "Ordem do formulário (PPTX slides 13–16): Catálogo primeiro → Parceria/Solução RO → "
                "demais campos. Grupo obrigatório. Cobrança inclui Única.",
                True,
            ),
        ],
    )

    # --- Portal / BPMN summary ---
    heading(doc, "10. Portais e fluxo (complemento da auditoria PPTX)", 2)
    items = [
        (
            "Slide 19 — Portal do Parceiro",
            "Incluir: parceiro monta catálogo rascunho, cadastra Grupos, visão planilha, "
            "CSV sem custo/markup/%. Enviar à MTI; não publica/apostila.",
        ),
        (
            "Slide 21 — Backoffice MTI",
            "Incluir toggle Licenciamento na homologação. Diferenciar Homologar / Publicar / Apostilar.",
        ),
        (
            "Slides 22, 24, 28 — BPMN",
            "Parceria gera casulo automático; parceiro monta catálogo; Categoria → Vertical.",
        ),
        (
            "Slides 27, 34",
            "Reformular: parceiro monta e envia; MTI homologa/publica/apostila.",
        ),
    ]
    for label, text in items:
        p = doc.add_paragraph(style="List Bullet")
        add_run(p, f"{label}: ", bold=True)
        add_run(p, text, red=True)

    heading(doc, "Referências", 2)
    refs = [
        "Validação Fase 2 — 17/08/2026 (fonte de verdade)",
        "Fase 2.pptx — 34 slides (auditoria em exports/_pptx_fase2_audit/)",
        "Protótipo: data/subprojects/atlas-prototipo/epics/prototipo/forms.json",
        "Script de delta: scripts/_apply_fase2_17ago_prototype.py",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    main()
