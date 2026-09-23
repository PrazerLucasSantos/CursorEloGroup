# -*- coding: utf-8 -*-
"""AJUSTE 1: apagar métodos da Métrica. AJUSTE 2: inserir Dados de Parceria por Produto."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

DOC_DIR = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor"
    r"\Especificador Sydle\exports\documentador"
)
SRC = DOC_DIR / "requisitos-atlas-fase-2-fonte-validacao-de-requisitos-fase-2-17-08-2026-.pptx"
OUT = SRC
OUT_STABLE = DOC_DIR / "requisitos-atlas-fase-2-ajustado-validacao-17082026.pptx"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def t_elms(p_elm):
    return p_elm.findall(f".//{{{A_NS}}}t")


def set_first_t(p_elm, text: str) -> None:
    ts = t_elms(p_elm)
    if not ts:
        return
    ts[0].text = text
    for extra in ts[1:]:
        extra.text = ""


def set_item_runs(p_elm, code: str, desc: str) -> None:
    ts = t_elms(p_elm)
    if len(ts) >= 2:
        ts[0].text = f"{code}: "
        ts[1].text = desc
        for extra in ts[2:]:
            extra.text = ""
    elif ts:
        ts[0].text = f"{code}: {desc}"


def replace_paragraphs(cell, new_ps) -> None:
    body = cell.text_frame._txBody
    for old in list(body.findall(qn("a:p"))):
        body.remove(old)
    for p in new_ps:
        body.append(p)


def set_shape_title(shape, text: str) -> None:
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text
    body = tf._txBody
    for extra in body.findall(qn("a:p"))[1:]:
        body.remove(extra)


def set_cell_plain(cell, text: str) -> None:
    tf = cell.text_frame
    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text
    body = tf._txBody
    for extra in body.findall(qn("a:p"))[1:]:
        body.remove(extra)


def remove_shape(shape) -> None:
    el = shape._element
    el.getparent().remove(el)


def copy_slide(prs, index):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    for shp in list(dest.shapes):
        remove_shape(shp)
    sp_tree = dest.shapes._spTree
    for child in list(source.shapes._spTree):
        tag = child.tag
        if tag.endswith("}sp") or tag.endswith("}pic") or tag.endswith("}graphicFrame") or tag.endswith(
            "}grpSp"
        ) or tag.endswith("}cxnSp"):
            sp_tree.append(deepcopy(child))
    return dest


def delete_slide(prs, index) -> None:
    sld_id_lst = prs.slides._sldIdLst
    sld_id = sld_id_lst[index]
    r_id = sld_id.get(R_ID)
    prs.part.drop_rel(r_id)
    sld_id_lst.remove(sld_id)


def move_slide(prs, old_index, new_index) -> None:
    sld_id_lst = prs.slides._sldIdLst
    el = sld_id_lst[old_index]
    sld_id_lst.remove(el)
    sld_id_lst.insert(new_index, el)


def find_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Classe"):
            return sh
    return None


def find_heading(slide, needle: str):
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    return None


def find_table(slide, cols=None, first_cell=None):
    for sh in slide.shapes:
        if not sh.has_table:
            continue
        t = sh.table
        if cols is not None and len(t.columns) != cols:
            continue
        if first_cell is not None and t.cell(0, 0).text_frame.text.strip() != first_cell:
            continue
        return sh
    return None


def set_n_rows(table, n: int) -> None:
    tbl = table._tbl
    while len(table.rows) > n:
        tbl.remove(tbl.tr_lst[-1])
    while len(table.rows) < n:
        tbl.append(deepcopy(tbl.tr_lst[-1]))


CONTEXTO_LINHAS = [
    "- Um registro por combinação parceria × solução × produto.",
    "- No Produto fica só o valor unitário de venda.",
    "- Aqui ficam custo, markup, período mínimo e percentuais.",
    "- Período mínimo vem no CSV do parceiro (12 / 24 / 36 / 48 / 60 meses). Perpétuo não entra neste campo (Perpétuo é Modelo de Venda + cobrança Única).",
    "- Custo do parceiro é informado à MTI à parte — não trafega no CSV do parceiro.",
    "- A MTI aplica o markup. O sistema calcula Distribuição do Parceiro (%) e Distribuição da MTI (%) em somente leitura. Fórmula exata pendente da MTI.",
]

RFS = [
    ("RF-DP-01", "Permitir registrar dados comerciais por parceria × solução × produto."),
    ("RF-DP-02", "Exigir período mínimo comercial com opções fechadas 12 / 24 / 36 / 48 / 60 meses."),
    ("RF-DP-03", "Permitir à MTI informar custo do parceiro e markup."),
    ("RF-DP-04", "Calcular e exibir distribuição percentual parceiro e MTI em somente leitura."),
]

RNS = [
    ("RN-DP-01", "Separar dados comerciais da parceria do preço de venda do produto."),
    ("RN-DP-02", "Markup não fica no cadastro do Produto; fica neste bloco, aplicado pela MTI."),
    ("RN-DP-03", "Custo do parceiro não vai no CSV do portal do parceiro."),
    ("RN-DP-04", "Período mínimo obrigatório; sem opção Perpétuo neste campo."),
    ("RN-DP-05", "Vigência textual não é usada; o período mínimo governa o aspecto comercial."),
    ("RN-DP-06", "% parceiro + % MTI = 100; campos somente leitura (cálculo a partir do markup). Fórmula exata = pendente MTI."),
]

CAMPOS = [
    ("Parceria", "reference", "Não", "Não", "Visível", "Parceria à qual se aplicam custo/markup."),
    ("Solução", "reference", "Sim", "Não", "Visível", "Solução associada às condições comerciais."),
    ("Produto", "reference", "Sim", "Não", "Visível", "Produto ao qual se aplicam as condições."),
    (
        "Período mínimo (comercial)",
        "textOptions",
        "Sim",
        "Não",
        "Visível",
        "12 / 24 / 36 / 48 / 60 meses. Sem Perpétuo. Vem no CSV.",
    ),
    (
        "Custo do parceiro",
        "number",
        "Sim",
        "Não",
        "Visível",
        "Custo de entrada informado à MTI à parte. Não é o valor unitário do Produto. Não vai no CSV do parceiro.",
    ),
    ("Markup", "number", "Sim", "Não", "Visível", "Markup aplicado pela MTI. Não cadastrar no Produto."),
    (
        "Distribuição do parceiro",
        "number",
        "Sim",
        "Sim",
        "Visível",
        "% da receita do parceiro. Somente leitura.",
    ),
    (
        "Distribuição da MTI",
        "number",
        "Sim",
        "Sim",
        "Visível",
        "% da receita da MTI. Somente leitura.",
    ),
    ("Ativo", "boolean", "Sim", "Não", "Visível", "Controla disponibilidade do registro."),
]


def fill_narrativa(slide, header_p, item_p, body_p, slide_height: int) -> None:
    set_shape_title(find_title(slide), "Classe – Dados de Parceria por Produto")
    to_drop = []
    for sh in slide.shapes:
        if sh.has_text_frame and "campos" in sh.text_frame.text.lower():
            to_drop.append(sh)
        elif sh.has_text_frame and "métodos" in sh.text_frame.text.lower():
            to_drop.append(sh)
        elif sh.has_table:
            hdr = sh.table.cell(0, 0).text_frame.text.strip()
            if hdr in ("Campo", "Método"):
                to_drop.append(sh)
    for sh in to_drop:
        remove_shape(sh)

    grid_sh = find_table(slide, cols=2, first_cell="Necessidade")
    t = grid_sh.table
    set_cell_plain(t.cell(0, 1), (
        "Registrar custo, markup, período mínimo comercial e divisão percentual "
        "(parceiro × MTI) fora do preço de venda do produto."
    ))

    ctx_ps = []
    for line in CONTEXTO_LINHAS:
        p = deepcopy(body_p)
        set_first_t(p, line)
        ctx_ps.append(p)
    replace_paragraphs(t.cell(1, 1), ctx_ps)

    req_ps = []
    hp = deepcopy(header_p)
    set_first_t(hp, "Requisitos funcionais")
    req_ps.append(hp)
    for code, desc in RFS:
        p = deepcopy(item_p)
        set_item_runs(p, code, desc)
        req_ps.append(p)
    hp2 = deepcopy(header_p)
    set_first_t(hp2, "Regras de negócio")
    req_ps.append(hp2)
    for code, desc in RNS:
        p = deepcopy(item_p)
        set_item_runs(p, code, desc)
        req_ps.append(p)
    replace_paragraphs(t.cell(2, 1), req_ps)

    # Preenche o espaço que era campos+métodos
    bottom_margin = 256032
    grid_sh.height = Emu(int(slide_height) - int(grid_sh.top) - bottom_margin)


def fill_campos(slide) -> None:
    title = find_title(slide)
    set_shape_title(title, "Classe – Dados de Parceria por Produto (cont.)")
    heading = find_heading(slide, "campos")
    set_shape_title(heading, "Dados de Parceria por Produto – campos")
    tbl_sh = find_table(slide, cols=6, first_cell="Campo")
    old_rows = len(tbl_sh.table.rows)
    old_h = int(tbl_sh.height)
    set_n_rows(tbl_sh.table, 1 + len(CAMPOS))
    tbl_sh.height = Emu(int(old_h * (1 + len(CAMPOS)) / old_rows))
    t = tbl_sh.table
    headers = ["Campo", "Tipo", "Obrg", "S. Leitura", "Visibilidade", "Regra"]
    for c, h in enumerate(headers):
        set_cell_plain(t.cell(0, c), h)
    for i, row in enumerate(CAMPOS, start=1):
        for c, val in enumerate(row):
            set_cell_plain(t.cell(i, c), val)


def slide_blob(slide) -> str:
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if sh.has_table:
            t = sh.table
            for r in range(len(t.rows)):
                for c in range(len(t.columns)):
                    parts.append(t.cell(r, c).text_frame.text)
    return "\n".join(parts)


def slide_title(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Classe"):
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def verify(prs) -> None:
    titles = [slide_title(s) for s in prs.slides]
    assert titles[0] == "Classe – Métrica", titles[0]
    assert "Métrica (cont.)" not in titles[1], titles[1]
    blob0 = slide_blob(prs.slides[0])
    assert "Criar Parceria" not in blob0
    assert "Gerar Catálogo" not in blob0
    assert "Métrica – métodos" not in blob0
    assert "RF-MET-01" in blob0 and "RN-MET-06" in blob0
    assert "Nome" in blob0 and "Relacionamentos" in blob0

    # Modelo de Venda then DP then Produto
    modelo_i = next(i for i, t in enumerate(titles) if t == "Classe – Modelo de Venda")
    dp_i = next(i for i, t in enumerate(titles) if t == "Classe – Dados de Parceria por Produto")
    dp_cont_i = next(i for i, t in enumerate(titles) if t == "Classe – Dados de Parceria por Produto (cont.)")
    prod_i = next(i for i, t in enumerate(titles) if t == "Classe – Produto")
    assert modelo_i < dp_i < dp_cont_i < prod_i, (modelo_i, dp_i, dp_cont_i, prod_i)

    dp = slide_blob(prs.slides[dp_i])
    assert "Registrar custo, markup, período mínimo comercial" in dp
    assert "RF-DP-01" in dp and "RF-DP-04" in dp
    assert "RN-DP-01" in dp and "RN-DP-06" in dp
    assert "Perpétuo não entra neste campo" in dp

    campos = slide_blob(prs.slides[dp_cont_i])
    for name in [
        "Parceria",
        "Solução",
        "Produto",
        "Período mínimo (comercial)",
        "Custo do parceiro",
        "Markup",
        "Distribuição do parceiro",
        "Distribuição da MTI",
        "Ativo",
    ]:
        assert name in campos, name
    assert "textOptions" in campos
    print("VERIFY OK")
    print(f"  Métrica: slide 1 (métodos removidos; cont. excluído)")
    print(f"  Dados de Parceria por Produto: slides {dp_i + 1} e {dp_cont_i + 1}")
    print(f"  total slides: {len(prs.slides)}")


def main() -> None:
    prs = Presentation(str(SRC))
    n0 = len(prs.slides)
    titles0 = [slide_title(s) for s in prs.slides]
    already = any(t == "Classe – Dados de Parceria por Produto" for t in titles0)
    metrica_has_methods = "Métrica – métodos" in slide_blob(prs.slides[0]) or "Criar Parceria" in slide_blob(
        prs.slides[0]
    )

    if already and not metrica_has_methods:
        print("já ajustado; só revalidando e gravando cópia estável")
        verify(prs)
        prs.save(str(OUT_STABLE))
        print(f"cópia estável: {OUT_STABLE}")
        return

    assert n0 == 32, f"esperado 32 slides no arquivo buggy, veio {n0}"
    assert metrica_has_methods, "arquivo não tem os métodos errados na Métrica"

    metrica = prs.slides[0]
    grid = find_table(metrica, cols=2, first_cell="Necessidade").table
    header_p = deepcopy(grid.cell(2, 1).text_frame.paragraphs[0]._p)
    item_p = deepcopy(grid.cell(2, 1).text_frame.paragraphs[1]._p)
    body_p = deepcopy(grid.cell(0, 1).text_frame.paragraphs[0]._p)

    slide_a = copy_slide(prs, 0)
    slide_b = copy_slide(prs, 10)
    fill_narrativa(slide_a, header_p, item_p, body_p, prs.slide_height)
    fill_campos(slide_b)

    # Novos slides estão no fim; mover para depois de Modelo de Venda (índice 7)
    move_slide(prs, 32, 8)
    move_slide(prs, 33, 9)

    # AJUSTE 1 — Métrica: apagar métodos e o slide de continuação
    metrica = prs.slides[0]
    drop = []
    for sh in metrica.shapes:
        if sh.has_text_frame and "métodos" in sh.text_frame.text.lower():
            drop.append(sh)
        elif sh.has_table and sh.table.cell(0, 0).text_frame.text.strip() == "Método":
            drop.append(sh)
    for sh in drop:
        remove_shape(sh)
    delete_slide(prs, 1)

    prs.save(str(OUT))
    prs.save(str(OUT_STABLE))

    prs2 = Presentation(str(OUT))
    print(f"salvo: {OUT}")
    print(f"cópia estável: {OUT_STABLE}")
    print(f"slides: {n0} -> {len(prs2.slides)}")
    for i, s in enumerate(prs2.slides):
        print(f"{i + 1:02d} {slide_title(s)}")
    verify(prs2)

    # reabrir do disco para garantir persistência
    prs3 = Presentation(str(OUT))
    verify(prs3)
    prs4 = Presentation(str(OUT_STABLE))
    verify(prs4)


if __name__ == "__main__":
    main()
