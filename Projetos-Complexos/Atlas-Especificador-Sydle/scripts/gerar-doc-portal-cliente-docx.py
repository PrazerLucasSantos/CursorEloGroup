#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Gera documentação do Portal do Cliente Atlas em .docx."""

import os
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(
    ROOT,
    "docs",
    "Atlas_Portal_do_Cliente.docx",
)

NAVY = RGBColor(0x00, 0x4A, 0x8D)
GREY = RGBColor(0x47, 0x55, 0x69)
HEADER_BG = "004A8D"


def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=10, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(size)
    if white:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(10)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        shade_cell(c, HEADER_BG)
        cell_text(c, h, bold=True, size=9, white=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell_text(table.rows[ri].cells[ci], val, size=9)
    if col_widths:
        for row in table.rows:
            for ci, w in enumerate(col_widths):
                row.cells[ci].width = Inches(w)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Portal do Cliente — Atlas")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s = sub.add_run(f"Documentação de requisitos e conteúdo das telas · {date.today().strftime('%d/%m/%Y')}")
    s.font.size = Pt(11)
    s.font.color.rgb = GREY

    doc.add_paragraph()

    add_heading(doc, "1. Visão geral", 1)
    add_para(
        doc,
        "O Portal do Cliente Atlas é a interface de autoconsumo para organizações clientes da MTI. "
        "Permite montar cotações a partir dos catálogos homologados, acompanhar contratos e saldos, "
        "consultar o catálogo MTI na íntegra (com indicação do que está ou não no contrato de gestão) "
        "e abrir ou acompanhar ordens de serviço vinculadas ao contrato.",
    )
    add_para(doc, "Estado atual: protótipo interativo navegável (React + Vite), com dados mock para demonstração.", bold=True)

    add_heading(doc, "1.1 Objetivo de negócio", 2)
    add_bullets(
        doc,
        [
            "Reduzir dependência de e-mail/WhatsApp para solicitações de cotação e abertura de OS.",
            "Dar transparência ao cliente sobre vigência, saldo contratual e consumo por item.",
            "Expor somente catálogos em Produção (após homologação/publicação pela MTI).",
            "Centralizar o acompanhamento de demandas (cotação, contrato, OS) em um único painel.",
        ],
    )

    add_heading(doc, "1.2 Fontes de requisitos", 2)
    add_table(
        doc,
        ["Fonte", "Trecho / requisito"],
        [
            ("Modelagem 22/05/2026", "Cliente acessa o portal e monta cotação com base nos catálogos de produtos, licenças e serviços."),
            ("Modelagem 26/05 — 5:30–5:47 (Luis)", "Cliente visualiza solicitações, contratos, emite OS, acompanha projetos e consumo."),
            ("PEAP F2 — Visão do Cliente", "Contratos do CNPJ, saldo, vigência, consumo global e por item."),
            ("PEAP F3 iv.1", "Página para ver na íntegra o catálogo MTI e o que está/ausente do contrato de gestão."),
            ("PEAP F3 iv.1 a", "Na abertura de OS, informar catálogo de serviços e respectiva versão."),
            ("PEAP item iii.a", "Publicar (homologar) move catálogo de Homologação para Produção — visível ao cliente."),
        ],
        col_widths=[1.8, 4.7],
    )

    add_heading(doc, "1.3 Como executar o protótipo", 2)
    add_bullets(
        doc,
        [
            "Comando: npm run dev:portal-cliente",
            "URL: http://localhost:5173/portal-cliente.html",
            "Arquivos principais: src/portalCliente/PortalClienteApp.tsx, portalClienteData.ts, portal-cliente.css",
            "Entrada Vite: portal-cliente.html",
        ],
    )

    add_heading(doc, "2. Estrutura de navegação", 1)
    add_para(doc, "O portal possui barra superior fixa (azul institucional) com seis destinos principais:")
    add_table(
        doc,
        ["Menu", "Rota interna", "Descrição resumida"],
        [
            ("Página inicial", "home", "Hub de serviços e busca rápida no catálogo."),
            ("Meu painel", "painel", "Lista unificada de solicitações (cotação, contrato, OS, demanda)."),
            ("Catálogo MTI", "catalogo", "Catálogo completo em Produção com status no contrato."),
            ("Nova cotação", "cotacao", "Montagem e envio de cotação por solução e itens de catálogo."),
            ("Meus contratos", "contratos", "Contratos vigentes, saldo e consumo por item."),
            ("Ordens de serviço", "ordens-servico", "Consulta e criação de OS com catálogo/versão."),
        ],
        col_widths=[1.5, 1.2, 3.8],
    )

    add_heading(doc, "3. Telas e conteúdo", 1)

    # --- Home ---
    add_heading(doc, "3.1 Página inicial", 2)
    add_para(doc, "Tela de entrada do portal. Compõe a experiência de descoberta dos serviços.")
    add_para(doc, "Elementos:", bold=True)
    add_bullets(
        doc,
        [
            "Hero: título PORTAL DO CLIENTE — ATLAS, texto institucional e campo de busca.",
            "Busca: filtra itens do catálogo em Produção; ao digitar, redireciona para Catálogo MTI.",
            "Cards de serviço (4): Nova cotação, Meus contratos, Catálogo MTI, Ordens de serviço.",
            "Nota de rodapé com referência às fontes (Modelagem 22/05, 26/05, PEAP F3 iv.1).",
            "Barra de acessibilidade (atalhos) e botão flutuante de ajuda.",
        ],
    )

    # --- Painel ---
    add_heading(doc, "3.2 Meu painel", 2)
    add_para(
        doc,
        "Inbox do cliente para acompanhar tudo que foi solicitado ou está em andamento. "
        "Atende ao requisito de visualizar solicitações já feitas (Modelagem 26/05).",
    )
    add_para(doc, "Elementos:", bold=True)
    add_bullets(
        doc,
        [
            "Breadcrumb: Início / Meu painel.",
            "Aba ativa: Todas as solicitações.",
            "Botão Nova cotação (atalho para tela de cotação).",
            "Filtros: Nome, Status (Pendente, Em análise, Em execução, Aprovado, Recusado), Tipo (Cotação, Contrato, OS, Demanda).",
            "Tabela: Tipo | Nome | Status | Situação | Atualizado em.",
            "Estado vazio: Nenhum item encontrado quando filtros não retornam linhas.",
        ],
    )
    add_para(doc, "Comportamento interativo:", bold=True)
    add_bullets(
        doc,
        [
            "Envio de cotação ou criação de OS adiciona nova linha ao painel automaticamente.",
            "Ordenação por data de atualização (mais recente primeiro).",
        ],
    )

    # --- Cotação ---
    add_heading(doc, "3.3 Nova cotação", 2)
    add_para(
        doc,
        "Tela principal de demanda comercial pelo cliente. Equivalente ao portal de cotação descrito na "
        "Modelagem 22/05 e ao formulário atlas-portal-cotacao no especificador.",
    )
    add_para(doc, "Elementos:", bold=True)
    add_bullets(
        doc,
        [
            "Seleção da solução / linha de cotação (ex.: MTI CLOUD, MTI Autonomy, MTI DevSec.Gov, MTI Simplifica).",
            "Filtro de catálogo de origem: Produtos vigentes, Licenças comerciais ou Serviços comerciais.",
            "Tabela Catálogo disponível (Produção): origem, código Atlas, descrição, solução, valor — botão Adicionar.",
            "Tabela Itens da cotação: origem, código, descrição, quantidade editável, valor unitário, total por linha, remover.",
            "Rodapé da tabela: Valor total da cotação (soma automática).",
            "Ações: Limpar itens | Enviar cotação.",
        ],
    )
    add_para(doc, "Regras de negócio (protótipo):", bold=True)
    add_bullets(
        doc,
        [
            "Somente itens de catálogo em ambiente Produção são exibidos.",
            "É obrigatório ao menos um item para enviar.",
            "Ao enviar, gera número COT-AAAA-NNNN e registra solicitação Pendente no Meu painel.",
        ],
    )

    # --- Contratos ---
    add_heading(doc, "3.4 Meus contratos", 2)
    add_para(
        doc,
        "Visão dos contratos vinculados ao CNPJ do cliente — saldo, vigência e consumo (PEAP F2 e Modelagem 26/05).",
    )
    add_para(doc, "Elementos:", bold=True)
    add_bullets(
        doc,
        [
            "Cards por contrato: número, status (Vigente / A vencer / Encerrado), parceria, vigência, itens contratados.",
            "Barra de progresso: percentual do saldo consumido (saldo consumido / saldo total).",
            "Botão Ver itens do contrato (atalho para Catálogo MTI filtrado).",
            "Tabela Consumo por item: itens com status No contrato ou Parcial, colunas de consumo (utilizado / contratado).",
        ],
    )

    # --- Catálogo ---
    add_heading(doc, "3.5 Catálogo MTI — visão do cliente", 2)
    add_para(
        doc,
        "Implementação do requisito PEAP F3 iv.1: catálogo na íntegra com indicação do que está relacionado "
        "ou ausente do contrato de gestão.",
    )
    add_para(doc, "Elementos:", bold=True)
    add_bullets(
        doc,
        [
            "Texto explicativo: exibe apenas catálogos homologados em Produção (pós Publicar pela MTI).",
            "Filtros: busca textual, No contrato (todos / no contrato / ausente), Solução.",
            "Legenda visual: No contrato (verde), Parcial (âmbar), Ausente (cinza).",
            "Tabela: Origem | Código Atlas | Descrição | Solução | Versão | Valor unit. | No contrato.",
        ],
    )
    add_table(
        doc,
        ["Status no contrato", "Significado"],
        [
            ("No contrato", "Item previsto e ativo no contrato de gestão do cliente."),
            ("Parcial", "Item com cobertura parcial ou consumo ainda não iniciado."),
            ("Ausente", "Item existe no catálogo MTI mas não está no contrato do cliente."),
        ],
        col_widths=[1.5, 5.0],
    )

    # --- OS ---
    add_heading(doc, "3.6 Ordens de serviço", 2)
    add_para(
        doc,
        "Consulta de OS emitidas e abertura de nova OS com associação obrigatória a catálogo e versão (PEAP F3 iv.1 a).",
    )
    add_para(doc, "Listagem:", bold=True)
    add_bullets(
        doc,
        [
            "Colunas: Número, Contrato, Catálogo/versão, Descrição, Status, Valor, Abertura.",
            "Status possíveis: Rascunho, Em execução, Aguardando assinatura, Concluída.",
        ],
    )
    add_para(doc, "Modal Nova OS:", bold=True)
    add_bullets(
        doc,
        [
            "Contrato (seleção entre contratos vigentes do cliente).",
            "Catálogo e versão (obrigatório — ex.: MTI CLOUD — v3.2).",
            "Descrição da OS.",
            "Itens do contrato a consumir (checkboxes — somente itens No contrato ou Parcial).",
            "Ao criar: gera número OS-AAAA-NNNN, status Rascunho, registra no painel.",
        ],
    )

    add_heading(doc, "4. Dados e integrações previstas", 1)
    add_para(
        doc,
        "No protótipo, os dados são mock (arquivo portalClienteData.ts). Em produção, as telas consomem "
        "as entidades já modeladas no Atlas:",
    )
    add_table(
        doc,
        ["Tela do portal", "Entidade / integração Atlas"],
        [
            ("Nova cotação", "Catálogo (Produção), Produto, Licença, Serviço; método Enviar cotação → Demanda/Processo"),
            ("Meus contratos", "Contrato, itens contratados, saldo e medição (HCMX / Protheus — Fase 2/3)"),
            ("Catálogo MTI", "Catálogo publicado + cruzamento com itens do Contrato de gestão"),
            ("Ordens de serviço", "Ordem de Serviço, vínculo Contrato, versão Catálogo, integração ServiceNow (projetos)"),
            ("Meu painel", "Agregação de instâncias de Processo (cotação, proposta, contrato, OS)"),
        ],
        col_widths=[1.6, 4.9],
    )

    add_heading(doc, "5. O que está fora do escopo deste protótipo", 1)
    add_bullets(
        doc,
        [
            "Autenticação real (MT Login / Gov.br) — prevista para integração futura.",
            "Assinatura digital de OS e contratos (fluxo MTI interno + Portal de assinaturas).",
            "Integrações Protheus, ServiceNow e OpenText HCMX (dados exibidos como mock).",
            "Portal do Parceiro e painéis MTI/DAFI (telas distintas por papel).",
            "Emissão de NF/DAR e central de pagamentos (PEAP Fase 2/3).",
        ],
    )

    add_heading(doc, "6. Resumo das telas", 1)
    add_table(
        doc,
        ["#", "Tela", "Requisito principal", "Ações do usuário"],
        [
            ("1", "Página inicial", "Hub e descoberta", "Buscar, navegar para serviços"),
            ("2", "Meu painel", "Acompanhar solicitações", "Filtrar, consultar status"),
            ("3", "Nova cotação", "Modelagem 22/05", "Montar itens, enviar cotação"),
            ("4", "Meus contratos", "PEAP F2 / Modelagem 26/05", "Ver saldo, vigência, consumo"),
            ("5", "Catálogo MTI", "PEAP F3 iv.1", "Comparar catálogo × contrato"),
            ("6", "Ordens de serviço", "PEAP F3 iv.1 a", "Listar OS, criar nova OS"),
        ],
        col_widths=[0.35, 1.4, 2.0, 2.75],
    )

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    doc.save(OUT_PATH)
    print("Documento gerado:", OUT_PATH)


if __name__ == "__main__":
    build()
