# -*- coding: utf-8 -*-
"""
Gera o documento completo de requisitos Atlas (PPTX) alinhado ao protótipo.
Campos embutidos/referência aparecem inline na aba. Exclui Currículo e convite por código (backlog).

Uso:
  python scripts/gerar-requisitos-atlas-completo.py
"""
import copy
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import atlas_classes_config as acc

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FORMS_PATH = os.path.join(ROOT, "data/subprojects/atlas-prototipo/epics/prototipo/forms.json")

TEMPLATE_CANDIDATES = [
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas - Final.pptx",
    os.path.join(ROOT, "docs/entregaveis/pptx/Requisitos Atlas - Final.pptx"),
    os.path.join(ROOT, "docs/entregaveis/pptx/Atlas requisitos.pptx"),
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas.pptx",
]

OUT_PATH = (
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas - Final.pptx"
)
OUT_PATH_FALLBACK = os.path.join(ROOT, "docs/entregaveis/pptx/Requisitos Atlas - Final.pptx")

BLUE = RGBColor(0x18, 0x22, 0xDC)
GREY_HDR = RGBColor(0xD9, 0xD9, 0xD9)
DARK = RGBColor(0x27, 0x27, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xBF, 0xBF)

F_TITLE = "PP Telegraf"
F_HEAD = "PP Telegraf"
F_LABEL = "PP Telegraf SemiBold"
F_BODY = "Arial"

ROWS_PER_SLIDE = 10
TABLE_TOP = Inches(1.05)
TABLE_LEFT = Inches(0.45)
COL_W = [Inches(2.15), Inches(1.55), Inches(0.75), Inches(5.05)]
HEADERS = ["Campo", "Tipo", "Obrig.", "Observação"]

# --- exclusões acordadas ---
EXCLUDE_SECTION_IDS = {
    "sec-patlasv4proto-pes-curriculo",
    "sec-patlasv4proto-pes-curr-trabalho",
    "sec-patlasv4proto-pes-curr-habilidades",
    "sec-patlasv4proto-pes-curr-exp-acad",
    "sec-patlasv4proto-pes-curr-exp-prof",
}

EXCLUDE_FIELD_IDS = {
    "patlasv4proto-uo-convites",
    "patlasv4proto-uo-meth-gerar-convite",
}

EXCLUDE_METHOD_FORM_IDS = set()

EXCLUDE_STANDALONE_FORMS = {
    "form-patlasv4-proto-convite-cadastro",
    "form-patlasv4-proto-metodo-gerar-convite",
    "form-patlasv4-proto-pessoa-habilidade",
    "form-patlasv4-proto-pessoa-exp-academica",
    "form-patlasv4-proto-pessoa-exp-profissional",
    "form-patlasv4-proto-metodo-enviar-protheus",
}

# Overrides de narrativa (sem Currículo / convite por CPF)
CFG_OVERRIDES = {
    "pessoa": {
        "requisitos": [
            r for r in acc.CONFIG["pessoa"]["requisitos"]
            if "currículo" not in r.lower() and "curriculo" not in r.lower()
        ],
        "contexto": (
            "A Pessoa é o cadastro único de indivíduos do Atlas. Não existe cadastro separado de "
            '"Cliente": cliente é uma organização do tipo Cliente; a Pessoa representa o indivíduo. '
            "O acesso ao sistema só é habilitado após uma atribuição de cargo, que cria, nos "
            "bastidores, o Servidor e o Ocupante (a tela de Servidor não é exposta ao usuário). "
            "Parceiro e cliente preenchem principalmente dados de contato e cadastrais; a MTI usa "
            "o cadastro completo de identificação, documentos e contatos — sem aba Currículo na Fase 1."
        ),
    },
    "convite": {
        "requisitos": [
            "Pré-cadastrar CPF de colaboradores na Organização, vinculando pessoa à unidade.",
            "Permitir acesso ao Portal do Parceiro via MT Login ou Gov.br, verificando CPF contra o pré-cadastro.",
            "Exibir dados retornados pelo provedor de identidade em modo somente leitura (sem senha no Atlas).",
            "Concluir autocadastro gerando Solicitação de vínculo pendente para aprovação do gestor.",
            "Permitir ao gestor confirmar ou recusar (com motivo) a solicitação de vínculo.",
            "Criar a Pessoa vinculada à organização após confirmação do gestor.",
        ],
        "regras": [
            "Somente CPF pré-cadastrado pela MTI pode iniciar autocadastro no Portal.",
            "Autenticação delegada ao MT Login ou Gov.br — dados pessoais não são digitados no Atlas.",
            "Cancelar o autocadastro não cria solicitação de vínculo.",
            "Solicitação de vínculo nasce com status Pendente até decisão do gestor.",
            "Confirmar vínculo cria Pessoa na organização; Recusar exige motivo.",
            "Convite por código/link permanece em backlog (fora da Fase 1).",
        ],
        "metodos_nota": (
            "Fluxo primário: Organização › Pré-cadastrar acesso (CPF) → Portal do Parceiro "
            "(MT Login/Gov.br) › Solicitação de vínculo (Confirmar/Recusar). Convite por código é backlog."
        ),
    },
    "organizacao": {
        "requisitos": [
            "Cadastrar unidades dos três tipos (MTI, Parceiro, Cliente) com hierarquia por unidade pai e caminho único.",
            "Controlar ativo/inativo e substituição de unidades (reestruturação), com migração de cargos da unidade extinta.",
            "Manter, na organização raiz, dados de cadastro, contas bancárias, documentos de habilitação, tributos e contatos.",
            "Controlar a conformidade documental (grupos MIPP) com progresso por grupo e bloqueio de acesso comercial enquanto incompleta.",
            "Permitir selecionar grupos de documento e adicionar, de uma vez, todos os tipos exigidos à grade de habilitação.",
            "Pré-cadastrar CPF de colaboradores externos e cadastrar cargos a partir da própria organização.",
            "Aprovar ou solicitar ajuste na habilitação documental do parceiro (ações Aprovar cadastro / Solicitar ajuste).",
            "Registrar validação tributária (selo MTI/DAFI) na organização raiz.",
            "Exibir, em modo leitura, os cargos e as pessoas vinculados à organização.",
            "Configurar os módulos visíveis da unidade, que definem (junto com o cargo) a permissão efetiva.",
        ],
        "fluxo": (
            "Fluxo típico: cadastrar a organização raiz (Empresa = Sim) → preencher dados de "
            "cadastro, contas e tributos → no caso de Parceiro, selecionar grupos MIPP e anexar "
            "documentos → acompanhar status da habilitação documental → cadastrar cargos → "
            "pré-cadastrar CPF dos colaboradores para acesso via Portal (MT Login / Gov.br)."
        ),
    },
}

# Módulos adicionais (catálogo + fluxo comercial)
EXTRA_MODULES = [
    {
        "section_title": "Catálogo & Produto",
        "section_subtitle": "Produtos, catálogos, parcerias e precificação",
        "titulo": "Catálogo & Produto — Fase 1",
        "historia": (
            "Como gestor de catálogo, quero manter produtos, catálogos homologados e dados de "
            "parceria com precificação e vigência, para alimentar propostas e contratos com itens "
            "padronizados e versionados."
        ),
        "contexto": (
            "O módulo de Catálogo centraliza Produto (licença ou serviço), Catálogo (versão "
            "homologada de itens), Catálogo Universal (matriz de complexidade) e Dados de Parceria "
            "por produto (markup, distribuição MTI/parceiro). Entidades de apoio (Parceria, "
            "Solução, Métrica, Tipo de Cobrança etc.) são referências internas consumidas pelas "
            "telas principais. Catálogos homologados alimentam a geração de propostas."
        ),
        "requisitos": [
            "Cadastrar produtos com identificação, tipo, valores, parceria e status.",
            "Montar catálogos versionados vinculados a produtos, com homologação e publicação.",
            "Manter catálogo universal com matriz de complexidade e conversores.",
            "Registrar dados de parceria por produto (vigência, markup, distribuição de receita).",
            "Importar/exportar catálogo via CSV e publicar versões homologadas.",
        ],
        "regras": [
            "Somente catálogos homologados/publicados entram em novas propostas.",
            "Produto Universal e Individualizado controlam visibilidade no catálogo de parceiro.",
            "Versão do catálogo é obrigatória e incrementada a cada homologação.",
            "Tipos de cobrança incluem Sob Demanda, Mensal, Anual e Pro-Rata (DIRC/PEAP).",
        ],
        "forms": [
            ("form-patlasv4-proto-cat-produto", "Produto"),
            ("form-patlasv4-proto-cat-catalogo", "Catálogo"),
            ("form-patlasv4-proto-cat-universal", "Catálogo Universal"),
            ("form-patlasv4-proto-cat-dados-parceria", "Dados de Parceria por Produto"),
            ("form-patlasv4-proto-cat-tipo-cobranca", "Tipo de Cobrança"),
        ],
        "methods": [],
        "anexos": [],
    },
    {
        "section_title": "Fluxo Comercial",
        "section_subtitle": "Processos, proposta, contrato, OS, homologação e RAER",
        "titulo": "Fluxo Comercial — Fase 1",
        "historia": (
            "Como gestor comercial, quero conduzir demandas desde a proposta até contrato, ordem "
            "de serviço e entrega, com assinatura digital e integrações, para formalizar a "
            "parceria ponta a ponta."
        ),
        "contexto": (
            "O hub Processos concentra o ciclo comercial (demanda, proposta, workflow, documentos, "
            "contrato, OS, homologação/RAER). Proposta, Contrato, Ordem de Serviço e Projeto "
            "(Homologação + RAER) são telas dedicadas com métodos de preparação, revisão de blocos, "
            "envio para assinatura e integração Protheus/ServiceNow. Documentos gerados mantêm PDF "
            "imutável e hash de integridade durante assinatura."
        ),
        "requisitos": [
            "Abrir e acompanhar processos comerciais versionados por tipo (Demanda, Proposta, OS, Projetos).",
            "Montar propostas a partir de template e itens de catálogo, com blocos por produto e revisão.",
            "Gerar contratos a partir de proposta aprovada, com OS, produtos contratados e integração Protheus.",
            "Emitir ordens de serviço vinculadas a contratos, com assinatura e registro SNOW.",
            "Registrar homologação e RAER no projeto vinculado à entrega.",
        ],
        "regras": [
            "Instâncias de processo preservam a versão de processo com que foram abertas.",
            "PDF travado enquanto houver envelope de assinatura ativo.",
            "Recusa na assinatura cancela o envelope inteiro; reabrir exige novo PDF.",
            "Proposta só avança quando todos os blocos por produto estiverem aprovados.",
        ],
        "forms": [
            ("form-patlasv4-proto-processos", "Processos (hub comercial)"),
            ("form-patlasv4-proto-proposta", "Proposta"),
            ("form-patlasv4-proto-documentos", "Documentos gerados"),
            ("form-patlasv4-proto-contrato", "Contrato"),
            ("form-patlasv4-proto-emissao-ordem-servico", "Ordem de Serviço"),
            ("form-patlasv4-proto-projeto", "Projeto · Homologação + RAER"),
        ],
        "methods": [
            ("form-patlasv4-proto-metodo-preparar-proposta", "Preparar proposta"),
            ("form-patlasv4-proto-metodo-aprovar-bloco-produto", "Aprovar/recusar bloco"),
            ("form-patlasv4-proto-metodo-criar-os", "Criar OS"),
            ("form-patlasv4-proto-metodo-zerar-pedido-venda", "Zerar pedido de venda"),
        ],
        "anexos": [],
    },
]

MODULE_ORDER = [
    ("organizacao", "Organização", "Cadastro hierárquico de MTI, Parceiros e Clientes"),
    ("pessoa", "Pessoa", "Cadastro único de indivíduos — MTI, Parceiro e Cliente"),
    ("cargo", "Cargo", "Permissões, assinatura e ocupantes por função"),
    ("convite", "Cadastro Base — Convite, Portal e Vínculo", "Autocadastro via MT Login / Gov.br (pré-cadastro CPF)"),
    ("assinatura", "Assinatura (Workflow)", "Configuração de fluxos e operação de assinatura digital"),
    ("notificacao", "Notificações", "Regras de alerta por evento ou data"),
    ("tipo-documento", "Tipo de Documento", "Catálogo MIPP — configuração de vencimento e alerta"),
    ("versoes-processo", "Versões de Processo", "Versionamento sem perda processual"),
    ("template", "Template Documental", "Modelos de documento por blocos parametrizáveis"),
    ("modelo-contrato", "Modelo de Contrato", "Associação template + workflow por parceria"),
    ("dossie", "Dossiê da Entrega de Valor", "Acompanhamento de outcome — UGEPV / Paulo Macedo"),
]

METHODS_BY_ORG = [
    ("form-patlasv4-proto-metodo-cadastrar-cargo", "Cadastrar cargo"),
    ("form-patlasv4-proto-metodo-migrar-cargos-uo", "Migrar cargos (UO extinta)"),
    ("form-patlasv4-proto-metodo-adicionar-grupos-mipp", "Adicionar grupos à grade MIPP"),
    ("form-patlasv4-proto-metodo-substituir-unidade", "Substituir unidade"),
    ("form-patlasv4-proto-metodo-aprovar-cadastro-org", "Aprovar cadastro"),
    ("form-patlasv4-proto-metodo-solicitar-ajuste-org", "Solicitar ajuste no cadastro"),
]

METHODS_BY_ASSINATURA = [
    ("form-patlasv4-proto-metodo-enviar-assinatura", "Enviar para assinatura"),
    ("form-patlasv4-proto-metodo-recusar-assinatura", "Recusar assinatura"),
    ("form-patlasv4-proto-metodo-selecionar-metodo-assinatura", "Selecionar método de assinatura"),
]

METHODS_BY_PESSOA = [
    ("form-patlasv4-proto-metodo-atribuir-cargo-pessoa", "Atribuir cargo"),
    ("form-patlasv4-proto-metodo-opcoes-pessoa", "Opções (desligamento / inativar)"),
    ("form-patlasv4-proto-metodo-precadastro-acesso", "Pré-cadastrar acesso (CPF)"),
    ("form-patlasv4-proto-metodo-criar-pessoa", "Criar pessoa"),
]

with open(FORMS_PATH, encoding="utf-8") as f:
    FORMS = json.load(f)
FORM_BY_ID = {fm["id"]: fm for fm in FORMS}
LABEL_BY_ID = {
    fld["id"]: fld.get("label", fld["id"])
    for fm in FORMS
    for fld in fm.get("fields", [])
}


def resolve_template():
    for p in TEMPLATE_CANDIDATES:
        if os.path.isfile(p):
            return p
    raise FileNotFoundError("Template PPTX não encontrado: " + " | ".join(TEMPLATE_CANDIDATES))


def friendly_type(fld):
    t = fld.get("type")
    base = {
        "text": "Texto longo" if fld.get("textLong") else "Texto",
        "number": "Número",
        "decimal": "Decimal",
        "date": "Data",
        "boolean": "Sim/Não",
        "file": "Arquivo",
        "html": "HTML",
        "alert": "Aviso",
    }.get(t)
    if base:
        return base
    if t == "textOptions":
        opts = fld.get("options") or []
        base = "Opções (várias)" if fld.get("multiple") else "Opções"
        if opts:
            amostra = " · ".join(opts[:8])
            if len(opts) > 8:
                amostra += "…"
            base += f". Ex.: {amostra}"
        return base
    if t == "reference":
        ref = FORM_BY_ID.get(fld.get("linkedFormId"))
        hint = f" → {ref['name']}" if ref else ""
        return ("Seleção (várias)" if fld.get("multiple") else "Seleção") + hint
    if t == "embeddedReference":
        return "Tabela (múltiplas linhas)"
    return t or "—"


def obrig_cell(fld):
    if fld.get("readOnly"):
        return "Leit."
    if fld.get("required"):
        return "Sim"
    if fld.get("hidden"):
        return "Cond."
    return "Não"


def _rule_desc(r):
    src = LABEL_BY_ID.get(r["sourceFieldId"], r["sourceFieldId"])
    if r.get("sourceKind") == "boolean":
        val = "Sim" if r.get("expectedBoolean") else "Não"
    elif r.get("expectedOptionText"):
        val = r["expectedOptionText"]
    elif r.get("expectedOptionTexts"):
        val = ", ".join(r["expectedOptionTexts"])
    else:
        val = "—"
    cmp = "≠" if r.get("operator") == "neq" else "="
    act = {
        "show": "Exibido",
        "hide": "Oculto",
        "readonly": "Somente leitura",
        "editable": "Editável",
    }.get(r.get("action"), r.get("action", ""))
    return f"{act} quando «{src}» {cmp} «{val}»."


def visibility_for_field(form, field_id):
    if not form or not field_id:
        return []
    return [
        _rule_desc(r)
        for r in form.get("fieldVisibilityRules", [])
        if field_id in r.get("targetFieldIds", [])
    ]


def obs_cell(fld, form=None):
    if fld.get("type") == "alert":
        return fld.get("alertMessage", "")
    parts = []
    if fld.get("spec"):
        parts.append(fld["spec"].strip())
    for vr in visibility_for_field(form, fld.get("id")):
        parts.append(vr)
    if fld.get("readOnly") and fld.get("required"):
        parts.append("Preenchimento obrigatório pelo sistema.")
    elif fld.get("readOnly"):
        parts.append("Calculado ou preenchido automaticamente.")
    if fld.get("multiple"):
        parts.append("Permite múltiplos valores.")
    return " ".join(parts)


def set_fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def no_fill(cell):
    cell.fill.background()


def bottom_border(cell, rgb=LINE, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB",):
        existing = tcPr.find(qn(tag))
        if existing is not None:
            tcPr.remove(existing)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])})
        fill.append(clr)
        ln.append(fill)
        tcPr.append(ln)


def style_cell(cell, text, *, bold=False, size=10, color=DARK, font=F_BODY,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text or "").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


class DeckBuilder:
    def __init__(self, template_path):
        self.prs = Presentation(template_path)
        self.layout = next((l for l in self.prs.slide_layouts if l.name == "Todo branco"), self.prs.slide_layouts[1])
        xml_slides = self.prs.slides._sldIdLst
        for sld in list(xml_slides):
            rId = sld.get(qn("r:id"))
            try:
                self.prs.part.drop_rel(rId)
            except Exception:
                pass
            xml_slides.remove(sld)

    def add_slide(self):
        return self.prs.slides.add_slide(self.layout)

    def add_section_divider(self, title, subtitle):
        slide = self.add_slide()
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(2.6), Inches(11.5), Inches(1.1))
        tf = box.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = Pt(36)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = WHITE
        box2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.65), Inches(11.5), Inches(0.5))
        p2 = box2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(14)
        r2.font.name = F_BODY
        r2.font.color.rgb = WHITE

    def add_title(self, slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.0), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = DARK

    def add_bar_header(self, slide, text):
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), self.prs.slide_width, Inches(0.52))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BLUE
        bar.line.fill.background()
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.08), Inches(12), Inches(0.4))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.name = F_HEAD
        r.font.color.rgb = WHITE

    def add_subtitle_bar(self, slide, text, top=Inches(0.58)):
        box = slide.shapes.add_textbox(Inches(0.45), top, Inches(12), Inches(0.35))
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.name = F_HEAD
        r.font.color.rgb = BLUE

    def build_intro_slide(self, module_name, cfg):
        slide = self.add_slide()
        self.add_title(slide, f"{module_name} — Visão geral")
        tbl_shape = slide.shapes.add_table(3, 2, Inches(0.45), Inches(0.72), Inches(12.3), Inches(5.95))
        table = tbl_shape.table
        table.columns[0].width = Inches(2.1)
        table.columns[1].width = Inches(10.2)
        labels = ["História", "Contexto da História", "Requisitos e Regras"]
        for i, lab in enumerate(labels):
            c = table.cell(i, 0)
            set_fill(c, BLUE)
            style_cell(c, lab, bold=True, size=11, color=WHITE, font=F_LABEL, anchor=MSO_ANCHOR.TOP)
        style_cell(table.cell(0, 1), cfg["historia"], size=10, anchor=MSO_ANCHOR.TOP)
        style_cell(table.cell(1, 1), cfg["contexto"], size=10, anchor=MSO_ANCHOR.TOP)
        c2 = table.cell(2, 1)
        c2.vertical_anchor = MSO_ANCHOR.TOP
        tf = c2.text_frame
        tf.word_wrap = True

        def add_par(text, bold=False, size=10, bullet=False, color=DARK, first=False):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            r = p.add_run()
            r.text = ("•  " if bullet else "") + text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = F_BODY
            r.font.color.rgb = color

        add_par("Requisitos Funcionais", bold=True, size=11, color=BLUE, first=True)
        for req in cfg.get("requisitos", []):
            add_par(req, bullet=True)
        add_par("Regras de Negócio", bold=True, size=11, color=BLUE)
        for rg in cfg.get("regras", []):
            add_par(rg, bullet=True)
        for i in range(3):
            bottom_border(table.cell(i, 0))
            bottom_border(table.cell(i, 1))

    def field_row(self, fld, prefix="", form=None):
        if fld.get("id") in EXCLUDE_FIELD_IDS:
            return None
        label = fld.get("label", "")
        if prefix:
            label = f"{prefix} › {label}"
        if fld.get("type") == "alert":
            return [fld.get("label", "Aviso"), "Aviso", "Não", obs_cell(fld, form)]
        return [label, friendly_type(fld), obrig_cell(fld), obs_cell(fld, form)]

    def emit_fields(self, blocks, field_list, form, prefix=""):
        for fld in field_list:
            if fld.get("type") == "embeddedReference":
                group = fld.get("label", "Tabela")
                linked = FORM_BY_ID.get(fld.get("linkedFormId"))
                if linked:
                    for sub in linked.get("fields", []):
                        row = self.field_row(sub, prefix=group, form=form)
                        if row:
                            blocks.append(("fields", row))
                else:
                    row = self.field_row(fld, prefix=prefix, form=form)
                    if row:
                        blocks.append(("fields", row))
            else:
                row = self.field_row(fld, prefix=prefix, form=form)
                if row:
                    blocks.append(("fields", row))

    def collect_rows_for_form(self, form, screen_name):
        blocks = []
        secs = form.get("sections", [])
        flds = form.get("fields", [])

        if not secs:
            blocks.append(("bar", f"Tela: {screen_name}"))
            self.emit_fields(blocks, flds, form)
            return blocks

        tops = [s for s in secs if not s.get("parentSectionId")]
        for sec in tops:
            if sec["id"] in EXCLUDE_SECTION_IDS:
                continue
            blocks.append(("bar", f"Tela: {screen_name}   ›   Aba: {sec.get('title', '').strip()}"))
            sec_ids = {sec["id"]}
            for sub in [s for s in secs if s.get("parentSectionId") == sec["id"]]:
                if sub["id"] not in EXCLUDE_SECTION_IDS:
                    sec_ids.add(sub["id"])
            self.emit_fields(blocks, [f for f in flds if f.get("sectionId") in sec_ids], form)
        return blocks

    def _split_tab_blocks(self, blocks):
        """Cada barra azul (aba) vira um grupo independente para paginação."""
        groups = []
        current = None
        for block in blocks:
            if block[0] == "bar":
                current = [block]
                groups.append(current)
            elif current is not None:
                current.append(block)
            else:
                groups.append([("bar", "Telas e Campos"), block])
        return groups

    def _render_single_tab_pages(self, tab_blocks, slide_title):
        bar = tab_blocks[0]
        fields = [b for b in tab_blocks[1:] if b[0] == "fields"]
        if not fields:
            return
        idx = 0
        while idx < len(fields):
            slide = self.add_slide()
            self.add_bar_header(slide, "Telas e Campos")
            self.add_subtitle_bar(slide, slide_title, top=Inches(0.58))
            page = [bar, ("colhdr", HEADERS)]
            count = 0
            while idx < len(fields) and count < ROWS_PER_SLIDE:
                page.append(fields[idx])
                idx += 1
                count += 1
            self._draw_table(slide, page, top=Inches(1.0))

    def render_table_pages(self, blocks, slide_title):
        for tab_blocks in self._split_tab_blocks(blocks):
            self._render_single_tab_pages(tab_blocks, slide_title)

    def _draw_table(self, slide, page, top=TABLE_TOP):
        rows = len(page)
        height = min(Inches(5.85), Inches(0.36 * rows + 0.15))
        tbl_shape = slide.shapes.add_table(rows, 4, TABLE_LEFT, top, sum(COL_W, Emu(0)), height)
        table = tbl_shape.table
        for ci, w in enumerate(COL_W):
            table.columns[ci].width = w
        for ri, (kind, payload) in enumerate(page):
            if kind in ("bar", "subbar"):
                c0 = table.cell(ri, 0)
                c3 = table.cell(ri, 3)
                c0.merge(c3)
                set_fill(c0, BLUE if kind == "bar" else GREY_HDR)
                style_cell(c0, payload, bold=True, size=10 if kind == "bar" else 9,
                           color=WHITE if kind == "bar" else DARK, font=F_HEAD)
            elif kind == "colhdr":
                for ci, h in enumerate(payload):
                    c = table.cell(ri, ci)
                    set_fill(c, GREY_HDR)
                    style_cell(c, h, bold=True, size=9, color=DARK, font=F_HEAD)
            else:
                for ci, val in enumerate(payload):
                    c = table.cell(ri, ci)
                    no_fill(c)
                    style_cell(c, val, bold=(ci == 0), size=9, font=(F_BODY if ci != 0 else F_HEAD))
                    bottom_border(c)

    def build_method_slide(self, org_name, method_form_id, method_label):
        fm = FORM_BY_ID.get(method_form_id)
        if not fm or method_form_id in EXCLUDE_METHOD_FORM_IDS:
            return
        slide = self.add_slide()
        self.add_bar_header(slide, "Telas e Campos")
        title = f"{org_name} › Método: {method_label}"
        self.add_subtitle_bar(slide, title, top=Inches(0.58))
        intro = fm.get("metadata") or ""
        if intro:
            box = slide.shapes.add_textbox(Inches(0.45), Inches(0.95), Inches(12.2), Inches(0.55))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = intro[:500]
            r.font.size = Pt(9)
            r.font.name = F_BODY
        blocks = self.collect_rows_for_form(fm, title)
        if blocks:
            page = [("colhdr", HEADERS)]
            for kind, payload in blocks:
                if kind == "fields":
                    page.append((kind, payload))
            if len(page) > 1:
                self._draw_table(slide, page, top=Inches(1.55 if intro else 1.0))

    def build_module_from_cfg(self, module_name, cfg_key):
        cfg = copy.deepcopy(acc.CONFIG[cfg_key])
        if cfg_key in CFG_OVERRIDES:
            cfg.update(CFG_OVERRIDES[cfg_key])
        titulo = cfg["titulo"].split("—")[0].strip()
        if titulo.lower().startswith("classe "):
            titulo = titulo[7:].strip()
        self.build_intro_slide(module_name, cfg)
        form_id = cfg.get("form_id")
        if form_id and form_id not in EXCLUDE_STANDALONE_FORMS:
            main = FORM_BY_ID.get(form_id)
            if main:
                blocks = self.collect_rows_for_form(main, titulo)
                self.render_table_pages(blocks, "Telas e Campos")
        for ax in cfg.get("anexos_forms", []):
            if ax["form_id"] in EXCLUDE_STANDALONE_FORMS:
                continue
            fm = FORM_BY_ID.get(ax["form_id"])
            if not fm:
                continue
            blocks = self.collect_rows_for_form(fm, ax["titulo"])
            self.render_table_pages(blocks, f"Telas relacionadas — {ax['titulo']}")

    def build_extra_module(self, mod):
        cfg = {k: mod[k] for k in ("historia", "contexto", "requisitos", "regras") if k in mod}
        self.build_intro_slide(mod["section_title"], cfg)
        for form_id, screen in mod.get("forms", []):
            if form_id in EXCLUDE_STANDALONE_FORMS:
                continue
            fm = FORM_BY_ID.get(form_id)
            if not fm:
                continue
            blocks = self.collect_rows_for_form(fm, screen)
            self.render_table_pages(blocks, "Telas e Campos")
        for form_id, label in mod.get("methods", []):
            self.build_method_slide(mod["section_title"], form_id, label)

    def build_org_methods(self):
        for form_id, label in METHODS_BY_ORG:
            self.build_method_slide("Organização", form_id, label)

    def build_pessoa_methods(self):
        for form_id, label in METHODS_BY_PESSOA:
            self.build_method_slide("Pessoa", form_id, label)

    def build_assinatura_methods(self):
        for form_id, label in METHODS_BY_ASSINATURA:
            self.build_method_slide("Assinatura", form_id, label)

    def build_pendencias(self):
        slide = self.add_slide()
        self.add_title(slide, "ATLAS — FASE 1 · Pendências e próximos passos")
        rows = [
            ["P-01", "Lista exata dos docs do grupo Compliance (MIPP 4º grupo)", "Luis / DPMN", "Alta"],
            ["P-02", "Convite por código/link (forma secundária) — manter em backlog", "Luis", "Média"],
            ["P-03", "RAER: validar campos finais com stakeholders", "Luis", "Média"],
            ["P-04", "Credenciais API assinatura (Gov.br, Certificado Digital)", "Gabriel / Robson", "Alta"],
            ["P-05", "Dossiê de Entrega de Valor: validar estrutura das 5 abas", "Paulo Macedo", "Média"],
            ["P-06", "Senha+MFA na assinatura: decisão formal de remoção ou manutenção", "Luis", "Média"],
            ["P-07", "UO no Cargo como referência dinâmica (não textOptions estático)", "Pablo + Gabriel", "Média"],
        ]
        tbl = slide.shapes.add_table(len(rows) + 1, 4, Inches(0.45), Inches(0.85), Inches(12.3), Inches(5.5)).table
        tbl.columns[0].width = Inches(0.7)
        tbl.columns[1].width = Inches(7.5)
        tbl.columns[2].width = Inches(2.5)
        tbl.columns[3].width = Inches(1.6)
        hdrs = ["ID", "Pendência", "Responsável", "Urgência"]
        for ci, h in enumerate(hdrs):
            c = tbl.cell(0, ci)
            set_fill(c, GREY_HDR)
            style_cell(c, h, bold=True, size=10)
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                style_cell(tbl.cell(ri, ci), val, size=9)

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.prs.save(path)


def main():
    template = resolve_template()
    deck = DeckBuilder(template)

    for cfg_key, title, subtitle in MODULE_ORDER:
        deck.add_section_divider(title, subtitle)
        deck.build_module_from_cfg(title, cfg_key)
        if cfg_key == "organizacao":
            deck.build_org_methods()
        if cfg_key == "pessoa":
            deck.build_pessoa_methods()
        if cfg_key == "assinatura":
            deck.build_assinatura_methods()

    for mod in EXTRA_MODULES:
        deck.add_section_divider(mod["section_title"], mod["section_subtitle"])
        deck.build_extra_module(mod)

    deck.build_pendencias()
    try:
        deck.save(OUT_PATH)
        print("PPTX gerado:", OUT_PATH)
    except PermissionError:
        deck.save(OUT_PATH_FALLBACK)
        print("PPTX gerado (cópia — arquivo original em uso):", OUT_PATH_FALLBACK)
        print("Feche o arquivo em Documentação de requisitos e execute novamente para atualizar o destino principal.")
    print("Slides:", len(deck.prs.slides))


if __name__ == "__main__":
    main()
