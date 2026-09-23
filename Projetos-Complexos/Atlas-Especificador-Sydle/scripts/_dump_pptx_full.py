"""Full PPTX text dump for senior BA audit."""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Requisitos_Finais_Catalogo_Produtos_Atlas.pptx"
)
OUT = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle"
    r"\exports\_pptx_full_audit.json"
)
OUT_TXT = Path(
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle"
    r"\exports\_pptx_full_audit.txt"
)


def iter_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(s.shapes)


def main():
    prs = Presentation(str(PPTX))
    slides = []
    lines = [f"PPTX: {PPTX}", f"Slides: {len(prs.slides)}", ""]
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in iter_shapes(slide.shapes):
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(
                        " | ".join(
                            c.text_frame.text.strip().replace("\n", " ") for c in row.cells
                        )
                    )
        body = "\n".join(texts)
        title = texts[0][:120] if texts else ""
        slides.append({"slide": i, "title": title, "text": body})
        lines.append("=" * 70)
        lines.append(f"SLIDE {i}: {title}")
        lines.append("=" * 70)
        lines.append(body)
        lines.append("")
    OUT.write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8")
    OUT_TXT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUT_TXT} slides={len(slides)}")


if __name__ == "__main__":
    main()
