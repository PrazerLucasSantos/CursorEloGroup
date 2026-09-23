# -*- coding: utf-8 -*-
import re
from pptx import Presentation

PPTX = (
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas - Final - em andamento.pptx"
)
prs = Presentation(PPTX)

# Visão geral slides (module intro)
print("=== VISÃO GERAL POR MÓDULO ===\n")
for si, slide in enumerate(prs.slides):
    for sh in slide.shapes:
        if sh.has_text_frame and "Visão geral" in sh.text_frame.text:
            mod = ""
            for sh2 in slide.shapes:
                if sh2.has_text_frame:
                    lines = [l.strip() for l in sh2.text_frame.text.split("\n") if l.strip()]
                    if lines and not lines[0].startswith("Como ") and "Visão" not in lines[0]:
                        if len(lines[0]) < 60:
                            mod = lines[0]
            text = sh.text_frame.text[:500].replace("\n", " | ")
            print(f"Slide {si+1} [{mod}]: {text[:400]}...")
            print()

# Search terms
blob_parts = []
for slide in prs.slides:
    for sh in slide.shapes:
        if sh.has_text_frame:
            blob_parts.append(sh.text_frame.text)
        if sh.has_table:
            for ri in range(len(sh.table.rows)):
                blob_parts.append(" ".join(sh.table.cell(ri, ci).text for ci in range(len(sh.table.columns))))
blob = "\n".join(blob_parts)

searches = [
    "Sigadoc", "sigadoc", "Matrícula", "matricula",
    "Tipo de telefone", "tipo telefone", "Celular", "WhatsApp",
    "OCR", "automático", "automática",
    "OAuth", "federado", "login federado",
    "metadado", "ICP-Brasil", "ICP", "A3", "certificado digital",
    "Ordem", "sequencial", "paralelo",
    "Convite", "convite", "código", "self-service",
    "Auditoria", "Histórico de alterações", "trilha",
    "Produto vinculado", "escopo do produto", "Filial de produto",
    "Fase 2", "Fase 3", "backlog",
    "Mustache", "variável", "variáveis",
    "Regra", "evento", "disparo",
]
print("=== BUSCA TERMO A TERMO ===")
for s in searches:
    hits = len(re.findall(re.escape(s), blob, re.I))
    ctx = ""
    m = re.search(r".{0,40}" + re.escape(s) + r".{0,40}", blob, re.I)
    if m:
        ctx = m.group(0).replace("\n", " ")[:90]
    print(f"  [{hits:2d}] {s:30s} {ctx}")

# Which modules have auditoria
print("\n=== AUDITORIA POR MÓDULO (slide range) ===")
ranges = [
    (1, 23, "Organização"),
    (24, 36, "Pessoa"),
    (37, 42, "Cargo"),
    (43, 56, "Assinatura"),
    (57, 60, "Notificações"),
    (61, 63, "Tipo Doc"),
    (72, 78, "Dossiê"),
]
for a, b, name in ranges:
    part = []
    for si in range(a - 1, b):
        for sh in prs.slides[si].shapes:
            if sh.has_text_frame:
                part.append(sh.text_frame.text)
            if sh.has_table:
                for ri in range(len(sh.table.rows)):
                    part.append(" ".join(sh.table.cell(ri, ci).text for ci in range(len(sh.table.columns))))
    p = "\n".join(part)
    aud = "Auditoria" in p or "Histórico" in p or "trilha" in p.lower()
    print(f"  {name} (slides {a}-{b}): {'TEM auditoria/histórico' if aud else 'sem auditoria explícita'}")
