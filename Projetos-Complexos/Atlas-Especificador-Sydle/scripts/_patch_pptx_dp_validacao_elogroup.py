# -*- coding: utf-8 -*-
"""Atualiza só a classe Dados de Parceria por Produto nos decks EloGroup do projeto
para o texto exato da Validação 17/08/2026. Não altera outras classes."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

BASE = Path(r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\Cursor\Especificador Sydle")

TARGETS = [
    BASE / r"exports\Projeto Atlas - Fase 2 - Documentacao Final de Requisitos - Revisada.pptx",
    BASE / r"exports\_pptx_fase2_audit\Fase 2.pptx",
    BASE / r"exports\_pptx_fase2_audit\Fase2_requisitos_atual.pptx",
]

NEC = (
    "Registrar custo, markup, período mínimo comercial e divisão percentual "
    "(parceiro × MTI) fora do preço de venda do produto."
)

CONTEXTO = [
    "Um registro por combinação parceria × solução × produto.",
    "No Produto fica só o valor unitário de venda.",
    "Aqui ficam custo, markup, período mínimo e percentuais.",
    "Período mínimo vem no CSV do parceiro (12 / 24 / 36 / 48 / 60 meses). Perpétuo não entra neste campo (Perpétuo é Modelo de Venda + cobrança Única).",
    "Custo do parceiro é informado à MTI à parte — não trafega no CSV do parceiro.",
    "A MTI aplica o markup. O sistema calcula Distribuição do Parceiro (%) e Distribuição da MTI (%) em somente leitura. Fórmula exata pendente da MTI.",
]

REQ_LINES = [
    ("h", "Requisitos funcionais"),
    ("i", "RF-DP-01: Permitir registrar dados comerciais por parceria × solução × produto."),
    ("i", "RF-DP-02: Exigir período mínimo comercial com opções fechadas 12 / 24 / 36 / 48 / 60 meses."),
    ("i", "RF-DP-03: Permitir à MTI informar custo do parceiro e markup."),
    ("i", "RF-DP-04: Calcular e exibir distribuição percentual parceiro e MTI em somente leitura."),
    ("h", "Regras de negócio"),
    ("i", "RN-DP-01: Separar dados comerciais da parceria do preço de venda do produto."),
    ("i", "RN-DP-02: Markup não fica no cadastro do Produto; fica neste bloco, aplicado pela MTI."),
    ("i", "RN-DP-03: Custo do parceiro não vai no CSV do portal do parceiro."),
    ("i", "RN-DP-04: Período mínimo obrigatório; sem opção Perpétuo neste campo."),
    ("i", "RN-DP-05: Vigência textual não é usada; o período mínimo governa o aspecto comercial."),
    ("i", "RN-DP-06: % parceiro + % MTI = 100; campos somente leitura (cálculo a partir do markup). Fórmula exata = pendente MTI."),
]

CAMPOS = [
    ["Campo", "Tipo", "Obrg", "S. Leitura", "Visibilidade", "Regra"],
    ["Parceria", "reference", "Não", "Não", "Visível", "Parceria à qual se aplicam custo/markup."],
    ["Solução", "reference", "Sim", "Não", "Visível", "Solução associada às condições comerciais."],
    ["Produto", "reference", "Sim", "Não", "Visível", "Produto ao qual se aplicam as condições."],
    [
        "Período mínimo (comercial)",
        "textOptions",
        "Sim",
        "Não",
        "Visível",
        "12 / 24 / 36 / 48 / 60 meses. Sem Perpétuo. Vem no CSV.",
    ],
    [
        "Custo do parceiro",
        "number",
        "Sim",
        "Não",
        "Visível",
        "Custo de entrada informado à MTI à parte. Não é o valor unitário do Produto. Não vai no CSV do parceiro.",
    ],
    ["Markup", "number", "Sim", "Não", "Visível", "Markup aplicado pela MTI. Não cadastrar no Produto."],
    [
        "Distribuição do parceiro",
        "number",
        "Sim",
        "Sim",
        "Visível",
        "% da receita do parceiro. Somente leitura.",
    ],
    [
        "Distribuição da MTI",
        "number",
        "Sim",
        "Sim",
        "Visível",
        "% da receita da MTI. Somente leitura.",
    ],
    ["Ativo", "boolean", "Sim", "Não", "Visível", "Controla disponibilidade do registro."],
]


def clear_runs(p):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)


def set_para_text(p, text: str, *, bold=False, size_pt=10.0, font_name=None):
    clear_runs(p)
    run = p.add_run()
    run.text = text
    run.font.bold = bold
    run.font.size = Pt(size_pt)
    if font_name:
        run.font.name = font_name


def set_cell_paragraphs(cell, lines: list[tuple[str, str]], *, size_pt=10.0):
    tf = cell.text_frame
    # keep first paragraph, drop extras
    while len(tf.paragraphs) > 1:
        tf._txBody.remove(tf.paragraphs[-1]._p)
    first = True
    for kind, text in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bold = kind == "h"
        set_para_text(p, text, bold=bold, size_pt=9.0 if bold else size_pt)


def set_cell_simple(cell, text: str, *, size_pt=8.0, bold=False):
    tf = cell.text_frame
    while len(tf.paragraphs) > 1:
        tf._txBody.remove(tf.paragraphs[-1]._p)
    set_para_text(tf.paragraphs[0], text, bold=bold, size_pt=size_pt)


def set_shape_text(shape, text: str):
    tf = shape.text_frame
    while len(tf.paragraphs) > 1:
        tf._txBody.remove(tf.paragraphs[-1]._p)
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        set_para_text(p, text, bold=True, size_pt=18)


def set_n_rows(table, n: int):
    tbl = table._tbl
    while len(table.rows) > n:
        tbl.remove(tbl.tr_lst[-1])
    while len(table.rows) < n:
        tbl.append(deepcopy(tbl.tr_lst[-1]))


def find_dp_slide(prs):
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and "Dados de Parceria por Produto" in sh.text_frame.text and sh.text_frame.text.strip().startswith(
                "Classe"
            ):
                return i, s
    # fallback: any shape with that class name as title-ish
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith("Classe") and "Dados de Parceria por Produto" in sh.text_frame.text:
                return i, s
    return None, None


def patch_dp_slide(slide) -> None:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Classe") and "Dados de Parceria" in sh.text_frame.text:
            set_shape_text(sh, "Classe - Dados de Parceria por Produto")
        if sh.has_text_frame and sh.text_frame.text.strip().startswith("Dados de Parceria por Produto"):
            # heading campos — keep dash style of deck
            if "campos" in sh.text_frame.text.lower():
                set_shape_text(sh, "Dados de Parceria por Produto - campos")

    grid = None
    campos = None
    for sh in slide.shapes:
        if not sh.has_table:
            continue
        t = sh.table
        if len(t.columns) == 2 and t.cell(0, 0).text_frame.text.strip() == "Necessidade":
            grid = t
        if len(t.columns) == 6 and t.cell(0, 0).text_frame.text.strip() == "Campo":
            campos = t
    assert grid is not None and campos is not None

    set_cell_simple(grid.cell(0, 1), NEC, size_pt=10)
    set_cell_paragraphs(grid.cell(1, 1), [("b", x) for x in CONTEXTO], size_pt=10)
    set_cell_paragraphs(grid.cell(2, 1), REQ_LINES, size_pt=10)

    set_n_rows(campos, len(CAMPOS))
    for r, row in enumerate(CAMPOS):
        for c, val in enumerate(row):
            set_cell_simple(campos.cell(r, c), val, size_pt=8, bold=(r == 0))


def verify_dp(slide) -> None:
    blob = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            blob.append(sh.text_frame.text)
        if sh.has_table:
            t = sh.table
            for r in range(len(t.rows)):
                for c in range(len(t.columns)):
                    blob.append(t.cell(r, c).text_frame.text)
    text = "\n".join(blob)
    assert "proposta" not in text.lower()
    assert "Perpétuo" not in text or "Sem Perpétuo" in text or "não entra" in text.lower()
    assert "Vigência" not in text or "Vigência textual não é usada" in text
    assert "RF-DP-01" in text and "RN-DP-06" in text
    assert "60 meses" in text
    assert "Ativo" in text
    assert "Produto" in text
    # no old wrong period option as field rule
    assert "48 meses · Perpétuo" not in text


def main() -> None:
    for path in TARGETS:
        if not path.exists():
            print("SKIP missing", path)
            continue
        prs = Presentation(str(path))
        idx, slide = find_dp_slide(prs)
        if slide is None:
            print("SKIP no DP slide", path.name)
            continue
        # AJUSTE 1 N/A nestes decks (sem métodos errados na Métrica) — só DP
        patch_dp_slide(slide)
        verify_dp(slide)
        prs.save(str(path))
        # reabrir
        prs2 = Presentation(str(path))
        _, slide2 = find_dp_slide(prs2)
        verify_dp(slide2)
        print(f"OK slide {idx + 1}: {path.relative_to(BASE)}")


if __name__ == "__main__":
    main()
