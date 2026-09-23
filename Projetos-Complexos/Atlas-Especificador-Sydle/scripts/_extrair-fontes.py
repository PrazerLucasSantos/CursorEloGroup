# -*- coding: utf-8 -*-
"""Extrai texto das fontes (docx/pdf/pptx/xlsx) para arquivos .txt em docs/_fontes_txt/."""
import os, glob, traceback

SRC = os.path.join("docs", "Fontes de Dados para analise")
OUT = os.path.join("docs", "_fontes_txt")
os.makedirs(OUT, exist_ok=True)


def w(name, text):
    p = os.path.join(OUT, name + ".txt")
    with open(p, "w", encoding="utf-8") as f:
        f.write(text or "")
    print(f"{name}.txt  ({len(text or '')} chars)")


def docx_text(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def pdf_text(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    return "\n".join((pg.extract_text() or "") for pg in r.pages)


def pptx_text(path):
    from pptx import Presentation
    pr = Presentation(path)
    out = []
    for i, s in enumerate(pr.slides, 1):
        out.append(f"--- slide {i} ---")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
            if sh.has_table:
                for row in sh.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== aba: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)


handlers = {".docx": docx_text, ".pdf": pdf_text, ".pptx": pptx_text, ".xlsx": xlsx_text}
for path in sorted(glob.glob(os.path.join(SRC, "*"))):
    ext = os.path.splitext(path)[1].lower()
    if ext not in handlers:
        continue
    name = os.path.splitext(os.path.basename(path))[0]
    try:
        w(name, handlers[ext](path))
    except Exception as e:
        print(f"ERRO {name}: {e}")
        traceback.print_exc()
print("\nOK")
