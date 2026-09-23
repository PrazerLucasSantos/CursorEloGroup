# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX = Path(os.environ.get("TEMP", r"C:\Users\LucasSantos\AppData\Local\Temp")) / "Fase2-atlas.pptx"
OUT_TXT = Path(__file__).resolve().parents[1] / "exports" / "_pptx_fase2_hoje.txt"
OUT_JSON = Path(__file__).resolve().parents[1] / "exports" / "_pptx_fase2_hoje.json"


def iter_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from iter_shapes(s.shapes)
            except Exception:
                pass


def shape_xy(shape):
    try:
        return (int(shape.top or 0), int(shape.left or 0))
    except Exception:
        return (0, 0)


def dump_table(shape):
    rows = []
    for row in shape.table.rows:
        rows.append([c.text_frame.text.strip().replace("\n", " ") for c in row.cells])
    return rows


def main() -> int:
    prs = Presentation(str(PPTX))
    slides = []
    lines = [
        f"PPTX: {PPTX}",
        f"Slides: {len(prs.slides)}",
        f"Size: {PPTX.stat().st_size}",
        f"W={prs.slide_width} H={prs.slide_height}",
        "",
    ]
    for i, slide in enumerate(prs.slides, 1):
        items = []
        for shape in sorted(iter_shapes(slide.shapes), key=shape_xy):
            rec = {"name": getattr(shape, "name", ""), "type": str(shape.shape_type)}
            try:
                rec["top"] = int(shape.top or 0)
                rec["left"] = int(shape.left or 0)
            except Exception:
                pass
            if getattr(shape, "has_table", False):
                rec["table"] = dump_table(shape)
            elif getattr(shape, "has_text_frame", False):
                t = shape.text_frame.text.strip()
                if t:
                    rec["text"] = t
            if "table" in rec or "text" in rec:
                items.append(rec)
        slides.append({"slide": i, "items": items})
        lines.append("=" * 72)
        lines.append(f"SLIDE {i}")
        lines.append("=" * 72)
        for it in items:
            if "table" in it:
                lines.append("[TABELA]")
                for row in it["table"]:
                    lines.append(" | ".join(row))
            else:
                lines.append(it["text"].replace("\r", ""))
            lines.append("")
    OUT_JSON.write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8")
    OUT_TXT.write_text("\n".join(lines), encoding="utf-8")
    print(f"slides={len(slides)} txt={OUT_TXT} json={OUT_JSON}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
