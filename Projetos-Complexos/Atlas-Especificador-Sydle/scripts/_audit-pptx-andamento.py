# -*- coding: utf-8 -*-
"""Auditoria F1: Requisitos Atlas - Final - em andamento.pptx"""
import re
from pptx import Presentation

PPTX = (
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas - Final - em andamento.pptx"
)

prs = Presentation(PPTX)
print(f"SLIDES: {len(prs.slides)}\n")

slides = []
all_text = []
field_rows = []
method_slides = []

for si, slide in enumerate(prs.slides):
    title = ""
    texts = []
    tabs = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                texts.append(t)
                if "Método:" in t or "Metodo:" in t:
                    method_slides.append((si + 1, t.replace("\n", " ")[:140]))
                if not title and len(t) < 150 and not t.startswith("Campo"):
                    first = t.split("\n")[0][:100]
                    if re.match(r"^\(\d+\)", first) or "Atlas" in first or first.startswith("Tela:"):
                        title = first
        if sh.has_table:
            tbl = sh.table
            for ri in range(len(tbl.rows)):
                c0 = tbl.cell(ri, 0).text.strip()
                if c0.startswith("Tela:") and "Aba:" in c0:
                    tabs.append(c0)
                if ri > 0 and c0 and c0 not in ("Campo",) and not c0.startswith("Tela:"):
                    try:
                        row = [tbl.cell(ri, ci).text.strip() for ci in range(min(4, len(tbl.columns)))]
                        if row[0]:
                            field_rows.append(row)
                            all_text.append(" ".join(row))
                    except Exception:
                        pass
    if not title and texts:
        title = texts[0].split("\n")[0][:100]
    slides.append({"n": si + 1, "title": title, "tabs": tabs})

for s in slides:
    tab_info = f" [{len(s['tabs'])} abas]" if s["tabs"] else ""
    print(f"{s['n']:3d}. {s['title']}{tab_info}")

full = "\n".join(all_text).lower()
full_all = full + "\n" + "\n".join(s["title"].lower() for s in slides)

print("\n=== KEYWORDS F1 ===")
keywords = {
    "Organização": r"organiza",
    "Pessoa": r"pessoa|cpf|pr[eé].?cadastr",
    "Cargo": r"cargo|ocupante|suplente|substituto",
    "Tipo Documento + OCR": r"tipo.*documento|modo autom[aá]tico|ocr",
    "Template": r"template|mustache|bloco",
    "Assinatura/Workflow": r"assinatura|workflow|envelope|ged",
    "Gov.br": r"gov\.?br",
    "MT Login": r"mt login|mtlogin",
    "Notificação email": r"notifica|e-mail|email",
    "WhatsApp/SMS": r"whatsapp|sms",
    "Auditoria": r"auditoria|hist[oó]rico|trilha",
    "Versões Processo": r"vers[aã]o.*processo|versões de processo|\(6\)",
    "Dossiê EV": r"dossi[eê]|entrega de valor|\(12\)",
    "Catálogo/Produto F3": r"cat[aá]logo|produto|dirc|universal|\(13\)|\(14\)",
    "Fluxo comercial F2": r"proposta|\(16\)|contrato|\(17\)|ordem de servi",
    "Convite F2": r"convite",
    "Filiação": r"filia",
    "Visibilidade produto cargo": r"produto.*cargo|visibilidade.*produto",
    "Migrar cargos": r"migrar cargo",
    "Substituir UO": r"substituir.*unidade",
    "Aprovar cadastro org": r"aprovar cadastro",
    "Solicitar ajuste org": r"solicitar ajuste",
    "MIPP Compliance": r"compliance|mipp",
    "Senha MFA": r"senha|mfa|certificado",
    "Pro-Rata": r"pro.?rata",
    "Autogestão F2": r"autogest|portal triplo",
}

for k, pat in keywords.items():
    m = re.search(pat, full_all, re.I)
    print(f"  [{'OK' if m else '--'}] {k}")

print(f"\nCampos em tabelas: {len(field_rows)}")
print("\n=== METHOD SLIDES ===")
for n, m in method_slides:
    print(f"  slide {n}: {m}")

# F1 modules expected from slide titles
f1_modules = [
    "(01) Organização", "(02) Pessoa", "(03) Cargo", "(04) Tipo",
    "(05)", "(06) Versões", "(07) Template", "(11) Assinatura",
    "(15) Workflow", "Notifica"
]
print("\n=== F1 MODULE COVERAGE BY SLIDE TITLE ===")
for mod in [
    "(01) Organização", "(02) Pessoa", "(03) Cargo", "(04) Tipo de Documento",
    "(05)", "(06) Versões de Processo", "(07) Template", "(11) Assinatura",
    "(15) Workflow", "Notificação", "(12) Dossiê", "Catálogo", "Produto",
    "Proposta", "Contrato", "Modelo", "Processos", "Versões",
]:
    hits = [s for s in slides if mod.lower() in s["title"].lower()]
    status = f"{len(hits)} slide(s)" if hits else "AUSENTE"
    print(f"  {mod}: {status}")
