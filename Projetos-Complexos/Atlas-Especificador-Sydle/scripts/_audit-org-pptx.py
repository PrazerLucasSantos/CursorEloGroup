# -*- coding: utf-8 -*-
import re
from pptx import Presentation

PPTX = (
    r"c:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho"
    r"\Documentação de requisitos\Requisitos Atlas - Final - em andamento.pptx"
)
prs = Presentation(PPTX)
rows = []
texts = []
methods = []
for si in range(0, 23):
    for sh in prs.slides[si].shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            texts.append(t)
            if "Método:" in t or "Metodo:" in t:
                methods.append((si + 1, t.replace("\n", " ")[:120]))
        if sh.has_table:
            for ri in range(len(sh.table.rows)):
                try:
                    r = [sh.table.cell(ri, ci).text.strip() for ci in range(min(4, len(sh.table.columns)))]
                except Exception:
                    continue
                if r[0] and r[0] not in ("Campo",) and not r[0].startswith("Tela:"):
                    rows.append((si + 1, r[0]))

blob = "\n".join(texts) + "\n" + "\n".join(x[1] for x in rows)

checks = [
    ("Empresa / org raiz", r"Empresa"),
    ("Tipo MTI/Parceiro/Cliente", r"Tipo do cadastro"),
    ("Poder (cliente)", r"Poder"),
    ("Unidade pai / caminho", r"Unidade pai|Caminho hier"),
    ("Unidade substituída por", r"substitu"),
    ("Código Protheus", r"Protheus"),
    ("Dados cadastro (razão, CNAE…)", r"Razão Social|CNAE|Natureza Jur"),
    ("Contas bancárias", r"Contas banc|bancária"),
    ("MIPP — 4 grupos", r"Compliance|Habilitação Jur"),
    ("Grade documentos MIPP", r"Documentos de habilita|Grupo do documento"),
    ("Modo vencimento / OCR", r"[Aa]utom[aá]tico|Modo de vencimento"),
    ("Status habilitação", r"Status da habilita"),
    ("Progresso por grupo MIPP", r"Progresso"),
    ("Acesso comercial bloqueado", r"Acesso comercial"),
    ("Tributos (regime, ICMS)", r"Regime|Tributo|ICMS"),
    ("Validação DAFI", r"DAFI"),
    ("Contatos", r"E-mail|Telefones|Endere"),
    ("Limite / contador usuários", r"Limite de usu|Usuários cadastrados"),
    ("Convites (F2)", r"Convite"),
    ("Espelho cargos/pessoas", r"Cargos atribu|Pessoas nos cargos"),
    ("Módulos visíveis UO", r"Módulos vis"),
    ("Docs execução parceria", r"execução da parceria|Documentos de execu"),
    ("Logo org (templates)", r"[Ll]ogo"),
    ("Responsável UO", r"Responsável"),
    ("Nível (auto)", r"Nível"),
    ("Recorrência documento", r"Recorrência"),
    ("Status validação doc linha", r"Status de validação"),
    ("Auditoria / Histórico", r"Histórico|Auditoria"),
    ("Produto/Parceria (F2)", r"Produto / Parceria"),
    ("Fluxo onboarding / BPM", r"apresentação|habilitação documental|potencial parceiro"),
]

print("=== ORG PPTX slides 1-23 vs fontes ===\n")
ok_n = 0
for name, pat in checks:
    ok = bool(re.search(pat, blob, re.I))
    if ok:
        ok_n += 1
    print("  [%s] %s" % ("OK" if ok else "--", name))
print("\nCobertura: %d/%d" % (ok_n, len(checks)))

print("\n=== MÉTODOS nos slides Org ===")
for n, m in methods:
    print("  slide %d: %s" % (n, m))

expected_methods = [
    "Criar", "Editar", "Salvar", "Aprovar cadastro", "Solicitar ajuste",
    "Migrar cargos", "Substituir unidade", "Adicionar grupos", "Gerar convite",
]
print("\n=== MÉTODOS esperados (fontes) ===")
mb = blob.lower()
for m in expected_methods:
    print("  [%s] %s" % ("OK" if m.lower() in mb else "--", m))

print("\n=== REGRAS DE NEGÓCIO no texto Org ===")
rules = [
    ("RN org raiz só MIPP", r"somente org raiz|Empresa = Sim"),
    ("Parceiro bloqueado sem MIPP", r"sem habilitação|comercial bloqueado|proposta"),
    ("Notificação vencimento", r"notifica|vencimento"),
    ("Filha herda tipo", r"filha|herda|unidade pai"),
    ("Módulos não herdam", r"não herda|livre"),
]
for name, pat in rules:
    print("  [%s] %s" % ("OK" if re.search(pat, blob, re.I) else "--", name))
