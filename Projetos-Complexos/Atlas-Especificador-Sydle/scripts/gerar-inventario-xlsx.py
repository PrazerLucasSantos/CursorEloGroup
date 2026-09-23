# -*- coding: utf-8 -*-
"""Gera planilha do inventário de desenvolvimentos do Atlas (F1/F2/F3)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs/entregaveis/xlsx/Atlas_Inventario_Desenvolvimentos.xlsx")

# nome, serve, solicitante, fonte, doc, fase
DATA = [
    ("Cadastro Organizacional (MTI/Parceiro/Cliente)", "Tela central da hierarquia; mesma tela atende os três tipos por qualificador.", "Luis / Lucas", "Discovery 15/06; 12/06", "Sim", "F1"),
    ("Hierarquia por unidade-pai + caminho/nível calculado", "Árvore recursiva sem classe 'Nível Organizacional'; caminho e nível automáticos.", "Lucas", "Reunião 12/06", "Sim", "F1"),
    ("Habilitação documental MIPP (grupos, progresso, bloqueio)", "Conformidade do parceiro; bloqueia acesso comercial enquanto incompleta.", "Luis", "Reunião 19/06", "Sim", "F1"),
    ("Selo / validação tributária (DAFI)", "MTI atesta os tributos do parceiro (status + responsável + data).", "Luis", "PEAP item p", "Sim", "F1"),
    ("Substituir unidade + migrar cargos", "Reestruturação: inativa unidade e migra cargos sem mexer em banco.", "Luis", "Discovery 15/06", "Sim", "F1"),
    ("Cargo (ocupantes, pode assinar, admin/módulos)", "Funções por organização/unidade; direciona fluxo e assinatura por cargo.", "Luis / Pablo", "Reunião 10/06", "Sim", "F1"),
    ("Permissão = módulos da UO ∩ módulos do Cargo", "Permissão efetiva por interseção; Administrador = acesso total.", "Luis", "Reunião 19/06", "Sim", "F1"),
    ("Pessoa (cadastro único + Atribuir cargo → Servidor)", "Cadastro único; acesso só após atribuição de cargo (Servidor no backend).", "Lucas / Luis", "25/05; 10/06", "Sim", "F1"),
    ("Pré-cadastro de acesso (CPF) em Pessoa", "MTI cadastra o CPF; pessoa acessa depois pelo MT Login/Gov.br.", "Luis", "Reunião 23/06", "Sim", "F1"),
    ("Portal do Parceiro via MT Login / Gov.br", "Autenticação delegada; CPF bate com pré-cadastro e mostra a organização.", "Luis / Pablo", "Discovery 15/06; 23/06", "Sim", "F1"),
    ("Solicitação de vínculo (aprovação do gestor, c/ cargo)", "Inbox do gestor para confirmar/recusar o vínculo, com info do cargo.", "Lucas", "Reunião 12/06; 23/06", "Sim", "F1"),
    ("Tipo de Documento (catálogo MIPP)", "Tipos exigidos, modo de vencimento e antecedência do alerta.", "Luis", "Discovery 15/06; 19/06", "Sim", "F1"),
    ("Versões de Processo (versionamento)", "Versiona configurações sem quebrar instâncias em andamento.", "Lucas / Pablo", "01/06", "Sim", "F1"),
    ("Template documental (blocos HTML + merge fields + cláusulas)", "Monta o modelo do documento por blocos com variáveis; cláusulas reutilizáveis.", "Luis", "Reunião 01/06", "Sim", "F1"),
    ("Modelo de Contrato (template + workflow por parceria)", "Associa template e workflow padrão para gerar o contrato.", "Lucas", "Requisitos.docx", "Sim", "F1"),
    ("OCR de vencimento de documento", "Backend lê a data de vencimento no arquivo e dispara alerta.", "Luis", "Discovery 15/06", "Parcial", "F1"),
    ("Sino de notificações no header (badge)", "Indicador visual de pendências (ex.: documento faltando).", "Luis", "Discovery 15/06", "Parcial", "F1"),
    ("Catálogo & Produto (Produto, Catálogo, Universal, Dados Parceria)", "Hierarquia de 3 camadas; base comercial das propostas.", "Luis / Gabriel", "01/06; DIRC", "Sim", "F1"),
    ("Entidades de suporte do catálogo (Parceria, Solução, Métrica…)", "Domínios configuráveis (Tipo Cobrança, Categoria, Grupo, Modelo de Venda).", "Gabriel", "DIRC", "Sim", "F1"),
    ("Importar / Exportar catálogo via CSV", "Alimentar o catálogo por download/upload no portal.", "Luis", "PEAP item 2.a", "Sim", "F1"),
    ("Homologação do catálogo (2 ambientes)", "Parceiro alimenta homologação; MTI publica para produção/visão do cliente.", "Luis", "PEAP item iii.a", "Sim", "F1"),
    ("Atributos de gestão do produto (riscos, restrições, marcos…)", "Artefatos de gestão de projeto do produto de catálogo.", "Luis", "PEAP itens ee–ii", "Sim", "F1"),
    ("Processos (hub)", "Acompanha cada instância (status, etapa, responsável, histórico).", "Lucas", "Requisitos.docx", "Sim", "F1"),
    ("Proposta (blocos por produto, gerar PDF, enviar p/ assinatura)", "Monta e aprova a proposta; PDF imutável após blocos aprovados.", "Luis", "Modelagem 22/05; 19/06", "Sim", "F1"),
    ("Extrato de publicação do contrato", "Registro mandatório para validade jurídica do contrato.", "Luis", "Modelagem 22/05", "Parcial", "F1"),
    ("Contrato multi-parceiro com visibilidade segmentada", "Cada parceiro aprova só seu escopo; não vê margem/dados de outros.", "Luis", "Modelagem 22/05; PEAP", "Não", "F1"),
    ("Documentos gerados (PDF + hash de integridade)", "Evidência legal; documento assinado é imutável.", "Lucas", "Requisitos.docx", "Sim", "F1"),
    ("Contrato (modelo, vigência, assinatura)", "Formaliza a proposta aprovada; produtos e OS vinculadas.", "Lucas", "Requisitos.docx", "Sim", "F1"),
    ("Ordem de Serviço (emissão, zerar pedido)", "Autoriza execução sobre contrato vigente.", "Lucas / UGPEN", "Escopo pptx", "Sim", "F1"),
    ("Projeto (Homologação + RAER)", "Unifica termo de homologação e relatório de resultados pós-OS.", "Luis", "Discovery 15/06", "Sim", "F1"),
    ("Workflow de Assinatura (paralelo; recusa cancela envelope)", "Define quem assina por cargo/unidade; envio simultâneo.", "Luis", "Discovery 15/06", "Sim", "F1"),
    ("Envelope GED + métodos (Certificado, Gov.br, MT Login)", "Agrupa documentos e conduz a assinatura digital.", "Albérico / Gabriel", "Assinatura 02/06", "Sim", "F1"),
    ("Assinatura por senha / presencial", "Métodos adicionais citados no contrato (a decidir).", "Luis (PEAP)", "PEAP; 02/06", "Parcial", "F1"),
    ("Templates de signatários reutilizáveis", "Grupo de papéis pré-montado para o envelope.", "Albérico", "Assinatura 02/06", "Parcial", "F1"),
    ("Painel de Assinaturas Pendentes", "Caixa de entrada do signatário (Aprovar/Recusar).", "Lucas", "Requisitos.docx", "Sim", "F1"),
    ("Dossiê da Entrega de Valor (UGEPV)", "Acompanha outcome: DTIC→UGEPV→ASCOM→Monitoramento→Timeline.", "Paulo Macedo / Luis", "Discovery 15/06", "Sim", "F1"),
    ("Notificações (regras, gatilho, cascata, canais)", "Regras configuráveis; destinatários em cascata até o e-mail.", "Luis", "Discovery 15/06", "Sim", "F1"),
    ("Template de Notificação (corpo HTML)", "Modelo de corpo reutilizável com variáveis.", "Luis", "Reunião 23/06", "Sim", "F1"),
    ("Campo Código Cliente/Parceiro (Protheus) — sem API", "Campo preparado na F1; integração só em fase futura.", "Luis", "Discovery 15/06", "Sim", "F1"),
    ("Painel de Autogestão Tríplice (MTI/Parceiro/Cliente)", "Ambiente de acesso por papel com visões próprias.", "Luis", "PEAP Fase 2", "Não", "F2"),
    ("Convite por código/link (autocadastro)", "Gestor gera código com usos/validade; alternativa ao pré-cadastro.", "Luis", "Discovery 15/06", "Sim (backlog)", "F2"),
    ("Portal do Cliente (contratos, saldos, vigência, consumo)", "Cliente vê contratos do seu CNPJ, saldo e consumo por item.", "Luis", "PEAP Fase 2", "Não", "F2"),
    ("Visão do Parceiro (projetos/objetos/saldo do seu escopo)", "Parceiro vê apenas o que é responsável no contrato.", "Luis", "PEAP Fase 2", "Não", "F2"),
    ("Gestão de usuários da unidade (cliente)", "Gestor do cliente cadastra/inativa usuários da sua equipe.", "Luis", "PEAP Fase 2", "Não", "F2"),
    ("Dashboard de contratos e renovação", "Visão geral de saldos, vigência e status de renovação.", "Luis", "Modelagem 26/05", "Não", "F2"),
    ("Forecast / previsão de faturamento", "Pipeline de venda e previsão de faturamento por mês.", "Ademir", "Modelagem 22/05", "Não", "F2"),
    ("Scraping do Diário Oficial (cargos)", "Atualização automática de cargos a partir do DO.", "Luis", "Discovery 15/06", "Não", "F2"),
    ("Catálogo dinâmico na Proposta (preço travado)", "Seleção de itens do catálogo na proposta com preço derivado.", "Gabriel / Lucas", "01/06; planilha DIRC", "Parcial", "F3"),
    ("Pedido de Venda (emissão e automação)", "Emite pedido por consumo medido; recorrência mensal/anual/pró-rata.", "Luis", "PEAP Fase 2/3", "Não", "F3"),
    ("Integração Protheus (contrato → pedido → NF)", "Envia contrato assinado e identifica NF emitidas.", "Luis", "PEAP; 18/06", "Não", "F3"),
    ("Integração ServiceNow (projetos/tarefas)", "Acompanhamento de projetos técnicos (output).", "Luis", "PEAP Fase 2/3", "Não", "F3"),
    ("Integração OpenText / HCMX (ativos consumidos)", "Inventário de ativos/licenças para cobrança.", "Luis", "PEAP Fase 2/3", "Não", "F3"),
    ("Nota Fiscal + DAR (emissão, download, comprovante)", "Disponibiliza NF/DAR ao cliente e registra pagamento.", "Luis", "PEAP Fase 2/3", "Não", "F3"),
    ("Central de Pagamentos e Documentos Fiscais", "Área financeira do cliente (download, alertas, comprovação).", "Luis", "PEAP Fase 2/3", "Não", "F3"),
]

HEADERS = ["Desenvolvimento", "Para que serve", "Quem solicitou", "Fonte", "Na documentação", "Fase"]
NAVY = "0C4A6E"
FASE_FILL = {"F1": "DDEBF7", "F2": "FCE4D6", "F3": "E2D9F3"}

wb = Workbook()
ws = wb.active
ws.title = "Desenvolvimentos"

thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center", horizontal="left")
    cell.border = border

for r, row in enumerate(DATA, 2):
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = border
        cell.font = Font(size=10, bold=(c == 1))
    fase = row[5]
    ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=FASE_FILL.get(fase, "FFFFFF"))

widths = [44, 60, 18, 24, 16, 8]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{len(DATA) + 1}"

# aba resumo
ws2 = wb.create_sheet("Resumo")
ws2["A1"] = "Resumo por fase"
ws2["A1"].font = Font(bold=True, size=12, color=NAVY)
ws2.append(["Fase", "Qtde"])
for f in ("F1", "F2", "F3"):
    ws2.append([f, sum(1 for d in DATA if d[5] == f)])
ws2.append(["Total", len(DATA)])
ws2.append([])
ws2.append(["Na documentação", "Qtde"])
for k in ("Sim", "Sim (backlog)", "Parcial", "Não"):
    n = sum(1 for d in DATA if d[4] == k)
    if n:
        ws2.append([k, n])
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 10

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("Planilha gerada:", OUT)
print("Itens:", len(DATA), "| F1:", sum(1 for d in DATA if d[5] == 'F1'), "F2:", sum(1 for d in DATA if d[5] == 'F2'), "F3:", sum(1 for d in DATA if d[5] == 'F3'))
