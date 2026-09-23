# -*- coding: utf-8 -*-
"""Gera PPTX — Portal do Parceiro · Gestão de Pessoas (Fase 1)."""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "docs/entregaveis/pptx/Atlas requisitos.pptx")
OUT_PATH = os.path.join(ROOT, "exports/Documentacao_Requisitos_F1_Portal_Pessoas.pptx")

BLUE = RGBColor(0x18, 0x22, 0xDC)
GREY_HDR = RGBColor(0xD9, 0xD9, 0xD9)
DARK = RGBColor(0x27, 0x27, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xBF, 0xBF)

F_TITLE = "PP Telegraf"
F_HEAD = "PP Telegraf"
F_LABEL = "PP Telegraf SemiBold"
F_BODY = "Arial"

COL_W_4 = [Inches(2.5), Inches(2.0), Inches(1.1), Inches(4.8)]
COL_W_3 = [Inches(3.0), Inches(2.5), Inches(5.0)]
HEADERS_4 = ["Campo", "Tipo", "Regras de Campo", "Observação"]
HEADERS_3 = ["Campo", "Tipo", "Regras e validação"]
ROWS_PER_SLIDE = 9


def set_fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def no_fill(cell):
    cell.fill.background()


def bottom_border(cell, rgb=LINE, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB",):
        existing = tcPr.find(qn(tag))
        if existing is not None:
            tcPr.remove(existing)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])})
        fill.append(clr)
        ln.append(fill)
        tcPr.append(ln)


def style_cell(cell, text, *, bold=False, size=10, color=DARK, font=F_BODY,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text or "").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        self.layout = next(
            (l for l in self.prs.slide_layouts if l.name == "Todo branco"),
            self.prs.slide_layouts[1],
        )
        xml_slides = self.prs.slides._sldIdLst
        for sld in list(xml_slides):
            rId = sld.get(qn("r:id"))
            try:
                self.prs.part.drop_rel(rId)
            except Exception:
                pass
            xml_slides.remove(sld)

    def add_slide(self):
        return self.prs.slides.add_slide(self.layout)

    def section_divider(self, title, subtitle):
        slide = self.add_slide()
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
        for y, txt, sz, bold in (
            (2.5, title, 34, True),
            (3.55, subtitle, 14, False),
        ):
            box = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(11.5), Inches(0.9))
            p = box.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.name = F_TITLE if bold else F_BODY
            r.font.color.rgb = WHITE

    def add_title(self, slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.0), Inches(0.45))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = DARK

    def intro_slide(self, title, historia, contexto, requisitos, regras):
        slide = self.add_slide()
        self.add_title(slide, title)
        tbl = slide.shapes.add_table(3, 2, Inches(0.45), Inches(0.72), Inches(12.3), Inches(5.95)).table
        tbl.columns[0].width = Inches(2.1)
        tbl.columns[1].width = Inches(10.2)
        for i, lab in enumerate(["História", "Contexto da História", "Requisitos e Regras"]):
            set_fill(tbl.cell(i, 0), BLUE)
            style_cell(tbl.cell(i, 0), lab, bold=True, size=11, color=WHITE, font=F_LABEL, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(0, 1), historia, size=10, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(1, 1), contexto, size=10, anchor=MSO_ANCHOR.TOP)
        c2 = tbl.cell(2, 1)
        c2.vertical_anchor = MSO_ANCHOR.TOP
        tf = c2.text_frame
        tf.word_wrap = True

        def par(text, *, bold=False, bullet=False, first=False, color=DARK):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            r = p.add_run()
            r.text = ("•  " if bullet else "") + text
            r.font.size = Pt(10)
            r.font.bold = bold
            r.font.name = F_BODY
            r.font.color.rgb = color

        par("Requisitos Funcionais", bold=True, color=BLUE, first=True)
        for req in requisitos:
            par(req, bullet=True)
        par("Regras de Negócio", bold=True, color=BLUE)
        for rg in regras:
            par(rg, bullet=True)
        for i in range(3):
            bottom_border(tbl.cell(i, 0))
            bottom_border(tbl.cell(i, 1))

    def action_intro(self, action_title, historia, contexto, requisitos, regras):
        self.intro_slide(action_title, historia, contexto, requisitos, regras)

    def proto_placeholder(self, label):
        slide = self.add_slide()
        self.add_title(slide, f"Protótipo — {label}")
        box = slide.shapes.add_textbox(Inches(0.6), Inches(2.8), Inches(11.5), Inches(1.0))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = "Captura de tela do protótipo — inserir na entrega."
        r.font.size = Pt(14)
        r.font.name = F_BODY
        r.font.color.rgb = DARK

    def _draw_table(self, slide, page, *, cols, col_w, top=Inches(1.05)):
        rows = len(page)
        height = min(Inches(5.85), Inches(0.36 * rows + 0.15))
        tbl_shape = slide.shapes.add_table(rows, cols, Inches(0.45), top, sum(col_w, Emu(0)), height)
        table = tbl_shape.table
        for ci, w in enumerate(col_w):
            table.columns[ci].width = w
        for ri, (kind, payload) in enumerate(page):
            if kind in ("bar", "subbar"):
                c0 = table.cell(ri, 0)
                c_last = table.cell(ri, cols - 1)
                c0.merge(c_last)
                set_fill(c0, BLUE if kind == "bar" else GREY_HDR)
                style_cell(
                    c0,
                    payload,
                    bold=True,
                    size=11 if kind == "bar" else 10,
                    color=WHITE if kind == "bar" else DARK,
                    font=F_HEAD,
                )
            elif kind == "colhdr":
                for ci, h in enumerate(payload):
                    set_fill(table.cell(ri, ci), GREY_HDR)
                    style_cell(table.cell(ri, ci), h, bold=True, size=10, color=DARK, font=F_HEAD)
            else:
                for ci, val in enumerate(payload):
                    no_fill(table.cell(ri, ci))
                    style_cell(
                        table.cell(ri, ci),
                        val,
                        bold=(ci == 0),
                        size=9,
                        font=F_HEAD if ci == 0 else F_BODY,
                    )
                    bottom_border(table.cell(ri, ci))

    def field_slides(self, slide_title, bar_label, rows, *, three_cols=False):
        headers = HEADERS_3 if three_cols else HEADERS_4
        cols = 3 if three_cols else 4
        col_w = COL_W_3 if three_cols else COL_W_4
        i = 0
        while i < len(rows):
            slide = self.add_slide()
            self.add_title(slide, slide_title)
            page = [("bar", bar_label), ("colhdr", headers)]
            count = 0
            while i < len(rows) and count < ROWS_PER_SLIDE:
                page.append(("fields", rows[i]))
                i += 1
                count += 1
            self._draw_table(slide, page, cols=cols, col_w=col_w)


# ---------- conteúdo ----------

VISAO = {
    "historia": (
        "Como responsável pela organização parceira, quero pré-cadastrar colaboradores e "
        "acompanhar seus dados no portal, para que cada pessoa conclua o cadastro com "
        "identidade verificada e a MTI valide o vínculo antes de liberar o acesso."
    ),
    "contexto": (
        "Somente a pessoa indicada no campo Responsável do cadastro da organização acessa "
        "a aba Pessoas e executa o pré-cadastro. O formulário é o mesmo do back-office "
        "(03.m4). O colaborador recebe e-mail, entra no portal, informa o CPF e completa "
        "a aba Meus dados (dados do MT Login/Gov.br em leitura). Ao finalizar, o cadastro "
        "vai para análise da MTI. O colaborador pode solicitar ajuste; o responsável "
        "edita apenas situação ativa e cargo."
    ),
    "requisitos": [
        "Restringir a aba Pessoas ao Responsável da organização.",
        "Reutilizar o formulário de pré-cadastro do back-office no portal.",
        "Enviar e-mail ao colaborador pré-cadastrado com link de acesso.",
        "Permitir ao colaborador finalizar o cadastro informando CPF e autenticando-se.",
        "Encaminhar à MTI após a conclusão pelo colaborador.",
        "Exibir Meus dados para o colaborador e Pessoas para o responsável.",
        "Permitir ao responsável editar somente Ativo e Cargo.",
        "Disponibilizar Solicitar ajuste em Meus dados, notificando responsável e MTI.",
        "Manter na MTI as ações Criar, Editar e Pré-cadastrar no back-office.",
    ],
    "regras": [
        "RN-PORT-PES-01: Aba Pessoas visível apenas se o usuário logado = Responsável da org.",
        "RN-PORT-PES-02: Campos do pré-cadastro idênticos entre portal e (03.m4).",
        "RN-PORT-PES-03: E-mail obrigatório no pré-cadastro para notificação.",
        "RN-PORT-PES-04: Colaborador não altera Ativo nem Cargo.",
        "RN-PORT-PES-05: Responsável não altera dados vindos do MT Login/Gov.br.",
        "RN-PORT-PES-06: Finalizar cadastro gera pendência para MTI (status Aguardando análise).",
        "RN-PORT-PES-07: Solicitar ajuste exige texto; notifica Responsável e equipe MTI.",
        "RN-PORT-PES-08: CPF do portal deve coincidir com o pré-cadastro.",
    ],
}

ABA_MEUS_DADOS = [
    ["CPF", "Texto", "Somente leitura", "Informado na entrada; deve bater com o pré-cadastro."],
    ["Nome completo", "Texto", "Somente leitura", "Retornado pelo MT Login / Gov.br."],
    ["Data de nascimento", "Data", "Somente leitura", "Provedor de identidade."],
    ["Estado civil", "Texto", "Somente leitura", "Provedor de identidade."],
    ["Nome social", "Texto", "Somente leitura", "Quando disponível no provedor."],
    ["Filiação", "Texto", "Somente leitura", "Provedor de identidade."],
    ["Celular", "Texto", "Somente leitura", "Provedor de identidade."],
    ["E-mail", "Texto", "Somente leitura", "Provedor ou pré-cadastro."],
    ["CEP", "Texto", "Somente leitura", "Endereço do provedor."],
    ["Logradouro", "Texto", "Somente leitura", ""],
    ["Número", "Texto", "Somente leitura", ""],
    ["Complemento", "Texto", "Somente leitura", ""],
    ["Bairro", "Texto", "Somente leitura", ""],
    ["Cidade (Município)", "Texto", "Somente leitura", ""],
    ["UF (Estado)", "Texto", "Somente leitura", ""],
    ["Foto", "Arquivo", "Somente leitura", "Quando disponível no provedor."],
    ["Organização vinculada", "Texto", "Somente leitura", "Definida no pré-cadastro."],
    ["Cargo", "Seleção", "Somente leitura", "Definido pelo responsável; colaborador não edita."],
    ["Ativo", "Sim/Não", "Somente leitura", "Controlado pelo responsável."],
    ["Solicitar ajuste", "Texto longo", "Editável", "Descrição do que precisa ser corrigido."],
]

ABA_PESSOAS_LISTA = [
    ["Nome", "Texto", "Somente leitura", "Lista de colaboradores da organização."],
    ["CPF", "Texto", "Somente leitura", ""],
    ["E-mail", "Texto", "Somente leitura", "Informado no pré-cadastro."],
    ["Cargo", "Seleção", "Editável (responsável)", "Alterável na ação Editar pessoa."],
    ["Condição", "Opções", "Somente leitura", "Titular, Substituto ou Suplente."],
    ["Ativo", "Sim/Não", "Editável (responsável)", "Alterável na ação Editar pessoa."],
    ["Status do cadastro", "Opções", "Somente leitura", "Pré-cadastrado · Em preenchimento · Aguardando MTI · Ativo · Ajuste solicitado."],
    ["Origem", "Opções", "Somente leitura", "MTI ou Responsável."],
]

PRECADASTRO_CAMPOS = [
    ["CPF", "Texto", "Obrigatório", "Documento principal; único no cadastro."],
    ["E-mail", "Texto", "Obrigatório", "Destino da notificação com link de acesso."],
    ["Organização", "Texto", "Somente leitura", "Parceria da sessão."],
    ["Unidade organizacional", "Seleção", "Opcional", "Filtrada pela organização."],
    ["Cargo", "Seleção", "Opcional", "Sugerido para o vínculo."],
    ["Condição", "Opções (uma)", "Opcional", "Titular, Substituto ou Suplente."],
    ["Região de atuação", "Opções (várias)", "Opcional", "Padrão Todos."],
]

ACTIONS = [
    {
        "title": "Portal — Ação: Pré-cadastrar pessoa",
        "proto": "Pré-cadastrar pessoa",
        "historia": (
            "Como responsável da parceria, quero pré-cadastrar um colaborador com os mesmos "
            "dados do back-office, para que ele receba o convite por e-mail e conclua o cadastro."
        ),
        "contexto": (
            "Disponível na aba Pessoas (portal) e no back-office (03.m4). Após confirmar, "
            "o sistema envia e-mail ao endereço informado e registra status Pré-cadastrado."
        ),
        "requisitos": [
            "Validar CPF e e-mail obrigatórios.",
            "Pré-preencher a organização da sessão.",
            "Enviar notificação por e-mail com link do portal.",
            "Registrar operador, data e organização de vínculo.",
        ],
        "regras": [
            "Executável somente pelo Responsável da organização (portal) ou perfil MTI (back-office).",
            "Não cria acesso ativo até conclusão pelo colaborador e análise da MTI.",
        ],
        "fields_bar": "Ação: Pré-cadastrar pessoa",
        "fields": PRECADASTRO_CAMPOS,
    },
    {
        "title": "Portal — Ação: Editar pessoa",
        "proto": "Editar pessoa (responsável)",
        "historia": (
            "Como responsável da parceria, quero atualizar a situação e o cargo de um "
            "colaborador já vinculado, sem alterar os dados pessoais de identidade."
        ),
        "contexto": (
            "Abre o registro a partir da aba Pessoas. Demais campos permanecem somente leitura "
            "para o responsável."
        ),
        "requisitos": [
            "Permitir alterar Ativo e Cargo.",
            "Bloquear edição de CPF, e-mail de identidade e dados do provedor.",
            "Registrar trilha de alteração.",
        ],
        "regras": [
            "Somente o Responsável da organização executa no portal.",
            "Inativar (Ativo = Não) impede acesso sem excluir o histórico.",
        ],
        "fields_bar": "Ação: Editar pessoa (responsável)",
        "fields": [
            ["Ativo", "Sim/Não", "Obrigatório", "Situação do vínculo na organização."],
            ["Cargo", "Seleção", "Opcional", "Lista filtrada por organização e perfil Parceiro."],
        ],
    },
    {
        "title": "Portal — Ação: Finalizar cadastro",
        "proto": "Finalizar cadastro (colaborador)",
        "historia": (
            "Como colaborador pré-cadastrado, quero informar meu CPF, autenticar-me e "
            "confirmar meus dados, para concluir o vínculo com a organização parceira."
        ),
        "contexto": (
            "Fluxo da aba Meus dados após o link recebido por e-mail. O colaborador entra "
            "com MT Login ou Gov.br; os dados pessoais são preenchidos automaticamente."
        ),
        "requisitos": [
            "Validar CPF contra o pré-cadastro.",
            "Autenticar via MT Login ou Gov.br.",
            "Exibir dados do provedor em leitura.",
            "Ao confirmar, encaminhar cadastro à MTI.",
            "Notificar responsável e MTI sobre nova submissão.",
        ],
        "regras": [
            "CPF divergente do pré-cadastro bloqueia a conclusão.",
            "Não há senha criada no Atlas.",
            "Status passa a Aguardando análise MTI.",
        ],
        "fields_bar": "Ação: Finalizar cadastro",
        "fields": [
            ["CPF", "Texto", "Obrigatório", "Digitado pelo colaborador na primeira tela."],
            ["Situação do CPF", "Opções", "Somente leitura", "Encontrado no pré-cadastro ou Não encontrado."],
            ["Organização vinculada", "Texto", "Somente leitura", "Organização do pré-cadastro."],
            ["Dados do provedor", "Vários", "Somente leitura", "Nome, nascimento, contato, endereço, foto."],
            ["Confirmar cadastro", "Botão", "—", "Submete à MTI."],
            ["Cancelar", "Botão", "—", "Encerra sem submeter."],
        ],
    },
    {
        "title": "Portal — Ação: Solicitar ajuste",
        "proto": "Solicitar ajuste (colaborador)",
        "historia": (
            "Como colaborador, quero informar o que está incorreto no meu cadastro, para "
            "que o responsável e a MTI corrijam os dados."
        ),
        "contexto": (
            "Campo na aba Meus dados. Gera notificação ao Responsável da organização e à "
            "equipe MTI; status do cadastro indica Ajuste solicitado."
        ),
        "requisitos": [
            "Exigir texto descritivo.",
            "Notificar Responsável e MTI.",
            "Registrar data e autor da solicitação.",
        ],
        "regras": [
            "Não altera automaticamente os dados — apenas abre pendência de correção.",
            "Colaborador continua sem editar Ativo e Cargo.",
        ],
        "fields_bar": "Ação: Solicitar ajuste",
        "fields": [
            ["Solicitar ajuste", "Texto longo", "Obrigatório", "Descreva o que precisa ser corrigido."],
            ["Status após envio", "Texto", "Somente leitura", "Ajuste solicitado."],
        ],
    },
    {
        "title": "Back-office — Ação: Analisar cadastro (MTI)",
        "proto": "Analisar cadastro MTI",
        "historia": (
            "Como analista MTI, quero aprovar ou solicitar ajuste no cadastro enviado pelo "
            "colaborador, para formalizar o vínculo na base do Atlas."
        ),
        "contexto": (
            "Inbox ou fila de cadastros com status Aguardando análise MTI, gerados ao "
            "finalizar o cadastro no portal."
        ),
        "requisitos": [
            "Exibir dados do pré-cadastro, do provedor e do responsável.",
            "Aprovar: criar/atualizar Pessoa e liberar vínculo.",
            "Solicitar ajuste: devolver com motivo ao responsável e colaborador.",
            "Registrar parecer e operador MTI.",
        ],
        "regras": [
            "Somente perfil MTI executa.",
            "Aprovação exige CPF validado e dados mínimos completos.",
        ],
        "fields_bar": "Ação: Analisar cadastro (MTI)",
        "fields": [
            ["Resumo do cadastro", "Texto", "Somente leitura", "Nome, CPF, org, cargo, responsável."],
            ["Novo status", "Opções", "Somente leitura", "Aprovado ou Ajuste solicitado."],
            ["Observação / motivo", "Texto longo", "Obrigatório se ajuste", "Enviado ao responsável e colaborador."],
            ["Anexo", "Arquivo", "Opcional", "Parecer ou evidência interna."],
        ],
    },
]

BACKOFFICE_PESSOA = {
    "title": "Back-office — (03) Pessoa — Ações existentes",
    "historia": (
        "Como analista MTI, quero manter o cadastro de pessoas no back-office, "
        "incluindo pré-cadastro direto quando o fluxo não passa pelo portal do responsável."
    ),
    "contexto": (
        "A classe (03) Pessoa permanece no back-office com cadastro completo. "
        "O pré-cadastro (03.m4) é compartilhado com o portal do responsável."
    ),
    "requisitos": [
        "Criar pessoa com dados mínimos e vínculo.",
        "Editar todos os campos permitidos ao perfil MTI.",
        "Pré-cadastrar acesso (CPF) com os mesmos campos do portal.",
    ],
    "regras": [
        "MTI publica direto no back-office, sem fila de aprovação.",
        "Parceiro (responsável) usa o portal; não edita a classe (03) no back-office.",
    ],
}


def main():
    deck = Deck()
    deck.section_divider(
        "Portal do Parceiro",
        "Gestão de Pessoas — Responsável, Colaborador e MTI",
    )

    deck.intro_slide(
        "Portal — Gestão de Pessoas — Visão geral",
        VISAO["historia"],
        VISAO["contexto"],
        VISAO["requisitos"],
        VISAO["regras"],
    )

    deck.field_slides(
        "Telas e Campos",
        "Portal do Parceiro › Aba: Meus dados (colaborador)",
        ABA_MEUS_DADOS,
    )
    deck.field_slides(
        "Telas e Campos",
        "Portal do Parceiro › Aba: Pessoas (somente Responsável)",
        ABA_PESSOAS_LISTA,
    )

    deck.intro_slide(
        BACKOFFICE_PESSOA["title"],
        BACKOFFICE_PESSOA["historia"],
        BACKOFFICE_PESSOA["contexto"],
        BACKOFFICE_PESSOA["requisitos"],
        BACKOFFICE_PESSOA["regras"],
    )

    deck.field_slides(
        "Telas e Campos",
        "Formulário compartilhado — Pré-cadastrar pessoa (portal e 03.m4)",
        PRECADASTRO_CAMPOS,
    )

    for act in ACTIONS:
        deck.action_intro(
            act["title"],
            act["historia"],
            act["contexto"],
            act["requisitos"],
            act["regras"],
        )
        deck.field_slides("Telas e Campos", act["fields_bar"], act["fields"])
        deck.proto_placeholder(act["proto"])

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    deck.prs.save(OUT_PATH)
    print("PPTX gerado:", OUT_PATH)
    print("Slides:", len(deck.prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
