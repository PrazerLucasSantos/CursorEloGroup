# -*- coding: utf-8 -*-
"""Extract Fase 2.pptx content for final-version audit."""
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import re
import json

SRC = Path(r"exports/_pptx_fase2_audit/Fase2_requisitos_atual.pptx")
OUT = Path(r"exports/_pptx_fase2_audit/fase2_atual_extract.txt")
OUT_JSON = Path(r"exports/_pptx_fase2_audit/fase2_atual_checks.json")

prs = Presentation(str(SRC))
lines = []
slides_meta = []

def shape_text(shape):
    chunks = []
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            chunks.append(t)
    if shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            chunks.append("\n".join(rows))
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            chunks.extend(shape_text(s))
    return chunks

for i, slide in enumerate(prs.slides, 1):
    parts = []
    for shape in slide.shapes:
        parts.extend(shape_text(shape))
    body = "\n".join(parts)
    lines.append(f"\n===== SLIDE {i} =====\n{body}")
    # title = first short line
    title = ""
    for ln in body.splitlines():
        ln = ln.strip()
        if ln and len(ln) < 140:
            title = ln
            break
    slides_meta.append({"n": i, "title": title, "chars": len(body)})

OUT.write_text("\n".join(lines), encoding="utf-8")

full = "\n".join(lines)
full_l = full.lower()

checks = {
    "slide_count": len(prs.slides),
    "file_size": SRC.stat().st_size,
    "titles": slides_meta,
    "flags": {
        "modelo_em_tipo_cobranca": bool(re.search(r"Tipo de Cobrança[\s\S]{0,800}\bModelo\b", full, re.I)),
        "campo_nome_cobranca": bool(re.search(r"Tipo de Cobrança[\s\S]{0,1200}\bNome\b", full, re.I)),
        "parceira_typo": "parceira" in full_l and "catálogo da parceira" in full_l,
        "solucao_parceria_obrig_nao": bool(re.search(r"Solução[\s\S]{0,1500}Parceria\s*\|\s*[^\n]*\|\s*Não", full, re.I)),
        "grupo_obrigatorio_absoluto": bool(re.search(r"RF-GRP-02|Grupo.*obrigat|obrigatório.*Grupo", full, re.I)),
        "luis_citacao": bool(re.search(r"\bLuís\b|\bLuis\b", full)),
        "sob_demanda_como_cobranca": bool(re.search(r"Sob demanda.*(Mensal|Anual|cobrança)", full, re.I)),
        "fator_cadastro_oficializado": bool(re.search(r"fator.*(cadastro|automático).*6 casas|cálculo automático.*fator", full_l)),
        "duas_metricas_produto": "métrica comum" in full_l and "métrica de comercialização" in full_l,
        "portal_cliente_consumo": bool(re.search(r"Portal do Cliente[\s\S]{0,2000}(cotação|ordem de serviço|OS|consumo)", full, re.I)),
        "fase_3_marcado": "fase 3" in full_l or "proposto" in full_l,
        "dados_parceria_proposta": bool(re.search(r"Dados de Parceria[\s\S]{0,400}proposta", full, re.I)),
        "rf_cat_00": "RF-CAT-00" in full or "cardinalidade" in full_l,
        "capa": bool(re.search(r"capa|objetivo do documento|glossário|critérios de aceite", full_l)),
        "slide_csv_campos": bool(re.search(r"CSV.*(campo|formulário|template)|Importação CSV", full, re.I)),
        "analise_mti_campos": bool(re.search(r"Análise.*(campo|justificativa)|formulário de análise", full, re.I)),
        "moeda_1_ro": bool(re.search(r"moeda universal[\s\S]{0,200}\b1\b", full_l)),
        "pendencias_secao": "pendência" in full_l or "pendente" in full_l,
        "homologar_publicar_apostilar": "apostilar" in full_l and "publicar" in full_l,
    },
}

# more precise cobranca field name from slide 2 table
m = re.search(r"===== SLIDE 2 =====([\s\S]*?)===== SLIDE 3 =====", full)
if m:
    s2 = m.group(1)
    checks["slide2_has_modelo_row"] = bool(re.search(r"(?m)^Modelo\s*\|", s2))
    checks["slide2_has_nome_row"] = bool(re.search(r"(?m)^Nome\s*\|", s2))

m5 = re.search(r"===== SLIDE 5 =====([\s\S]*?)===== SLIDE 6 =====", full)
if m5:
    s5 = m5.group(1)
    checks["slide5_parceria_row"] = re.findall(r"(?m)^Parceria\s*\|[^\n]+", s5)

m6 = re.search(r"===== SLIDE 6 =====([\s\S]*?)===== SLIDE 7 =====", full)
if m6:
    s6 = m6.group(1)
    checks["slide6_snippets"] = [ln for ln in s6.splitlines() if re.search(r"obrig|RF-GRP|Nome", ln, re.I)][:15]

m20 = re.search(r"===== SLIDE 20 =====([\s\S]*?)===== SLIDE 21 =====", full)
if m20:
    checks["slide20_preview"] = m20.group(1)[:1200]

OUT_JSON.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
print("slides", checks["slide_count"])
print(json.dumps(checks["flags"], ensure_ascii=False, indent=2))
print("slide2 modelo", checks.get("slide2_has_modelo_row"), "nome", checks.get("slide2_has_nome_row"))
print("slide5", checks.get("slide5_parceria_row"))
