# -*- coding: utf-8 -*-
"""Gera PPTX — Portal do Parceiro (Checkpoint Atlas 08/07/2026)."""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "docs/entregaveis/pptx/Atlas requisitos.pptx")
OUT_PATH = os.path.join(ROOT, "exports/Documentacao_Requisitos_F1_Portal_Parceiro.pptx")

BLUE = RGBColor(0x18, 0x22, 0xDC)
GREY_HDR = RGBColor(0xD9, 0xD9, 0xD9)
DARK = RGBColor(0x27, 0x27, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xBF, 0xBF)

F_TITLE = "PP Telegraf"
F_HEAD = "PP Telegraf"
F_LABEL = "PP Telegraf SemiBold"
F_BODY = "Arial"

COL_W_4 = [Inches(2.5), Inches(2.0), Inches(1.1), Inches(4.8)]
COL_W_3 = [Inches(3.0), Inches(2.5), Inches(5.0)]
HEADERS_4 = ["Campo", "Tipo", "Regras de Campo", "Observação"]
HEADERS_3 = ["Campo", "Tipo", "Regras e validação"]
ROWS_PER_SLIDE = 9


def set_fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def no_fill(cell):
    cell.fill.background()


def bottom_border(cell, rgb=LINE, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB",):
        existing = tcPr.find(qn(tag))
        if existing is not None:
            tcPr.remove(existing)
        ln = tcPr.makeelement(qn(tag), {"w": str(w), "cap": "flat"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])})
        fill.append(clr)
        ln.append(fill)
        tcPr.append(ln)


def style_cell(cell, text, *, bold=False, size=10, color=DARK, font=F_BODY,
               align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text or "").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        self.layout = next(
            (l for l in self.prs.slide_layouts if l.name == "Todo branco"),
            self.prs.slide_layouts[1],
        )
        xml_slides = self.prs.slides._sldIdLst
        for sld in list(xml_slides):
            rId = sld.get(qn("r:id"))
            try:
                self.prs.part.drop_rel(rId)
            except Exception:
                pass
            xml_slides.remove(sld)

    def add_slide(self):
        return self.prs.slides.add_slide(self.layout)

    def section_divider(self, title, subtitle):
        slide = self.add_slide()
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
        for y, txt, sz, bold in (
            (2.5, title, 34, True),
            (3.55, subtitle, 14, False),
        ):
            box = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(11.5), Inches(0.9))
            p = box.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.name = F_TITLE if bold else F_BODY
            r.font.color.rgb = WHITE

    def add_title(self, slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.0), Inches(0.45))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = DARK

    def intro_slide(self, title, historia, contexto, requisitos, regras):
        slide = self.add_slide()
        self.add_title(slide, title)
        tbl = slide.shapes.add_table(3, 2, Inches(0.45), Inches(0.72), Inches(12.3), Inches(5.95)).table
        tbl.columns[0].width = Inches(2.1)
        tbl.columns[1].width = Inches(10.2)
        for i, lab in enumerate(["História", "Contexto da História", "Requisitos e Regras"]):
            set_fill(tbl.cell(i, 0), BLUE)
            style_cell(tbl.cell(i, 0), lab, bold=True, size=11, color=WHITE, font=F_LABEL, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(0, 1), historia, size=10, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(1, 1), contexto, size=10, anchor=MSO_ANCHOR.TOP)
        c2 = tbl.cell(2, 1)
        c2.vertical_anchor = MSO_ANCHOR.TOP
        tf = c2.text_frame
        tf.word_wrap = True

        def par(text, *, bold=False, bullet=False, first=False, color=DARK):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            r = p.add_run()
            r.text = ("•  " if bullet else "") + text
            r.font.size = Pt(10)
            r.font.bold = bold
            r.font.name = F_BODY
            r.font.color.rgb = color

        par("Requisitos Funcionais", bold=True, color=BLUE, first=True)
        for req in requisitos:
            par(req, bullet=True)
        par("Regras de Negócio", bold=True, color=BLUE)
        for rg in regras:
            par(rg, bullet=True)
        for i in range(3):
            bottom_border(tbl.cell(i, 0))
            bottom_border(tbl.cell(i, 1))

    def action_intro(self, action_title, historia, contexto, requisitos, regras):
        self.intro_slide(action_title, historia, contexto, requisitos, regras)

    def proto_placeholder(self, label):
        slide = self.add_slide()
        self.add_title(slide, f"Protótipo — {label}")
        box = slide.shapes.add_textbox(Inches(0.6), Inches(2.8), Inches(11.5), Inches(1.0))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = "Captura de tela do protótipo — inserir na entrega."
        r.font.size = Pt(14)
        r.font.name = F_BODY
        r.font.color.rgb = DARK

    def _draw_table(self, slide, page, *, cols, col_w, top=Inches(1.05)):
        rows = len(page)
        height = min(Inches(5.85), Inches(0.36 * rows + 0.15))
        tbl_shape = slide.shapes.add_table(rows, cols, Inches(0.45), top, sum(col_w, Emu(0)), height)
        table = tbl_shape.table
        for ci, w in enumerate(col_w):
            table.columns[ci].width = w
        for ri, (kind, payload) in enumerate(page):
            if kind in ("bar", "subbar"):
                c0 = table.cell(ri, 0)
                c_last = table.cell(ri, cols - 1)
                c0.merge(c_last)
                set_fill(c0, BLUE if kind == "bar" else GREY_HDR)
                style_cell(
                    c0,
                    payload,
                    bold=True,
                    size=11 if kind == "bar" else 10,
                    color=WHITE if kind == "bar" else DARK,
                    font=F_HEAD,
                )
            elif kind == "colhdr":
                for ci, h in enumerate(payload):
                    set_fill(table.cell(ri, ci), GREY_HDR)
                    style_cell(table.cell(ri, ci), h, bold=True, size=10, color=DARK, font=F_HEAD)
            else:
                for ci, val in enumerate(payload):
                    no_fill(table.cell(ri, ci))
                    style_cell(
                        table.cell(ri, ci),
                        val,
                        bold=(ci == 0),
                        size=9,
                        font=F_HEAD if ci == 0 else F_BODY,
                    )
                    bottom_border(table.cell(ri, ci))

    def field_slides(self, slide_title, bar_label, rows, *, three_cols=False):
        headers = HEADERS_3 if three_cols else HEADERS_4
        cols = 3 if three_cols else 4
        col_w = COL_W_3 if three_cols else COL_W_4
        i = 0
        while i < len(rows):
            slide = self.add_slide()
            self.add_title(slide, slide_title)
            page = [("bar", bar_label), ("colhdr", headers)]
            count = 0
            while i < len(rows) and count < ROWS_PER_SLIDE:
                page.append(("fields", rows[i]))
                i += 1
                count += 1
            self._draw_table(slide, page, cols=cols, col_w=col_w)


# ---------- conteúdo (Checkpoint Atlas 08/07/2026) ----------

FONTE = "Fonte: Checkpoint Atlas — 08/07/2026 (Meeting Recording)."

VISAO = {
    "historia": (
        "Como parceiro da MTI, quero acessar um portal para completar o cadastro da "
        "organização, anexar documentos, gerenciar pessoas e manter os dados da parceria "
        "atualizados, para que a MTI analise e aprove sem precisar preencher tudo no back-office."
    ),
    "contexto": (
        f"{FONTE} Na Fase 1 o portal cobre documentação, dados da parceria e pessoas. "
        "A MTI inicia o cadastro no back-office (nome, CNPJ, sigla, documentos exigidos, "
        "responsável) e executa Disponibilizar para parceiro. O responsável acessa o portal "
        "(match por CPF), completa os dados e envia para análise. Colaboradores pré-cadastrados "
        "veem apenas Meus dados. Alterações pós-aprovação geram solicitação à MTI sem bloquear "
        "o acesso. Layout: aplicação UI Sydle (modelo Service Desk / Clube do Servidor); "
        "prioridade em campos e métodos."
    ),
    "requisitos": [
        "Disponibilizar portal exclusivo para parceiros pré-cadastrados pela MTI.",
        "Permitir ao responsável completar dados da organização e documentação.",
        "Exibir progresso de documentos por grupo (ex.: 3/10 anexados e validados).",
        "Permitir pré-cadastro de pessoas (CPF + e-mail) pelo responsável.",
        "Exibir Meus dados ao colaborador e Pessoas ao responsável.",
        "Encaminhar cadastros e alterações à MTI (aprovar / solicitar ajuste / reprovar).",
        "Gerar solicitação de alteração pós-aprovação sem bloquear o acesso ao portal.",
        "Priorizar campos e métodos; layout bonito e usável dentro do padrão Sydle.",
    ],
    "regras": [
        "RN-PORT-01: Acesso somente com CPF pré-cadastrado na classe Pessoa (não é cadastro aberto).",
        "RN-PORT-02: Quem envia o cadastro organizacional inicial é o Responsável da organização.",
        "RN-PORT-03: Parceiro não cria grupo/tipo de documento — apenas faz upload do configurado.",
        "RN-PORT-04: MTI aprova, solicita ajuste ou reprova; sempre com notificação (e-mail opcional).",
        "RN-PORT-05: Alteração pós-aprovação fica Pendente até a MTI validar; não bloqueia o portal.",
        "RN-PORT-06: Fase 1 = documentação + parceria + pessoas; Fase 2 = catálogo comercial.",
        "RN-PORT-07: Formulários do portal reutilizam regras do back-office; campos ocultos quando padrão.",
    ],
}

DISPONIBILIZAR = {
    "title": "Back-office — Ação: Disponibilizar para parceiro",
    "historia": (
        "Como analista MTI, quero cadastrar a organização parceira com dados mínimos e "
        "disponibilizá-la no portal, para que o responsável complete o cadastro e a documentação."
    ),
    "contexto": (
        "A MTI informa nome, sigla, CNPJ, se é empresa (organização raiz), tipo de cadastro = "
        "Parceiro do portal, documentos obrigatórios e o Responsável (já na classe Pessoa com CPF). "
        "Parceiro não tem unidade subordinada nesta fase. Produto/parceria comercial ainda não entra. "
        "O método Disponibilizar para parceiro notifica o responsável para acessar o portal."
    ),
    "requisitos": [
        "Cadastrar organização com dados mínimos no back-office.",
        "Selecionar grupos e tipos de documento que a organização deve enviar.",
        "Vincular Responsável já cadastrado (CPF informado).",
        "Executar Disponibilizar para parceiro e notificar o responsável.",
        "Garantir que o login do responsável bata com o CPF pré-cadastrado.",
    ],
    "regras": [
        "Não é qualquer pessoa que acessa o portal — somente quem a MTI pré-cadastrou.",
        "A MTI não preenche todos os dados da organização; o parceiro completa no portal.",
        "Organização parceira: sem unidade subordinada nesta fase.",
    ],
    "fields": [
        ["Nome da organização", "Texto", "Obrigatório", "Informado pela MTI no back-office."],
        ["Sigla", "Texto", "Obrigatório", ""],
        ["CNPJ", "Texto", "Obrigatório", "Identificador da empresa."],
        ["É empresa / org. raiz", "Sim/Não", "Obrigatório", "Parceiro = organização raiz."],
        ["Tipo de cadastro", "Opções", "Obrigatório", "Parceiro do portal."],
        ["Documentos exigidos", "Seleção múltipla", "Obrigatório", "Grupos e tipos MIPP configurados."],
        ["Responsável", "Referência Pessoa", "Obrigatório", "CPF já informado; recebe notificação."],
        ["Unidade subordinada", "—", "Não se aplica", "Parceiro não possui nesta fase."],
        ["Produto / parceria", "—", "Fora do escopo F1", "Entra em fases seguintes."],
    ],
}

LOGIN = {
    "title": "Portal — Login e acesso",
    "historia": (
        "Como responsável ou colaborador pré-cadastrado, quero autenticar-me no portal "
        "para visualizar apenas as abas e dados permitidos ao meu perfil."
    ),
    "contexto": (
        "O acesso não é aberto. A MTI já informou o CPF na classe Pessoa. No login, o sistema "
        "valida o CPF e libera a visualização. Responsável vê dados da parceria, documentos, "
        "pessoas e seus dados. Colaborador vê apenas Meus dados até concluir e ser aprovado."
    ),
    "requisitos": [
        "Autenticar usuário e validar CPF contra pré-cadastro.",
        "Liberar abas conforme perfil (Responsável × Colaborador).",
        "Bloquear acesso se CPF não estiver pré-cadastrado.",
    ],
    "regras": [
        "RN-PORT-LOGIN-01: Sem CPF pré-cadastrado não há acesso ao portal.",
        "RN-PORT-LOGIN-02: Colaborador sem envio para análise: acesso Pendente; só Meus dados.",
        "RN-PORT-LOGIN-03: Status do vínculo funcional permanece Pendente até o envio para análise.",
    ],
}

ORG_VISAO = {
    "title": "Portal — Dados da parceria (Organização) — Visão geral",
    "historia": (
        "Como responsável da organização, quero visualizar e completar os dados da parceria "
        "no portal, para enviar o cadastro completo à MTI sem que a MTI preencha tudo no back-office."
    ),
    "contexto": (
        "No portal são exibidos: dados de cadastro, documentos, tributos e encargos, contato "
        "e usuários. Não exibir: AZUS (unidade) e permissões. Integração por CNPJ (Receita/KPI) "
        "para preencher dados cadastrais. Conta bancária: reaproveitar opções já usadas nas "
        "demais classes (conta corrente, poupança, salário, pagamento)."
    ),
    "requisitos": [
        "Exibir seções permitidas da organização no portal.",
        "Ocultar AZUS e permissões.",
        "Integrar consulta por CNPJ para preencher dados de cadastro.",
        "Permitir Salvar (rascunho) e Enviar para análise.",
        "Permitir Ajustar quando a MTI solicitar correção.",
    ],
    "regras": [
        "RN-PORT-ORG-01: Somente o Responsável envia o cadastro organizacional inicial.",
        "RN-PORT-ORG-02: Campos de origem (perfil, org, unidade) = somente leitura quando já definidos.",
        "RN-PORT-ORG-03: Após envio, MTI aprova, solicita ajuste ou reprova com notificação.",
    ],
}

ORG_CAMPOS = [
    ["Dados de unidade", "Seção", "Parcial", "AZUS não é exibido no portal."],
    ["Permissões", "Seção", "Oculta", "Não exibida ao parceiro."],
    ["Dados de cadastro", "Seção", "Editável", "Razão social, nome, natureza etc.; CNPJ com integração."],
    ["CNPJ", "Texto", "Obrigatório", "Dispara integração para preencher dados cadastrais."],
    ["Documentos", "Seção", "Upload", "Grupos/tipos definidos pela MTI; só anexar arquivos."],
    ["Tributos e encargos", "Seção", "Editável", "Parceiro preenche no portal."],
    ["Dados de contato", "Seção", "Editável", ""],
    ["Usuários / Pessoas", "Seção", "Conforme perfil", "Responsável gerencia; ver aba Pessoas."],
    ["Tipo de conta bancária", "Opções", "Conforme classe", "Corrente, poupança, salário, pagamento."],
    ["Progresso documental", "Indicador", "Somente leitura", "Ex.: 3/10 documentos anexados e validados por grupo."],
]

DOCS_VISAO = {
    "title": "Portal — Documentos da organização",
    "historia": (
        "Como responsável, quero anexar os documentos exigidos pela MTI e acompanhar o "
        "progresso por grupo, para finalizar a habilitação da parceria."
    ),
    "contexto": (
        "A MTI configura quais grupos e tipos a organização deve enviar. No portal o parceiro "
        "não cria grupo nem tipo — apenas faz upload. Exemplos de documentos da apresentação "
        "da parceira: papel timbrado, contrato, planos de negócio, planos de operação, logomarca. "
        "Validação por OCR/IA (vencimento, assinatura, CNPJ) fica para alinhamento posterior "
        "(portfólio Simplifica); status previsto: Aprovado, Pendente ou Recusado com motivo."
    ),
    "requisitos": [
        "Listar grupos e tipos configurados pela MTI.",
        "Permitir upload por tipo de documento.",
        "Exibir progresso por grupo (anexados e validados / total).",
        "Bloquear criação de novo grupo ou tipo pelo parceiro.",
        "Preparar status de validação (Aprovado / Pendente / Recusado) para OCR/IA.",
    ],
    "regras": [
        "RN-PORT-DOC-01: Upload somente dos documentos configurados no back-office.",
        "RN-PORT-DOC-02: Progresso visual obrigatório por grupo documental.",
        "RN-PORT-DOC-03: OCR/IA não bloqueia o início da Fase 1; alinhar depois com Simplifica.",
    ],
}

DOCS_CAMPOS = [
    ["Grupo de documento", "Texto", "Somente leitura", "Ex.: Habilitação jurídica; definido pela MTI."],
    ["Tipo de documento", "Texto", "Somente leitura", "Ex.: papel timbrado, contrato, logomarca."],
    ["Arquivo", "Upload", "Obrigatório por tipo", "Parceiro anexa o arquivo."],
    ["Progresso do grupo", "Indicador", "Somente leitura", "Ex.: 1/10 anexados e validados."],
    ["Status do documento", "Opções", "Somente leitura", "Aprovado · Pendente · Recusado (OCR/IA — fase posterior)."],
    ["Motivo da recusa", "Texto", "Somente leitura", "Ex.: vencido, sem assinatura, CNPJ divergente."],
]

ENVIO_ORG = {
    "title": "Portal — Ação: Enviar cadastro da organização para análise",
    "historia": (
        "Como responsável, quero enviar o cadastro completo da organização para a MTI "
        "analisar, para obter aprovação, solicitação de ajuste ou reprovação."
    ),
    "contexto": (
        "Após preencher dados e anexar documentos, o responsável usa Enviar para análise. "
        "Métodos no back-office: Aprovar, Solicitar ajuste, Reprovar — com notificação "
        "(e-mail opcional). Salvar mantém rascunho. Ajustar é usado quando a MTI devolve "
        "com correção pendente."
    ),
    "requisitos": [
        "Disponibilizar Salvar (rascunho).",
        "Disponibilizar Enviar para análise.",
        "Disponibilizar Ajustar após solicitação da MTI.",
        "Notificar responsável e MTI a cada decisão.",
    ],
    "regras": [
        "Somente o Responsável envia o cadastro organizacional.",
        "MTI sempre notifica a decisão (e-mail opcional).",
    ],
    "fields": [
        ["Salvar", "Botão", "—", "Mantém rascunho sem enviar à MTI."],
        ["Enviar para análise", "Botão", "—", "Gera pendência no back-office."],
        ["Ajustar", "Botão", "Condicional", "Após solicitação de ajuste da MTI."],
        ["Status", "Opções", "Somente leitura", "Rascunho · Aguardando análise · Aprovado · Ajuste · Reprovado."],
    ],
}

ALTERACAO = {
    "title": "Portal — Ação: Solicitar alteração pós-aprovação",
    "historia": (
        "Como responsável, quero alterar dados da organização já aprovada (ex.: endereço) "
        "e enviar à MTI, para que a alteração só valha após parecer, sem perder o acesso ao portal."
    ),
    "contexto": (
        "Após aprovação, qualquer alteração gera solicitação de alteração de dados de "
        "organização (tipo processo/ticket). Enquanto a MTI não aprova, a alteração fica "
        "Pendente. O acesso ao portal NÃO é bloqueado — apenas a alteração aguarda validação. "
        "No back-office o analista vê o que mudou e aprova ou reprova (parecer). Tela do "
        "portal: formulário Sydle embedado, pré-preenchido, em tela cheia (não precisa modal)."
    ),
    "requisitos": [
        "Abrir formulário pré-preenchido com dados vigentes.",
        "Ao confirmar, gerar solicitação/processo no back-office.",
        "Manter status Pendente na alteração até parecer da MTI.",
        "Não bloquear o acesso do parceiro ao portal.",
        "Permitir aprovar ou reprovar no back-office com parecer.",
    ],
    "regras": [
        "RN-PORT-ALT-01: Alteração pós-aprovação ≠ update direto na classe — gera processo.",
        "RN-PORT-ALT-02: Acesso ao portal permanece liberado durante a pendência.",
        "RN-PORT-ALT-03: Só a alteração específica fica inválida até aprovação.",
    ],
    "fields": [
        ["Dados alterados", "Formulário", "Editável", "Pré-preenchido com dados vigentes."],
        ["Enviar / Confirmar", "Botão", "—", "Abre solicitação de alteração."],
        ["Status da alteração", "Opções", "Somente leitura", "Pendente · Aprovada · Reprovada."],
        ["Parecer MTI", "Texto", "Back-office", "Aprovar ou reprovar com observação."],
    ],
}

MEUS_DADOS = {
    "title": "Portal — Aba: Meus dados (pessoa)",
    "historia": (
        "Como responsável ou colaborador, quero ver e completar meus dados pessoais e "
        "documentos no portal, para finalizar meu vínculo com a organização parceira."
    ),
    "contexto": (
        "Exibe nome, foto, dados pessoais, contato e documentos (grupo/tipo + upload). "
        "Perfil padrão = Parceiro (não editável). Dados complementares: sem acesso/modificação "
        "pelo parceiro. Cargo e vínculo org/unidade definidos pela MTI. Dados do MT Login/Gov.br "
        "preenchidos automaticamente (a detalhar). Colaborador: primeiro login após pré-cadastro "
        "abre Meus dados; Finalizar = Enviar para análise. Salvar permite continuar depois "
        "(ex.: retorno com solicitação de ajuste)."
    ),
    "requisitos": [
        "Exibir dados pessoais, contato e documentos da pessoa.",
        "Permitir upload de foto e documentos exigidos.",
        "Bloquear edição de Perfil e dados complementares.",
        "Disponibilizar Salvar e Enviar para análise / Finalizar cadastro.",
    ],
    "regras": [
        "RN-PORT-MD-01: Perfil = Parceiro; somente leitura.",
        "RN-PORT-MD-02: Aba complementares inacessível ao parceiro.",
        "RN-PORT-MD-03: Sem envio para análise, vínculo funcional = Pendente; só Meus dados.",
    ],
}

MEUS_DADOS_CAMPOS = [
    ["Nome", "Texto", "Editável / provedor", "Pode subir foto e completar dados."],
    ["Foto", "Arquivo", "Editável", "Upload pelo usuário."],
    ["Dados MT Login / Gov.br", "Vários", "Somente leitura", "Puxados automaticamente (a detalhar)."],
    ["Perfil", "Texto", "Somente leitura", "Parceiro — não alterável."],
    ["Dados complementares", "Seção", "Sem acesso", "Parceiro não modifica."],
    ["Cargo", "Texto", "Somente leitura", "Cadastrado pela MTI (ex.: Gerente de Projetos)."],
    ["Organização / Unidade", "Texto", "Somente leitura", "Definidos no pré-cadastro / MTI."],
    ["Documentos pessoais", "Upload", "Conforme config.", "Mesmo padrão grupo + tipo da organização."],
    ["Contato", "Seção", "Editável", "Informações padrão da classe Pessoa."],
    ["Solicitar ajuste", "Texto longo", "Editável", "Quando aplicável; notifica responsável e MTI."],
    ["Salvar", "Botão", "—", "Rascunho; continua posteriormente."],
    ["Enviar para análise", "Botão", "—", "Finaliza cadastro e encaminha à MTI."],
]

PESSOAS = {
    "title": "Portal — Aba: Pessoas (somente Responsável)",
    "historia": (
        "Como responsável da organização, quero pré-cadastrar colaboradores e gerenciar "
        "cargo e situação ativa, para que cada pessoa conclua o cadastro e a MTI aprove o vínculo."
    ),
    "contexto": (
        "Aba Pessoas aparece só para o Responsável. Colaborador vê apenas Meus dados "
        "(ex.: Fernanda = responsável → Pessoas; analista = colaborador → Meus dados). "
        "Pré-cadastro: CPF + e-mail; organização/unidade pré-preenchidas; cargo automático "
        "se houver só um habilitado, senão o responsável seleciona. Notificação por e-mail "
        "com link. Responsável edita só desativar e cargo — não altera dados pessoais. "
        "Responsável não visualiza a fila de aprovação da MTI."
    ),
    "requisitos": [
        "Listar pessoas da organização.",
        "Pré-cadastrar com CPF e e-mail e notificar.",
        "Pré-preencher organização e unidade.",
        "Aplicar regra de cargo (1 cargo = default; N cargos = seleção).",
        "Permitir editar apenas Ativo (desativar) e Cargo.",
        "Encaminhar cadastro finalizado pelo colaborador à MTI.",
    ],
    "regras": [
        "RN-PORT-PES-01: Aba Pessoas só para o Responsável da organização.",
        "RN-PORT-PES-02: Cargo não pode ficar em branco no pré-cadastro.",
        "RN-PORT-PES-03: Responsável não altera informações pessoais nem documentos de terceiros.",
        "RN-PORT-PES-04: Aprovação/reprovação de pessoa é exclusiva da MTI.",
        "RN-PORT-PES-05: Colaborador sem envio: acesso Pendente; só Meus dados.",
    ],
}

PESSOAS_LISTA = [
    ["Nome", "Texto", "Somente leitura", "Lista de colaboradores da organização."],
    ["CPF", "Texto", "Somente leitura", ""],
    ["E-mail", "Texto", "Somente leitura", "Informado no pré-cadastro."],
    ["Cargo", "Seleção", "Editável (responsável)", "Alterável na ação Editar."],
    ["Ativo", "Sim/Não", "Editável (responsável)", "Desativar vínculo."],
    ["Status do cadastro", "Opções", "Somente leitura", "Pré-cadastrado · Em preenchimento · Aguardando MTI · Ativo · Ajuste."],
    ["Vínculo funcional", "Opções", "Somente leitura", "Pendente até envio para análise."],
]

PRECADASTRO = {
    "title": "Portal — Ação: Pré-cadastrar pessoa",
    "historia": (
        "Como responsável, quero pré-cadastrar um colaborador informando CPF e e-mail, "
        "para que ele receba o link e complete Meus dados no portal."
    ),
    "contexto": (
        "Formulário resumido: CPF, e-mail, organização e unidade pré-preenchidas, cargo "
        "condicional. Ao confirmar, notifica a pessoa. No primeiro acesso ela completa "
        "dados (MT Login/Gov.br) e documentos e envia para análise da MTI."
    ),
    "requisitos": [
        "Exigir CPF e e-mail.",
        "Pré-preencher organização e unidade.",
        "Aplicar regra de cargo (único vs múltiplo).",
        "Enviar notificação com link do portal.",
    ],
    "regras": [
        "Executável somente pelo Responsável.",
        "Não cria acesso pleno até Finalizar + aprovação MTI.",
    ],
    "fields": [
        ["CPF", "Texto", "Obrigatório", "Documento principal do colaborador."],
        ["E-mail", "Texto", "Obrigatório", "Destino da notificação com link."],
        ["Organização", "Texto", "Somente leitura", "Da sessão do responsável."],
        ["Unidade organizacional", "Seleção", "Pré-preenchida", "Filtrada pela organização."],
        ["Cargo", "Seleção", "Obrigatório", "1 cargo = default; N cargos = seleção do responsável."],
    ],
}

FINALIZAR = {
    "title": "Portal — Ação: Finalizar cadastro (colaborador)",
    "historia": (
        "Como colaborador pré-cadastrado, quero completar Meus dados e documentos e "
        "enviar para análise, para que a MTI aprove meu vínculo."
    ),
    "contexto": (
        "Após o link do e-mail, o colaborador autentica-se, completa dados e documentos "
        "e usa Enviar para análise (Finalizar). Salvar mantém rascunho (útil após solicitação "
        "de ajuste). Enquanto não envia, acesso permanece Pendente e restrito a Meus dados."
    ),
    "requisitos": [
        "Preencher/confirmar dados pessoais e documentos.",
        "Salvar rascunho.",
        "Enviar para análise da MTI.",
        "Atualizar status do vínculo funcional após o envio.",
    ],
    "regras": [
        "Sem envio: vínculo = Pendente; só Meus dados.",
        "MTI: aprovar, solicitar ajuste ou reprovar.",
        "Responsável não visualiza a fila de aprovação.",
    ],
    "fields": [
        ["Dados pessoais", "Vários", "Conforme regras", "Provedor + complementos permitidos."],
        ["Documentos", "Upload", "Conforme config.", "Grupo e tipo definidos."],
        ["Salvar", "Botão", "—", "Continuar posteriormente."],
        ["Enviar para análise", "Botão", "—", "Finaliza e encaminha à MTI."],
    ],
}

EDITAR_PESSOA = {
    "title": "Portal — Ação: Editar pessoa (responsável)",
    "historia": (
        "Como responsável, quero desativar um colaborador ou alterar o cargo, "
        "sem modificar dados pessoais ou documentos."
    ),
    "contexto": (
        "Na aba Pessoas, a edição pelo responsável limita-se a Ativo e Cargo. "
        "Informações pessoais e documentação permanecem no fluxo do colaborador / MTI."
    ),
    "requisitos": [
        "Permitir alterar Ativo e Cargo.",
        "Bloquear edição de dados pessoais e documentos.",
    ],
    "regras": [
        "Somente o Responsável executa no portal.",
        "Inativar impede acesso sem excluir histórico.",
    ],
    "fields": [
        ["Ativo", "Sim/Não", "Obrigatório", "Desativar vínculo."],
        ["Cargo", "Seleção", "Obrigatório", "Lista habilitada para a organização."],
    ],
}

ANALISE_MTI = {
    "title": "Back-office — Ação: Analisar cadastro (MTI)",
    "historia": (
        "Como analista MTI, quero aprovar, solicitar ajuste ou reprovar cadastros "
        "enviados pelo portal (organização e pessoas), para manter o controle de qualidade."
    ),
    "contexto": (
        "Aplica-se ao cadastro inicial da organização, ao cadastro de pessoas e às "
        "solicitações de alteração pós-aprovação. Sempre com notificação. O parecer "
        "de alteração de organização já existe no back-office, mas depende do portal "
        "para ficar funcional."
    ),
    "requisitos": [
        "Exibir dados enviados pelo portal.",
        "Aprovar, solicitar ajuste ou reprovar.",
        "Notificar o parceiro (e-mail opcional).",
        "Registrar parecer e operador.",
    ],
    "regras": [
        "Somente perfil MTI executa.",
        "Responsável do portal não visualiza a fila de aprovação.",
        "Alteração pós-aprovação: aprovar/reprovar como solicitação/processo.",
    ],
    "fields": [
        ["Resumo do envio", "Texto", "Somente leitura", "Org ou pessoa; dados e documentos."],
        ["Decisão", "Opções", "Obrigatório", "Aprovar · Solicitar ajuste · Reprovar."],
        ["Observação / motivo", "Texto longo", "Obrigatório se ajuste/reprovação", "Enviado ao parceiro."],
        ["Notificação", "Sistema", "Automática", "E-mail opcional."],
    ],
}

TECNICO = {
    "title": "Portal — Decisões técnicas e UX",
    "historia": (
        "Como time de implementação, quero construir o portal no Sydle com o máximo "
        "de usabilidade possível, priorizando campos e métodos da Fase 1."
    ),
    "contexto": (
        "Referência visual: Clube do Servidor (aplicação UI no Sydle). Base: modelo "
        "Service Desk. Personalização limitada no plano padrão (componentes, carrossel, "
        "seções, header). Formulários podem ser embedados no portal com cara de Sydle, "
        "em tela cheia. UI custom (Stencil/HTML/CSS) só se a OS cobrir; por ora, padrão "
        "Service Desk. Luís exige portal bonito e acessível. Protótipo Cursor/Especificador "
        "Sydle como referência de campos."
    ),
    "requisitos": [
        "Criar aplicação UI no Sydle baseada em Service Desk / Clube do Servidor.",
        "Embedar formulários Sydle no portal (tela cheia).",
        "Priorizar campos e métodos sobre customização visual avançada.",
        "Usar protótipo do Especificador como referência de campos e regras.",
    ],
    "regras": [
        "RN-PORT-UX-01: Prioridade = campos e métodos corretos.",
        "RN-PORT-UX-02: Layout o mais bonito/usável possível dentro do padrão contratado.",
        "RN-PORT-UX-03: Regras de campo = mesmas do back-office; portal só oculta o que for padrão.",
    ],
}

ROADMAP = {
    "title": "Portal — Roadmap de fases",
    "historia": (
        "Como stakeholder, quero entender o que entra na Fase 1 e o que fica para depois, "
        "para priorizar documentação, parceria e pessoas antes do catálogo comercial."
    ),
    "contexto": (
        "Fase 1: documentação do cadastro e dados da parceria + pessoas. Fase 2: parceiro "
        "alimenta catálogo (ex.: Simplifica), envia à MTI, aprovação e cotação/OS. Futuro: "
        "processo comercial completo. Validação OCR/IA de documentos: alinhar depois com "
        "portfólio Simplifica — não bloqueia o início."
    ),
    "requisitos": [
        "Entregar Fase 1: login, dados da parceria, documentos, Meus dados, Pessoas.",
        "Planejar Fase 2: catálogo no portal com aprovação MTI.",
        "Alinhar OCR/IA de documentos em momento posterior.",
    ],
    "regras": [
        "RN-PORT-FASE-01: Fase 1 não inclui alimentação de catálogo comercial.",
        "RN-PORT-FASE-02: OCR/IA é desejável, mas não bloqueante para iniciar F1.",
    ],
}

DEMANDAS = [
    ["Login", "Autenticação + match CPF pré-cadastrado", "Fase 1", "Demanda listada na reunião."],
    ["Dados da parceria", "Editar, Salvar, Enviar para análise, Ajustar", "Fase 1", "Organização no portal."],
    ["Documentos", "Upload + progresso + (depois) OCR/IA", "Fase 1 / pós", "Parte mais complexa e primordial."],
    ["Meus dados", "Dados pessoais: Editar, Salvar, Enviar para análise", "Fase 1", "Responsável e colaborador."],
    ["Pessoas", "Listar, Pré-cadastrar (CPF+e-mail), Editar (Ativo+Cargo)", "Fase 1", "Somente Responsável."],
    ["Alteração pós-aprovação", "Solicitação/processo no back-office", "Fase 1", "Não bloqueia acesso."],
    ["Disponibilizar para parceiro", "Método back-office + notificação", "Fase 1", "Libera acesso do responsável."],
    ["Catálogo no portal", "Parceiro alimenta itens e envia à MTI", "Fase 2", "Ex.: Simplifica."],
]


def main():
    deck = Deck()

    # 1. Capa / visão
    deck.section_divider(
        "Portal do Parceiro",
        "Fase 1 — Checkpoint Atlas 08/07/2026 · Documentação, Parceria e Pessoas",
    )
    deck.intro_slide(
        "Portal do Parceiro — Visão geral",
        VISAO["historia"],
        VISAO["contexto"],
        VISAO["requisitos"],
        VISAO["regras"],
    )

    # 2. Disponibilizar (back-office)
    deck.section_divider("Back-office MTI", "Disponibilizar organização para o portal")
    deck.action_intro(
        DISPONIBILIZAR["title"],
        DISPONIBILIZAR["historia"],
        DISPONIBILIZAR["contexto"],
        DISPONIBILIZAR["requisitos"],
        DISPONIBILIZAR["regras"],
    )
    deck.field_slides("Telas e Campos", "Back-office › Disponibilizar para parceiro", DISPONIBILIZAR["fields"])
    deck.proto_placeholder("Disponibilizar para parceiro")

    # 3. Login
    deck.section_divider("Portal", "Login, acesso e abas por perfil")
    deck.intro_slide(
        LOGIN["title"],
        LOGIN["historia"],
        LOGIN["contexto"],
        LOGIN["requisitos"],
        LOGIN["regras"],
    )

    # 4. Organização / Parceria
    deck.section_divider("Portal › Dados da parceria", "Organização, documentos e envio")
    deck.intro_slide(
        ORG_VISAO["title"],
        ORG_VISAO["historia"],
        ORG_VISAO["contexto"],
        ORG_VISAO["requisitos"],
        ORG_VISAO["regras"],
    )
    deck.field_slides("Telas e Campos", "Portal › Dados da parceria — Seções e campos", ORG_CAMPOS)

    deck.intro_slide(
        DOCS_VISAO["title"],
        DOCS_VISAO["historia"],
        DOCS_VISAO["contexto"],
        DOCS_VISAO["requisitos"],
        DOCS_VISAO["regras"],
    )
    deck.field_slides("Telas e Campos", "Portal › Documentos da organização", DOCS_CAMPOS)
    deck.proto_placeholder("Documentos da organização")

    deck.action_intro(
        ENVIO_ORG["title"],
        ENVIO_ORG["historia"],
        ENVIO_ORG["contexto"],
        ENVIO_ORG["requisitos"],
        ENVIO_ORG["regras"],
    )
    deck.field_slides("Telas e Campos", ENVIO_ORG["title"].replace("Portal — ", ""), ENVIO_ORG["fields"])
    deck.proto_placeholder("Enviar organização para análise")

    # 5. Alteração pós-aprovação
    deck.action_intro(
        ALTERACAO["title"],
        ALTERACAO["historia"],
        ALTERACAO["contexto"],
        ALTERACAO["requisitos"],
        ALTERACAO["regras"],
    )
    deck.field_slides("Telas e Campos", "Portal › Alteração pós-aprovação", ALTERACAO["fields"])
    deck.proto_placeholder("Alteração pós-aprovação")

    # 6. Meus dados
    deck.section_divider("Portal › Meus dados", "Pessoa — responsável e colaborador")
    deck.intro_slide(
        MEUS_DADOS["title"],
        MEUS_DADOS["historia"],
        MEUS_DADOS["contexto"],
        MEUS_DADOS["requisitos"],
        MEUS_DADOS["regras"],
    )
    deck.field_slides("Telas e Campos", "Portal › Aba: Meus dados", MEUS_DADOS_CAMPOS)
    deck.proto_placeholder("Meus dados")

    # 7. Pessoas
    deck.section_divider("Portal › Pessoas", "Pré-cadastro e gestão pelo Responsável")
    deck.intro_slide(
        PESSOAS["title"],
        PESSOAS["historia"],
        PESSOAS["contexto"],
        PESSOAS["requisitos"],
        PESSOAS["regras"],
    )
    deck.field_slides("Telas e Campos", "Portal › Aba: Pessoas (lista)", PESSOAS_LISTA)

    for act in (PRECADASTRO, FINALIZAR, EDITAR_PESSOA):
        deck.action_intro(act["title"], act["historia"], act["contexto"], act["requisitos"], act["regras"])
        deck.field_slides("Telas e Campos", act["title"].replace("Portal — ", ""), act["fields"])
        deck.proto_placeholder(act["title"].split("—")[-1].strip())

    # 8. Análise MTI
    deck.section_divider("Back-office MTI", "Análise de cadastros e alterações")
    deck.action_intro(
        ANALISE_MTI["title"],
        ANALISE_MTI["historia"],
        ANALISE_MTI["contexto"],
        ANALISE_MTI["requisitos"],
        ANALISE_MTI["regras"],
    )
    deck.field_slides("Telas e Campos", "Back-office › Analisar cadastro (MTI)", ANALISE_MTI["fields"])
    deck.proto_placeholder("Analisar cadastro MTI")

    # 9. Técnico + roadmap + demandas
    deck.section_divider("Implementação", "UX Sydle, roadmap e demandas Fase 1")
    deck.intro_slide(
        TECNICO["title"],
        TECNICO["historia"],
        TECNICO["contexto"],
        TECNICO["requisitos"],
        TECNICO["regras"],
    )
    deck.intro_slide(
        ROADMAP["title"],
        ROADMAP["historia"],
        ROADMAP["contexto"],
        ROADMAP["requisitos"],
        ROADMAP["regras"],
    )
    deck.field_slides(
        "Telas e Campos",
        "Demandas Fase 1 (listadas na reunião) + foresight Fase 2",
        DEMANDAS,
    )

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    deck.prs.save(OUT_PATH)
    print("PPTX gerado:", OUT_PATH)
    print("Slides:", len(deck.prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
