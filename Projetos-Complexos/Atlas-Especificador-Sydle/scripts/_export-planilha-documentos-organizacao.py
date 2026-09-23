"""Gera a planilha canônica da aba Documentos da Organização."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "exports" / "Planilha_Canonica_Documentos_Organizacao.xlsx"

HEADERS = [
    "Grupo",
    "Campo",
    "Opções / Tipo",
    "Obrigatório",
    "Edição",
    "Visibilidade",
    "Regra / Validação",
    "ID técnico",
]

SHEETS = {
    "Organização - Documentos": [
        [
            "Documentos",
            "Etapas a incluir",
            "Multiseleção: Habilitação documental; Pré-rito parceria; Execução da parceria",
            "Não",
            "MTI",
            "Organização raiz = Sim",
            "Selecionar etapas ativas e clicar em Adicionar. Na Organização não se escolhe grupo; grupos e tipos vêm do catálogo da etapa.",
            "patlasv4proto-uo-etapas-selecao",
        ],
        [
            "Documentos",
            "Adicionar",
            "Ação",
            "—",
            "MTI",
            "Organização raiz = Sim",
            "Inclui as etapas selecionadas, com seus grupos e tipos. Não deve duplicar etapa já incluída.",
            "patlasv4proto-uo-etapas-adicionar",
        ],
        [
            "Documentos",
            "Etapas incluídas",
            "Acordeão / formulário embutido",
            "Condicional",
            "MTI / Parceiro",
            "Organização raiz = Sim",
            "Cada etapa exibe grupos e documentos copiados do catálogo. O parceiro anexa arquivos; não cria tipos.",
            "patlasv4proto-uo-etapas-documentacionais",
        ],
        [
            "Documentos",
            "Progresso geral / da etapa",
            "Texto calculado, ex.: 0/6",
            "—",
            "Somente leitura",
            "Quando houver etapas",
            "Calculado como documentos válidos sobre total exigido. Exibido no cabeçalho da etapa.",
            "patlasv4proto-uoetapa-progresso",
        ],
        [
            "Habilitação",
            "Status de habilitação documental",
            "Incompleta; Em apresentação; Completa – aguardando MTI; Aprovada pela MTI; Reprovada pela MTI",
            "Não",
            "Sistema / MTI",
            "Tipo de organização = Parceiro",
            "Incompleta → Disponibilizar → Em apresentação → Enviar → Completa – aguardando MTI → Aprovar/Reprovar. Solicitar ajuste retorna para Em apresentação.",
            "patlasv4proto-uo-habilitacao-status",
        ],
        [
            "Habilitação",
            "NDA assinado",
            "Sim; Não",
            "Não",
            "Conforme permissão",
            "Tipo de organização = Parceiro",
            "Controla a visualização, pelo parceiro, de documentos anexados pela MTI.",
            "patlasv4proto-uo-nda-assinado",
        ],
        [
            "Habilitação",
            "Acesso comercial",
            "Sim; Não",
            "Não",
            "Sistema / MTI",
            "Tipo de organização = Parceiro",
            "Sim significa acesso liberado. Em geral permanece Não até o cadastro ficar Aprovado pela MTI.",
            "patlasv4proto-uo-acesso-comercial",
        ],
        [
            "Habilitação",
            "Progresso — Habilitação jurídica",
            "Texto calculado, ex.: 12/19",
            "—",
            "Somente leitura",
            "Tipo = Parceiro",
            "Quantidade válida sobre total do grupo Habilitação Jurídica.",
            "patlasv4proto-uo-prog-juridica",
        ],
        [
            "Habilitação",
            "Progresso — Qualificação técnica",
            "Texto calculado",
            "—",
            "Somente leitura",
            "Tipo = Parceiro",
            "Quantidade válida sobre total do grupo Qualificação Técnica.",
            "patlasv4proto-uo-prog-tecnica",
        ],
        [
            "Habilitação",
            "Progresso — Qualificação financeira",
            "Texto calculado",
            "—",
            "Somente leitura",
            "Tipo = Parceiro",
            "Quantidade válida sobre total do grupo Qualificação Econômica e Financeira.",
            "patlasv4proto-uo-prog-financeira",
        ],
        [
            "Habilitação",
            "Progresso — Compliance",
            "Texto calculado",
            "—",
            "Somente leitura",
            "Tipo = Parceiro",
            "Quantidade válida sobre total do grupo Compliance.",
            "patlasv4proto-uo-prog-compliance",
        ],
        [
            "Execução (legado)",
            "Documentos de execução da parceria",
            "Tabela embutida antiga",
            "Não",
            "Não utilizar no modelo novo",
            "Oculto",
            "Campo legado. A modelagem correta usa a etapa Execução da parceria. Não duplicar documentos neste bloco.",
            "patlasv4proto-uo-docs-execucao",
        ],
    ],
    "Etapa incluída": [
        [
            "Etapa",
            "Etapa documentacional",
            "Habilitação documental; Pré-rito parceria; Execução da parceria",
            "Sim",
            "Somente leitura",
            "Título do acordeão",
            "Definida ao clicar em Adicionar; identifica a etapa incluída.",
            "patlasv4proto-uoetapa-nome",
        ],
        [
            "Etapa",
            "Ordem",
            "Número",
            "Sim",
            "Automática",
            "Oculto",
            "Ordem herdada do catálogo para exibição das etapas.",
            "patlasv4proto-uoetapa-ordem",
        ],
        [
            "Etapa",
            "Ativo",
            "Sim; Não",
            "Sim",
            "Somente leitura",
            "Oculto",
            "Definido na inclusão da etapa. Etapas inativas do catálogo não entram em novos cadastros.",
            "patlasv4proto-uoetapa-ativo",
        ],
        [
            "Etapa",
            "Progresso da etapa",
            "Texto calculado, ex.: 22/40",
            "—",
            "Somente leitura",
            "Cabeçalho do acordeão",
            "Agrega a contagem de todos os grupos da etapa.",
            "patlasv4proto-uoetapa-progresso",
        ],
        [
            "Etapa",
            "Grupos e documentos da etapa",
            "Formulários de Grupo na etapa",
            "Condicional",
            "Anexo pelo responsável / análise MTI",
            "Corpo do acordeão",
            "Carregado do catálogo. A Organização não seleciona grupos manualmente.",
            "patlasv4proto-uoetapa-grupos",
        ],
    ],
    "Grupo na etapa": [
        [
            "Grupo",
            "Grupo de documentos",
            "Habilitação Jurídica; Qualificação Técnica; Qualificação Econômica e Financeira; Compliance (adendo); Pré-rito parceria; Documentos de execução",
            "Sim",
            "Somente leitura",
            "Título do acordeão",
            "Vem da composição da Etapa documentacional. Não é selecionado na Organização.",
            "patlasv4proto-uogrpetapa-grupo",
        ],
        [
            "Grupo",
            "Progresso",
            "Texto calculado, ex.: 0/3",
            "—",
            "Somente leitura",
            "Cabeçalho do acordeão",
            "Quantidade de documentos válidos sobre o total do grupo.",
            "patlasv4proto-uogrpetapa-progresso",
        ],
        [
            "Grupo",
            "Documentos",
            "Tabela de Documento habilitação MIPP",
            "Condicional",
            "Parceiro / MTI conforme origem",
            "Corpo do acordeão",
            "Tipos vêm do catálogo. Permite upload e acompanhamento do status de validação.",
            "patlasv4proto-uogrpetapa-tipos",
        ],
    ],
    "Documento MIPP": [
        [
            "Documento",
            "Grupo do documento",
            "Referência ao Grupo de Documento",
            "Não",
            "Somente leitura",
            "Oculto",
            "Mantido apenas para rastreio; o grupo já é representado pelo acordeão.",
            "mqlh6yafch5jzb",
        ],
        [
            "Documento",
            "Tipo de documento",
            "Referência ao catálogo Tipo de Documento",
            "Sim",
            "Origem catálogo",
            "Grade",
            "Parceiro não cria tipo. O tipo já vem configurado no grupo da etapa.",
            "patlasv4proto-uodoc-tipo",
        ],
        [
            "Documento",
            "Número",
            "Texto",
            "Não",
            "Parceiro / MTI",
            "Grade",
            "Identificador do documento, como CPF, RG, CNPJ ou protocolo, quando aplicável.",
            "patlasv4proto-uodoc-numero",
        ],
        [
            "Documento",
            "Arquivo",
            "Upload de arquivo",
            "Sim para documento exigido",
            "Parceiro / MTI conforme tipo",
            "Grade",
            "Upload obrigatório para concluir o documento.",
            "patlasv4proto-uodoc-arquivo",
        ],
        [
            "Documento",
            "Data de anexo",
            "Data",
            "—",
            "Automática",
            "Grade",
            "Preenchida automaticamente no upload.",
            "patlasv4proto-uodoc-data-anexo",
        ],
        [
            "Documento",
            "Data de vencimento",
            "Data",
            "Condicional",
            "OCR / cálculo / manual",
            "Grade",
            "Automático: OCR. Calculado: periodicidade. Manual: usuário informa. Documento vencido impede envio.",
            "patlasv4proto-uodoc-data-vencimento",
        ],
        [
            "Documento",
            "Status de validação",
            "Pendente; Aprovado; Recusado; Vencido",
            "Sim",
            "Sistema / MTI",
            "Grade",
            "Pendente/Vencido são definidos pelo sistema; Aprovado/Recusado pela MTI. Obter vencimento não aprova o documento.",
            "patlasv4proto-uodoc-status-validacao",
        ],
        [
            "Documento",
            "Modo de vencimento",
            "Automático; Calculado; Manual",
            "Não",
            "Somente leitura",
            "Oculto",
            "Herdado do Tipo de Documento. Automático = OCR; Calculado = periodicidade; Manual = data informada.",
            "patlasv4proto-uodoc-modo-vencimento",
        ],
        [
            "Documento",
            "Periodicidade",
            "Anual; Semestral; Trimestral; Mensal",
            "Condicional",
            "Somente leitura",
            "Oculto",
            "Herdada do tipo e utilizada apenas quando Modo = Calculado.",
            "patlasv4proto-uodoc-periodicidade",
        ],
        [
            "Documento",
            "Origem do vencimento",
            "OCR (automático); Calculado (periodicidade); Manual; Manual (correção OCR)",
            "Não",
            "Somente leitura",
            "Oculto",
            "Registra como a data de vencimento foi definida.",
            "patlasv4proto-uodoc-origem-vencimento",
        ],
        [
            "Documento",
            "OCR — revisão necessária",
            "Alerta",
            "Condicional",
            "—",
            "Origem = Manual (correção OCR)",
            "Exibido quando o OCR não encontra a validade e exige correção manual.",
            "patlasv4proto-uodoc-ocr-alerta",
        ],
    ],
    "Execução da parceria": [
        [
            "Execução",
            "Tipo",
            "Papel timbrado; Contrato de Parceria; Plano de Negócio; Plano de Operação; Logo Marca / LOGO",
            "Sim",
            "MTI / Parceiro",
            "Etapa Execução da parceria",
            "Documento operacional. Não participa da habilitação MIPP.",
            "patlasv4proto-uodexec-tipo",
        ],
        [
            "Execução",
            "Arquivo",
            "PDF / DOCX",
            "Sim",
            "MTI / Parceiro",
            "Grade",
            "Upload do modelo operacional.",
            "patlasv4proto-uodexec-arquivo",
        ],
        [
            "Execução",
            "Data do anexo",
            "Data",
            "—",
            "Automática",
            "Grade",
            "Preenchida automaticamente no upload.",
            "patlasv4proto-uodexec-data-anexo",
        ],
        [
            "Execução",
            "Ativo",
            "Sim; Não",
            "Sim",
            "MTI / Parceiro",
            "Grade",
            "Documentos inativos não devem ser utilizados em novos processos.",
            "patlasv4proto-uodexec-ativo",
        ],
        [
            "Execução",
            "Observações",
            "Texto longo",
            "Não",
            "MTI / Parceiro",
            "Detalhes",
            "Informação complementar do documento.",
            "patlasv4proto-uodexec-observacao",
        ],
        [
            "Execução",
            "URL",
            "Texto / URL pública",
            "Condicional",
            "MTI / Parceiro",
            "Grade",
            "Obrigatória para logo institucional; opcional para os demais tipos.",
            "patlasv4proto-uodexec-url",
        ],
    ],
}

RULES = [
    ["RN-DOC-01", "Na Organização seleciona-se Etapa, nunca Grupo avulso."],
    ["RN-DOC-02", "Grupos e tipos são copiados do catálogo Etapa documentacional ao clicar em Adicionar."],
    ["RN-DOC-03", "O parceiro anexa arquivos, mas não cria novos tipos de documento."],
    ["RN-DOC-04", "Progresso = quantidade válida / total exigido, calculado por grupo e etapa."],
    ["RN-DOC-05", "Status da Organização é independente do Status de cada Documento."],
    ["RN-DOC-06", "Documento obrigatório Vencido impede o envio do cadastro."],
    ["RN-DOC-07", "Pendente/Vencido são estados sistêmicos; Aprovado/Recusado são decisões humanas da MTI na Fase 1."],
    ["RN-DOC-08", "A extração ou cálculo do vencimento não aprova o conteúdo do documento."],
    ["RN-DOC-09", "NDA controla a visibilidade, pelo parceiro, dos anexos incluídos pela MTI."],
    ["RN-DOC-10", "Acesso comercial = Sim significa liberado; normalmente ocorre após Aprovada pela MTI."],
    ["RN-DOC-11", "Etapas e campos jurídicos da raiz ficam ocultos quando Organização raiz = Não."],
    ["RN-DOC-12", "Se uma seção ficar sem campos visíveis, o accordion inteiro deve ser ocultado."],
    ["RN-DOC-13", "Documentos de execução pertencem à etapa Execução da parceria e não à grade MIPP."],
    ["RN-DOC-14", "Etapas e grupos inativos não aparecem em novos cadastros; instâncias históricas são preservadas."],
]


def style_sheet(ws, widths=None):
    dark_blue = "173F67"
    light_blue = "EAF4FB"
    white = "FFFFFF"
    border_color = "CBD5E1"
    thin = Side(style="thin", color=border_color)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 28

    for row in range(2, ws.max_row + 1):
        if row % 2 == 0:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=light_blue)
        for cell in ws[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

    default_widths = [18, 34, 48, 16, 23, 28, 88, 48]
    for index, width in enumerate(widths or default_widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def add_data_sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    return ws


def add_readme(wb):
    ws = wb.active
    ws.title = "Leia-me"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 105

    entries = [
        ["PLANILHA CANÔNICA", "ABA DOCUMENTOS DA ORGANIZAÇÃO"],
        ["Modelo", "Etapa documentacional → Grupo de documentos → Tipo/Documento"],
        ["Fonte", "Protótipo Atlas · forms.json · modelo vigente em 30/07/2026"],
        ["Arquivo de origem", "data/subprojects/atlas-prototipo/epics/prototipo/forms.json"],
        ["Objetivo", "Consolidar campos, visibilidade, obrigatoriedade e regras de negócio da tela atual."],
        ["Mudança principal", "Na Organização seleciona-se Etapa. Os grupos e tipos vêm do catálogo e não são escolhidos manualmente."],
        ["Status da habilitação", "Incompleta → Em apresentação → Completa – aguardando MTI → Aprovada pela MTI / Reprovada pela MTI"],
        ["Atenção", "Acesso comercial usa semântica positiva: Sim = acesso liberado."],
        ["Legado", "O bloco antigo Documentos de execução da parceria deve ser substituído pela etapa Execução da parceria."],
    ]
    for row in entries:
        ws.append(row)

    ws.merge_cells("A1:B1")
    ws["A1"] = "PLANILHA CANÔNICA — DOCUMENTOS DA ORGANIZAÇÃO"
    ws["A1"].fill = PatternFill("solid", fgColor="173F67")
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).font = Font(bold=True, color="173F67")
        ws.cell(row, 1).fill = PatternFill("solid", fgColor="EAF4FB")
        for column in (1, 2):
            ws.cell(row, column).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row, column).border = Border(
                bottom=Side(style="thin", color="CBD5E1")
            )


def add_rules_sheet(wb):
    ws = wb.create_sheet("Regras de negócio")
    ws.append(["ID", "Regra"])
    for row in RULES:
        ws.append(row)
    style_sheet(ws, [18, 120])
    return ws


def validate_workbook(path):
    workbook = load_workbook(path, data_only=False)
    expected = ["Leia-me", *SHEETS.keys(), "Regras de negócio"]
    assert workbook.sheetnames == expected
    assert all(workbook[name].max_row > 1 for name in expected)
    return sum(workbook[name].max_row - 1 for name in expected if name != "Leia-me")


def main():
    workbook = Workbook()
    add_readme(workbook)
    for sheet_name, rows in SHEETS.items():
        add_data_sheet(workbook, sheet_name, rows)
    add_rules_sheet(workbook)

    workbook.properties.title = "Planilha Canônica — Documentos da Organização"
    workbook.properties.subject = "Campos e regras de negócio do protótipo Atlas"
    workbook.properties.creator = "Atlas / Especificador Sydle"
    workbook.properties.created = datetime.now()
    workbook.save(OUTPUT)

    total = validate_workbook(OUTPUT)
    print(f"Gerado: {OUTPUT}")
    print(f"Abas: {len(workbook.sheetnames)} | Linhas de especificação: {total}")


if __name__ == "__main__":
    main()
