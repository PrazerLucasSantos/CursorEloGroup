# -*- coding: utf-8 -*-
"""PPTX — Minuta de requisitos Catálogo & Produtos (espelha o HTML v0.9)."""
from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_CANDIDATES = [
    Path(r"C:\Users\LucasSantos\OneDrive - EloGroup\Área de Trabalho\docs\entregaveis\pptx\Atlas requisitos.pptx"),
    ROOT / "docs" / "entregaveis" / "pptx" / "Atlas requisitos.pptx",
]
OUT = ROOT / "exports" / "Requisitos_Finais_Catalogo_Produtos_Atlas.pptx"

BLUE = RGBColor(0x18, 0x1F, 0xDB)
DARK = RGBColor(0x27, 0x27, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xBF, 0xBF, 0xBF)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
SOFT = RGBColor(0xEE, 0xF0, 0xFF)
GREY_HDR = RGBColor(0xD9, 0xD9, 0xD9)
GREEN = RGBColor(0x16, 0x65, 0x34)
ORANGE = RGBColor(0xC2, 0x41, 0x0C)
RED = RGBColor(0x99, 0x1B, 0x1B)
PARC = RGBColor(0x25, 0x63, 0xEB)
CLI = RGBColor(0x7C, 0x3A, 0xED)

F_TITLE = "PP Telegraf"
F_LABEL = "PP Telegraf SemiBold"
F_BODY = "Arial"


def find_template() -> Path:
    for p in TEMPLATE_CANDIDATES:
        if p.exists() and p.stat().st_size > 100_000:
            return p
    for p in TEMPLATE_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError("Template Atlas requisitos.pptx não encontrado")


def clear_slides(prs: Presentation) -> None:
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        r_id = sld.get(qn("r:id"))
        try:
            prs.part.drop_rel(r_id)
        except Exception:
            pass
        xml_slides.remove(sld)


def set_fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def bottom_border(cell, rgb=LINE, w=6350):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("a:lnB"))
    if existing is not None:
        tc_pr.remove(existing)
    ln = tc_pr.makeelement(qn("a:lnB"), {"w": str(w), "cap": "flat"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (rgb[0], rgb[1], rgb[2])})
    fill.append(clr)
    ln.append(fill)
    tc_pr.append(ln)


def style_cell(
    cell,
    text,
    *,
    bold=False,
    size=10,
    color=DARK,
    font=F_BODY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    tf = cell.text_frame
    tf.word_wrap = True
    lines = str(text or "").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(1)
        p.space_before = Pt(0)
        for r in list(p.runs):
            r.text = ""
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color


def strip_html(s: str) -> str:
    return (
        str(s or "")
        .replace("<strong>", "")
        .replace("</strong>", "")
        .replace("<em>", "")
        .replace("</em>", "")
        .replace("<br/>", "\n")
        .replace("<br>", "\n")
        .replace("&amp;", "&")
    )


class Deck:
    def __init__(self):
        tpl = find_template()
        self.prs = Presentation(str(tpl))
        clear_slides(self.prs)
        layouts = {l.name: l for l in self.prs.slide_layouts}
        self.lay_white = layouts.get("Todo branco") or layouts.get("Blank") or self.prs.slide_layouts[0]
        self.lay_black = layouts.get("Todo preto") or self.lay_white
        self.lay_capa = layouts.get("Capa tema preto") or self.lay_black
        self.sw = self.prs.slide_width
        self.n = 0

    def add(self, layout=None):
        self.n += 1
        return self.prs.slides.add_slide(layout or self.lay_white)

    def title(self, slide, text):
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.16), Inches(12.2), Inches(0.42))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.name = F_TITLE
        r.font.color.rgb = DARK

    def section(self, title: str, subtitle: str):
        slide = self.add(self.lay_black)
        # ensure dark bg if layout is blank-ish
        try:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(0x22, 0x22, 0x22)
        except Exception:
            pass
        for y, txt, sz, bold in ((2.4, title, 32, True), (3.5, subtitle, 14, False)):
            box = slide.shapes.add_textbox(Inches(0.6), Inches(y), Inches(11.8), Inches(0.85))
            r = box.text_frame.paragraphs[0].add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.name = F_TITLE if bold else F_BODY
            r.font.color.rgb = WHITE

    def capa(self):
        slide = self.add(self.lay_capa)
        box = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(11.5), Inches(3.2))
        tf = box.text_frame
        tf.word_wrap = True
        lines = [
            ("Documentação de Requisitos", 28, True),
            ("Catálogo e Produtos", 28, True),
            ("", 12, False),
            ("Minuta para validação · v0.9 — 04/08/2026", 14, False),
            ("Projeto Atlas · MTI Simplifica", 13, False),
            ("Backoffice · Portal do Parceiro · Portal do Cliente", 12, False),
            ("", 10, False),
            ("Base: Discovery 08/07/2026 · protótipo · portais", 11, False),
            ("Origem: Confirmado · Protótipo · Proposto · QA", 11, False),
        ]
        first = True
        for text, sz, bold in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = PP_ALIGN.CENTER
            p.space_after = Pt(4)
            if not text:
                continue
            r = p.add_run()
            r.text = text
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.name = F_TITLE if bold else F_BODY
            r.font.color.rgb = WHITE

    def story(
        self,
        title: str,
        necessidade: str,
        contexto: str,
        requisitos: list[str],
        regras: list[str],
        extra_req_header: str | None = None,
        extra_reqs: list[str] | None = None,
    ):
        """Layout padrão Atlas: Necessidade | Como é usado | Requisitos e Regras."""
        slide = self.add()
        self.title(slide, title)
        tbl = slide.shapes.add_table(3, 2, Inches(0.4), Inches(0.62), Inches(12.4), Inches(6.05)).table
        tbl.columns[0].width = Inches(2.05)
        tbl.columns[1].width = Inches(10.35)
        labels = ["Necessidade", "Como é usado / contexto", "Requisitos e Regras"]
        for i, lab in enumerate(labels):
            set_fill(tbl.cell(i, 0), BLUE)
            style_cell(
                tbl.cell(i, 0),
                lab,
                bold=True,
                size=11,
                color=WHITE,
                font=F_LABEL,
                anchor=MSO_ANCHOR.TOP,
            )
            bottom_border(tbl.cell(i, 0))
            bottom_border(tbl.cell(i, 1))
        style_cell(tbl.cell(0, 1), strip_html(necessidade), size=10, anchor=MSO_ANCHOR.TOP)
        style_cell(tbl.cell(1, 1), strip_html(contexto), size=10, anchor=MSO_ANCHOR.TOP)

        c2 = tbl.cell(2, 1)
        c2.vertical_anchor = MSO_ANCHOR.TOP
        tf = c2.text_frame
        tf.word_wrap = True

        def par(text, *, bold=False, bullet=False, first=False, color=DARK, size=9):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            p.space_before = Pt(1 if not first else 0)
            p.space_after = Pt(0)
            r = p.add_run()
            r.text = ("•  " if bullet else "") + strip_html(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = F_BODY
            r.font.color.rgb = color

        par("Requisitos funcionais", bold=True, color=BLUE, first=True, size=10)
        for req in requisitos:
            par(req, bullet=True)
        if regras:
            par("Regras de negócio", bold=True, color=BLUE, size=10)
            for rg in regras:
                par(rg, bullet=True)
        if extra_req_header and extra_reqs:
            par(extra_req_header, bold=True, color=BLUE, size=10)
            for x in extra_reqs:
                par(x, bullet=True)

    def table_slide(self, title: str, headers: list[str], rows: list[list[str]], note: str = ""):
        slide = self.add()
        self.title(slide, title)
        y0 = Inches(0.62)
        if note:
            box = slide.shapes.add_textbox(Inches(0.45), Inches(0.55), Inches(12.2), Inches(0.35))
            r = box.text_frame.paragraphs[0].add_run()
            r.text = note
            r.font.size = Pt(10)
            r.font.name = F_BODY
            r.font.color.rgb = MUTED
            y0 = Inches(0.92)

        n_rows = 1 + len(rows)
        n_cols = len(headers)
        height = min(Inches(0.32) * n_rows + Inches(0.15), Inches(5.9))
        tbl = slide.shapes.add_table(n_rows, n_cols, Inches(0.4), y0, Inches(12.4), height).table

        # column widths
        widths = {
            2: (2.2, 10.2),
            3: (2.0, 7.5, 2.9),
            4: (3.0, 2.0, 1.2, 6.2),
        }.get(n_cols)
        if widths:
            for i, w in enumerate(widths):
                tbl.columns[i].width = Inches(w)
        else:
            each = 12.4 / n_cols
            for i in range(n_cols):
                tbl.columns[i].width = Inches(each)

        for j, h in enumerate(headers):
            set_fill(tbl.cell(0, j), GREY_HDR)
            style_cell(tbl.cell(0, j), h, bold=True, size=10, font=F_LABEL)
            bottom_border(tbl.cell(0, j))
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                style_cell(tbl.cell(i, j), strip_html(val), size=9, anchor=MSO_ANCHOR.TOP)
                bottom_border(tbl.cell(i, j))

    def bullets_slide(self, title: str, items: list[str], lead: str = ""):
        slide = self.add()
        self.title(slide, title)
        y = 0.6
        if lead:
            box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12.2), Inches(0.5))
            r = box.text_frame.paragraphs[0].add_run()
            r.text = lead
            r.font.size = Pt(11)
            r.font.name = F_BODY
            r.font.color.rgb = MUTED
            y = 1.1
        box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12.2), Inches(5.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(3)
            p.space_after = Pt(2)
            r = p.add_run()
            r.text = "•  " + strip_html(item)
            r.font.size = Pt(12)
            r.font.name = F_BODY
            r.font.color.rgb = DARK

    def flow_status_slide(self):
        slide = self.add()
        self.title(slide, "Status do fluxo (espelho portal ↔ MTI)")
        steps = [
            ("Rascunho", MUTED),
            ("Aguardando", BLUE),
            ("Ajuste", ORANGE),
            ("Homologado", GREEN),
            ("Publicado", GREEN),
        ]
        x = Inches(0.45)
        y = Inches(0.85)
        bw, bh, gap = Inches(2.0), Inches(0.7), Inches(0.22)
        for i, (lab, color) in enumerate(steps):
            sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh)
            sh.fill.solid()
            sh.fill.fore_color.rgb = SOFT
            sh.line.color.rgb = color
            tf = sh.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            r.text = lab
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.name = F_BODY
            r.font.color.rgb = color
            x += bw + gap
            if i < len(steps) - 1:
                ar = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    x - gap + Inches(0.02),
                    y + Inches(0.26),
                    Inches(0.18),
                    Inches(0.18),
                )
                ar.fill.solid()
                ar.fill.fore_color.rgb = BLUE
                ar.line.fill.background()

        # alt outcomes
        box = slide.shapes.add_textbox(Inches(0.45), Inches(1.8), Inches(12.2), Inches(0.35))
        r = box.text_frame.paragraphs[0].add_run()
        r.text = "Desvios: Reprovado · Paralisado  →  parceiro pode abrir novo rascunho"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.name = F_BODY
        r.font.color.rgb = RED

        rows = [
            ["Rascunho / Ajuste", "Enviar à MTI", "Aguardando", "Proposto · ≥1 produto; trava edição"],
            ["Aguardando", "Solicitar ajuste", "Ajuste", "Proposto · justificativa"],
            ["Aguardando", "Aprovar", "Homologado", "Proposto"],
            ["Homologado", "Publicar", "Ativo (publicado)", "Proposto — Homologado→Publicado a validar"],
            ["Aguardando", "Reprovar", "Reprovado", "Proposto"],
            ["Elegível", "Paralisar", "Paralisado", "Proposto · bloqueia novos usos"],
            ["Reprovado / Paralisado", "Novo rascunho", "Rascunho", "Proposto"],
        ]
        tbl = slide.shapes.add_table(1 + len(rows), 4, Inches(0.4), Inches(2.15), Inches(12.4), Inches(4.35)).table
        hdrs = ["De", "Ação", "Para", "Validação / origem"]
        for j, h in enumerate(hdrs):
            tbl.columns[j].width = Inches([2.4, 2.6, 2.2, 5.2][j])
            set_fill(tbl.cell(0, j), GREY_HDR)
            style_cell(tbl.cell(0, j), h, bold=True, size=10, font=F_LABEL)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                style_cell(tbl.cell(i, j), v, size=9, anchor=MSO_ANCHOR.TOP)
                bottom_border(tbl.cell(i, j))

    def lanes_slide(self):
        slide = self.add()
        self.title(slide, "Quem faz o quê — três canais")
        lanes = [
            (BLUE, "MTI (backoffice)", [
                "Listas básicas",
                "Parceria e solução",
                "Catálogo universal",
                "Cadastro interno",
                "Análise do parceiro",
                "Publicar / versionar",
                "Apostilar no contrato",
            ]),
            (PARC, "Parceiro (portal)", [
                "Rascunho do catálogo",
                "Produtos manual/CSV",
                "Visão planilha",
                "Enviar à MTI",
                "Corrigir após ajuste",
                "Status espelhado",
                "Não publica / não apostila",
            ]),
            (CLI, "Cliente (portal)", [
                "Versão apostilada",
                "Consulta produtos",
                "PJ + secretarias/unidades",
                "Orçamento (serviço prévio)",
                "Autorizar → OS + snapshot",
                "Não cadastra catálogo",
                "Cotação → demanda (proposto)",
            ]),
        ]
        w, g = Inches(3.9), Inches(0.25)
        x0 = Inches(0.45)
        for i, (color, head, items) in enumerate(lanes):
            x = x0 + i * (w + g)
            hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(0.7), w, Inches(0.45))
            hdr.fill.solid()
            hdr.fill.fore_color.rgb = color
            hdr.line.fill.background()
            tf = hdr.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = tf.paragraphs[0].add_run()
            r.text = head
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.name = F_BODY
            r.font.color.rgb = WHITE

            body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.15), w, Inches(5.4))
            body.fill.solid()
            body.fill.fore_color.rgb = WHITE
            body.line.color.rgb = LINE
            tf = body.text_frame
            tf.word_wrap = True
            for j, it in enumerate(items):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.space_before = Pt(6)
                r = p.add_run()
                r.text = f"{j+1}.  {it}"
                r.font.size = Pt(12)
                r.font.name = F_BODY
                r.font.color.rgb = DARK

    def build(self):
        self.capa()

        # INTRO
        self.section("Introdução", "Objetivo · Como ler · Escopo")
        self.table_slide(
            "Origem do requisito (legenda)",
            ["Marca", "Significado"],
            [
                ["Confirmado (Discovery)", "Fala explícita na reunião 08/07/2026"],
                ["Protótipo", "Presente em forms/portais; validar com MTI"],
                ["Proposto", "Workflow/UI candidata — não fechada na Discovery"],
                ["QA / teste", "Critério de verificação; não é regra de negócio"],
            ],
            note="Status do fluxo = proposto (andamento portal↔MTI), distinto do status do objeto (Ativo / Homologado / Paralisado).",
        )
        self.story(
            "Objetivo e escopo",
            "A MTI precisa cadastrar, versionar, homologar e disponibilizar catálogos/produtos para contratação nos modelos Tipo 1, 2 e 3 (CGS), com participação do parceiro e consumo pelo cliente.",
            "Cardápio comercial: listas → parceria/solução → catálogo/produtos → parecer → publicação → apostila → consumo. Lei 14.133 / SGDI citadas na Discovery — base normativa pendente de validação jurídica.\n\nConfiguração (agora): cadastros, CSV, parecer, versão, fator, portais.\nConsumo (Discovery/proposta): objeto no contrato, apostila, demanda/orçamento/OS com snapshot.",
            [
                "RF-ESC-01: Cobrir backoffice MTI, Portal do Parceiro e Portal do Cliente.",
                "RF-ESC-02: Classes do protótipo — origem protótipo.",
                "RF-ESC-03: Espelhar status do fluxo portal↔MTI — proposto/protótipo.",
                "RF-ESC-04: Preparar vínculo contrato/OS (versão + apostila + snapshot).",
            ],
            [
                "RN-ESC-01: Parceiro não publica nem cria versão contratual — proposto.",
                "RN-ESC-02: Cliente nunca cadastra catálogo/produto — proposto.",
                "RN-ESC-03: Produto originado do parceiro exige validação/parecer MTI.",
            ],
        )

        # NEGOCIO
        self.section("Negócio e configuração", "Tipos · Métricas · Papéis · Ordem de cadastro")
        self.story(
            "Três formas de contratar (Tipos 1, 2 e 3)",
            "A contratação de TI não se resume a “um item = um preço”. Na Discovery: vínculo com Lei 14.133 e orientação SGDI (base normativa a validar).",
            "Flags Universal / Individualizado e a métrica definem elegibilidade. O tipo de contratação decorre do objeto definido no contrato (não de escolha isolada no cadastro).",
            [
                "RF-TIPO-01: Indicadores Universal e Individualizado (independentes).",
                "RF-TIPO-02: Exigir métrica cadastrada em todo produto.",
                "RF-TIPO-03: Calcular/exibir fator quando Universal=Sim e houver moeda Tipo 3.",
            ],
            [
                "RN-TIPO-01: Tipo 1 — direito só ao item contratado.",
                "RN-TIPO-02: Tipo 2 — saldo (USN/UST/HST) consome itens do catálogo na métrica.",
                "RN-TIPO-03: Tipo 3 — créditos dão acesso a produtos ativos universais.",
                "RN-TIPO-04: Tipo 3 só com métricas universais (USN/UST/HST) — elegibilidade; não pré-condição contratual genérica.",
                "RN-TIPO-05: Só Individualizado (Universal=Não) → sem fator Tipo 3.",
                "RN-TIPO-06: Forma de contratação decorre do objeto contratado (não “MTI escolhe o tipo” no cadastro).",
            ],
        )
        self.table_slide(
            "Resumo dos tipos de contratação",
            ["Tipo", "O que se contrata", "O que se consome"],
            [
                ["1", "Produto/licença específica", "Somente aquele item"],
                ["2", "Objeto-moeda ligado a um catálogo", "Itens do catálogo na métrica"],
                ["3 · CGS", "MTI Créditos de TI e/ou Serviço", "Produtos ativos universais (com fator)"],
            ],
        )
        self.story(
            "Métricas universais (USN, UST, HST)",
            "Sem métrica padronizada não há saldo de contrato nem conversão no Tipo 3.",
            "MTI cadastra a lista. Todo produto escolhe uma. No Tipo 3, o fator usa a moeda universal correspondente.",
            [
                "RF-MET-01: Cadastro parametrizável (sigla + descrição).",
                "RF-MET-02: Produto referencia métrica (não texto livre).",
            ],
            [
                "RN-MET-01: USN → licenciamento; fator tipicamente 1 no Tipo 3.",
                "RN-MET-02: UST → serviço técnico; não linear com hora.",
                "RN-MET-03: HST → hora de serviço técnico (~hora de profissional).",
                "RN-MET-04: “Horas” legadas migram para HST.",
            ],
        )
        self.lanes_slide()
        self.table_slide(
            "Matriz de permissões",
            ["Ação", "MTI", "Parceiro", "Cliente"],
            [
                ["Listas / parceria / universal", "Sim", "Não", "Não"],
                ["Rascunho catálogo/produto", "Sim", "Sim (sua)", "Não"],
                ["CSV import/export", "Sim", "Sim", "Não"],
                ["Enviar análise", "Interno", "Sim", "Não"],
                ["Aprovar / ajuste / reprovar", "Sim", "Não", "Não"],
                ["Publicar / nova versão", "Sim", "Não*", "Não"],
                ["Apostilar contrato", "Sim", "Não", "Não"],
                ["Consultar catálogo", "Sim", "Limitado", "Sim (apostilado)"],
                ["Cotação / OS", "Apoia", "Pode orçar", "Sim"],
            ],
            note="* Parceiro pode abrir novo rascunho após reprovação; versão contratual é governança MTI.",
        )
        self.bullets_slide(
            "Como configurar — ordem recomendada",
            [
                "1. Listas básicas (MTI): Métrica, Cobrança, Categoria, Grupo, Modelo de Venda.",
                "2. Organização parceiro (Fase 1) — CNPJ homologado.",
                "3. Parceria vinculada à Organização (N parcerias / mesmo CNPJ).",
                "4. Solução — fabricante em texto + documentos de apoio.",
                "5. Catálogo — identificador + versão; cobrança; Unidade DTIC; parceria.",
                "6. Produtos — página única; 1 catálogo/produto = candidata; CSV. Primeiro carregamento: parceira + validação MTI.",
                "7. Dados de Parceria — custo, markup, distribuição % (=100) — protótipo.",
                "8. Catálogo Universal — agrega catálogos elegíveis ao Tipo 3.",
                "9. Parecer / Publicar / Nova versão / Apostila no contrato.",
            ],
            lead="Markup NÃO fica no Produto — fica em Dados de Parceria (protótipo).",
        )

        # FLUXO
        self.section("Fluxo operacional", "Passo a passo · Status · Validações · Fator")
        self.bullets_slide(
            "Fluxo ponta a ponta",
            [
                "1. Configurar base (MTI) — listas, parceria, solução.",
                "2. Montar catálogo — MTI no backoffice ou parceiro no portal (rascunho).",
                "3. Incluir produtos — manual, planilha ou CSV (template oficial).",
                "4. Definir markup — Dados de Parceria.",
                "5. Enviar à MTI — status Aguardando; edição travada.",
                "6. Analisar — Aprovar · Ajuste (com motivo) · Reprovar · Paralisar.",
                "7. Loop de ajuste — parceiro corrige → reenvia.",
                "8. Publicar — só após Homologado → Ativo (publicado).",
                "9. Nova versão — copia produtos; congela a anterior.",
                "10. Apostilar no contrato — cliente passa a ver a versão.",
                "11. Consumo — cotação/demanda → orçamento → OS (snapshot).",
            ],
        )
        self.flow_status_slide()
        self.story(
            "Ações do fluxo — resumo operacional",
            "Cada etapa precisa de ação clara e validação, para o parceiro não furar a governança e o cliente não ver versão errada.",
            "Parceiro opera no portal; analista MTI no backoffice (análise + publicar + apostilar); cliente só consome no portal dele.",
            [
                "RF-FLUX-01: Histórico de status — proposto/protótipo.",
                "RF-FLUX-02: Justificativa MTI no portal (Ajuste/Reprovado/Paralisado) — proposto.",
                "RF-FLUX-03: Publicar após parecer favorável — proposto (Homologado→Publicado a validar).",
                "RF-FLUX-04: Apostila libera visão do cliente sem alterar conteúdo versionado — confirmado.",
            ],
            [
                "RN-FLUX-01: Produtos da parceira exigem validação/parecer MTI — confirmado.",
                "RN-FLUX-02: Ajuste/reprovação com justificativa — proposto.",
                "RN-FLUX-03: Publicação após homologação — proposto.",
                "RN-VER-03: Cliente vê nova versão só após apostilamento — confirmado.",
            ],
            extra_req_header="Validações",
            extra_reqs=[
                "Enviar: ≥1 produto — proposto.",
                "Ajuste/Reprovar: justificativa — proposto.",
                "Fator: Universal=Sim + moeda; senão oculto/nulo — conceito confirmado; UI proposta.",
                "Distribuição % = 100 — origem protótipo.",
                "CSV: template round-trip — confirmado; rejeição por linha a especificar.",
            ],
        )
        self.table_slide(
            "Fator de conversão (Tipo 3) — estrutura (exemplo ilustrativo)",
            ["ID / Variável", "Regra / valor"],
            [
                ["RN-FAT-01", "Fator a partir de unitário e moeda universal (fórmula oficial a fechar)"],
                ["RN-FAT-02", "USN no Tipo 3: relação tipicamente 1:1"],
                ["RN-FAT-03", "Consumo: moeda × fator × qtde [× complexidade] [× peso]"],
                ["Valor unitário", "Preço do produto (exemplo — não oficializar)"],
                ["Moeda universal", "Valor da moeda Tipo 3 (exemplo — não oficializar)"],
                ["Fator", "unitário ÷ moeda (não oficializar coeficiente até planilha/áudio)"],
                ["Consumo", "fator × qtde [× coef.] [× peso]"],
            ],
            note="Não usar R$ 22.679 / R$ 282,58 / fator 80,26 como valores oficiais até validação. Botão Recalcular = UX protótipo.",
        )

        # CLASSES
        self.section("Classes e cadastros", "Campos · RF · RN · Ações · Validações")
        self.table_slide(
            "Mapa de classes do domínio",
            ["#", "Classe", "Quem configura", "Depende de"],
            [
                ["1", "Métrica", "MTI", "—"],
                ["2", "Tipo de Cobrança", "MTI", "—"],
                ["3", "Categoria de Serviços", "MTI", "—"],
                ["4", "Grupo", "MTI", "—"],
                ["5", "Modelo de Venda", "MTI", "—"],
                ["6", "Parceria", "MTI", "Organização (F1)"],
                ["7", "Solução", "MTI", "Parceria"],
                ["8", "Catálogo", "MTI / Parceiro", "Listas + Parceria"],
                ["9", "Produto", "MTI / Parceiro", "Catálogo + listas"],
                ["10", "Dados de Parceria por Produto", "MTI (/parceiro)", "Produto + Solução"],
                ["11", "Catálogo Universal / objetos", "MTI", "Catálogos + objetos contratados"],
            ],
        )

        classes = [
            (
                "Métrica",
                "Padronizar a unidade comercial dos produtos e do saldo contratual.",
                "Cadastro de sigla (USN, UST, HST…) e descrição. MTI mantém; produto referencia.",
                [
                    "RF-MET-01: Criar, editar e listar (parametrizável).",
                    "RF-MET-02: Produto referencia métrica (não texto livre).",
                    "RF-MET-03: Sigla única — proposto/integridade.",
                    "RF-MET-04: Não excluir se usada em produto homologado — proposto.",
                ],
                ["RN-MET-01…04: ver métricas universais (Discovery; siglas a conferir)."],
                [["Identificador (sigla)", "Texto", "Sim", "USN, UST, HST…"], ["Descrição", "Texto", "Não", ""]],
            ),
            (
                "Tipo de Cobrança",
                "Informar como o item/catálogo é cobrado.",
                "Modelo + recorrência. Obrigatório em Catálogo e Produto. Não destacar no card do produto.",
                [
                    "RF-COB-01: CRUD da lista.",
                    "RF-COB-02: Obrigatório em Catálogo e Produto.",
                    "RF-COB-03: Não excluir se referenciado homologado.",
                ],
                ["RN-COB-01: Cobrança no formulário; não no card."],
                [["Modelo", "Texto", "Sim", "Sob Demanda, Mensal…"], ["Recorrência", "Texto", "Não", ""]],
            ),
            (
                "Categoria de Serviços",
                "Classificar produtos para filtro e organização.",
                "Opcional no Produto; filtro na visão planilha.",
                [
                    "RF-CATG-01: CRUD.",
                    "RF-CATG-02: Vincular opcionalmente ao Produto.",
                    "RF-CATG-03: Filtro na planilha.",
                    "RF-CATG-04: Não excluir se homologada em uso.",
                ],
                ["RN-CATG-01: Não excluir categoria referenciada homologada."],
                [["Identificador", "Texto", "Sim", ""], ["Descrição", "Texto", "Não", ""]],
            ),
            (
                "Grupo",
                "Agrupar produtos para leitura rápida na lista.",
                "Obrigatório no Produto. No card: nome + tipo + tag de status (Discovery); Grupo no card = candidata.",
                [
                    "RF-GRP-01: CRUD.",
                    "RF-GRP-02: Obrigatório no Produto.",
                    "RF-GRP-03: Destacar no card — candidata.",
                    "RF-GRP-04: Filtrar na planilha.",
                ],
                ["RN-GRP-01: Grupo obrigatório; visibilidade no card = a confirmar."],
                [["Identificador", "Texto", "Sim", ""], ["Descrição", "Texto", "Não", ""]],
            ),
            (
                "Modelo de Venda",
                "Indicar o modelo comercial do produto.",
                "Obrigatório no Produto (Por Licença, Serviço, Homologação…).",
                ["RF-MV-01: CRUD.", "RF-MV-02: Obrigatório no Produto."],
                ["RN-MV-01: Modelo de venda obrigatório no produto."],
                [["Identificador", "Texto", "Sim", ""], ["Descrição", "Texto", "Não", ""]],
            ),
            (
                "Parceria",
                "Separar o acordo comercial da Organização (CNPJ). Mesma empresa pode ter N parcerias.",
                "Primeiro cadastro da entidade Parceria = MTI. Não confundir com primeiro carregamento de produtos (parceira + validação MTI). Define escopo do Portal.",
                [
                    "RF-PAR-01: Criar/editar no backoffice.",
                    "RF-PAR-02: Vincular Organização.",
                    "RF-PAR-03: N parcerias no mesmo CNPJ.",
                    "RF-PAR-04: Status Ativa/Homologada/Paralisada.",
                    "RF-PAR-05: Associar soluções.",
                    "RF-PAR-06: Restringir portal ao escopo da parceria.",
                ],
                [
                    "RN-PAR-01: Primeiro cadastro da entidade Parceria = MTI.",
                    "RN-PAR-02: Parceria ≠ Organização.",
                    "RN-PAR-03: Paralisada bloqueia novos produtos na fila.",
                ],
                [
                    ["Identificador", "Texto", "Sim", "MTI SIMPLIFICA…"],
                    ["Parceiro", "Ref. Org", "Sim", "CNPJ"],
                    ["Status", "Lista", "Sim", "Ativa/Homologada/…"],
                    ["Soluções", "Ref. múltipla", "Sim", ""],
                    ["Observações", "Texto", "Não", ""],
                ],
            ),
            (
                "Solução",
                "Descrever a oferta da parceria com fabricante e documentos.",
                "Fabricante: nome e contato em texto (direção Discovery; não formalizar ausência definitiva de classe até confirmação). Documentos anexáveis.",
                [
                    "RF-SOL-01: Criar/editar vinculada à parceria.",
                    "RF-SOL-02: Fabricante em texto (provisório).",
                    "RF-SOL-03: Anexar documentos.",
                    "RF-SOL-04: Listar no contexto catálogo/demanda.",
                ],
                ["RN-SOL-01: Fabricante em texto (provisório).", "RN-SOL-02: Solução pertence a uma parceria."],
                [
                    ["Identificador", "Texto", "Sim", ""],
                    ["Parceria", "Ref.", "Não*", "Vazio = própria MTI"],
                    ["Fabricante nome/contato", "Texto", "Não", "Texto livre"],
                    ["Documentos de apoio", "Arquivo", "Não", "Múltiplos"],
                ],
            ),
            (
                "Catálogo",
                "Empacotar produtos versionados para contratação e consumo.",
                "Criado pela MTI ou pelo parceiro (rascunho). Importa produtos; envia à análise; MTI publica e versiona. Versão também é atributo do contrato.",
                [
                    "RF-CTL-01: Criar/editar rascunho com versão.",
                    "RF-CTL-02: Importar CSV (template).",
                    "RF-CTL-03: Exportar CSV e PDF.",
                    "RF-CTL-04: Visão planilha com filtros.",
                    "RF-CTL-05: Enviar à análise.",
                    "RF-CTL-06: Aprovar / ajuste / reprovar.",
                    "RF-CTL-07: Publicar após homologar.",
                    "RF-CTL-08: Nova versão sem sobrescrever.",
                    "RF-CTL-09: Paralisar.",
                    "RF-CTL-10: Espelhar status no portal.",
                ],
                [
                    "RN-VER-01: Versão é atributo do contrato.",
                    "RN-VER-02: Nova versão não sobrescreve.",
                    "RN-VER-03: Cliente vê só após apostila.",
                    "RN-CTL-01: Homologada não se edita no lugar.",
                    "RN-CSV-01: Template CSV round-trip.",
                ],
                [
                    ["Identificador / Versão", "Texto", "Sim", "Versão = contrato"],
                    ["Versão anterior", "Texto RO", "Não", "Auto"],
                    ["Preencher códigos manual?", "Sim/Não", "Sim", "Default Não"],
                    ["Cobrança / Valor unitário", "Ref / Nº", "Sim", "Pacote Tipo 2"],
                    ["Unidade DTIC", "Ref.", "Sim", ""],
                    ["Parceria", "Ref.", "Não", ""],
                    ["Status objeto / fluxo", "Lista", "Sim/Sist.", "Distintos"],
                ],
            ),
            (
                "Produto",
                "Item comercializável (licença ou serviço) dentro de um catálogo.",
                "Página única. Card Discovery: nome + tipo + tag de status (sem cobrança/métrica). Grupo no card = candidata. Base CSV, fator Tipo 3 e consumo.",
                [
                    "RF-PROD-01: Página única.",
                    "RF-PROD-02: Exigir catálogo, tipo, métrica, cobrança, modelo, valor, flags.",
                    "RF-PROD-03: Categoria opcional — protótipo.",
                    "RF-PROD-04: Moeda + fator se Universal=Sim (Recalcular = UX protótipo).",
                    "RF-PROD-05: Complexidade só se Serviço.",
                    "RF-PROD-06: CSV/PDF.",
                    "RF-PROD-07: Enviar / aprovar / ajuste / inativar.",
                    "RF-PROD-08: Card nome + tipo + tag; Grupo candidata.",
                    "RF-PROD-09: Observações.",
                ],
                [
                    "RN-P-01: Produto referencia um catálogo (1:1 candidata — formalizar).",
                    "RN-P-02: Universal e Individualizado independentes.",
                    "RN-P-03/04: Sem fator se só individualizado ou fora de métrica universal.",
                    "RN-P-05: Markup não fica no produto.",
                    "RN-FAT-01: Fator a partir de unitário e moeda (fórmula a fechar).",
                    "RN-FAT-02: USN Tipo 3 tipicamente 1:1.",
                    "RN-FAT-03: Consumo Tipo 3: moeda × fator × qtde [× complexidade] [× peso].",
                ],
                [
                    ["Nome / SKU / Tipo", "Texto / Lista", "Cand.", "Card: nome+tipo"],
                    ["Grupo / Catálogo", "Ref.", "Cand./Sim", "1 catálogo candidata"],
                    ["Universal / Individualizado", "Sim/Não", "Sim", "Indep."],
                    ["Métrica / Cobrança / Modelo", "Ref.", "Sim", "Não no card"],
                    ["Valor unitário / Fator", "Nº / RO", "Sim/Cond.", "Fórmula a fechar"],
                    ["Complexidade/peso", "Cond.", "Não", "Só Serviço"],
                    ["Unidade DTIC", "Ref.", "Sim", ""],
                ],
            ),
            (
                "Dados de Parceria por Produto",
                "Guardar custo, markup e repartição sem misturar com o preço de venda (classe do protótipo).",
                "Por combinação parceria/solução/produto. Soma % = 100 — origem protótipo; não apresentar como regra fechada só da Discovery.",
                [
                    "RF-DP-01: Criar/editar por produto/solução — protótipo.",
                    "RF-DP-02: Validar % parceiro + % MTI = 100 — protótipo.",
                    "RF-DP-03: Não gravar markup no Produto — protótipo (confirmar vs Discovery).",
                ],
                ["RN-DP-01: Soma = 100% — protótipo.", "RN-P-05: Markup neste cadastro — protótipo."],
                [
                    ["Solução / Vigência", "Ref / Texto", "Sim", "Protótipo"],
                    ["Custo / Markup", "Número", "Sim", "Protótipo"],
                    ["Dist. parceiro % / MTI %", "Número", "Sim", "Soma 100"],
                ],
            ),
            (
                "Catálogo Universal / objetos",
                "Agregar catálogos elegíveis ao Tipo 3 (CGS) e permitir seleção de objetos no contrato.",
                "MTI mantém. No contrato: objetos específicos, objetos universais ou catálogo (Discovery F3).",
                [
                    "RF-UNI-01: Criar/editar.",
                    "RF-UNI-02: Vincular catálogos — ativo = proposto.",
                    "RF-UNI-03: Publicar/paralisar — proposto.",
                    "RF-UNI-04: Bloquear vínculo de catálogo paralisado em novos contratos — proposto.",
                    "RF-UNI-05: No contrato, selecionar fonte: objetos específicos, universais ou catálogo.",
                ],
                ["RN-UNI-01: Só catálogos ativos.", "RN-UNI-02: Objetos universais distintos do agregado de catálogos.", "RN-TIPO-04: Elegibilidade Tipo 3 = USN/UST/HST."],
                [
                    ["Identificador", "Texto", "Sim", ""],
                    ["Status", "Lista", "Sim", ""],
                    ["Catálogos", "Ref. múltipla", "Sim", "Camada 2"],
                ],
            ),
        ]

        for nome, nec, ctx, rf, rn, fields in classes:
            self.story(f"Classe — {nome}", nec, ctx, rf, rn)
            self.table_slide(
                f"{nome} — campos",
                ["Campo", "Tipo", "Obrig.", "Observação"],
                fields,
            )

        # PORTAIS
        self.section("Portal do Parceiro", "Rascunho · CSV · Envio · Ajuste")
        self.story(
            "Portal do Parceiro — visão geral",
            "O parceiro precisa montar e enviar o catálogo sem acessar o backoffice, com status transparente.",
            "Home com KPIs → Meus catálogos → Detalhe → Enviados/ajustes. Análise é da MTI; o portal espelha.",
            [
                "RF-PARC-01: Salvar rascunho da própria parceria.",
                "RF-PARC-02: Produtos manual ou CSV.",
                "RF-PARC-03: Exportar CSV/PDF.",
                "RF-PARC-04: Enviar à MTI (travar edição).",
                "RF-PARC-05: Corrigir após ajuste.",
                "RF-PARC-06: Status espelhado.",
                "RF-PARC-07: Visão planilha + filtros.",
                "RF-PARC-08…14: Template, fila, justificativa, novo rascunho, histórico, KPIs, mensagem.",
            ],
            [
                "RN-PARC-01: Parecer MTI obrigatório (origem portal).",
                "RN-PARC-02: Ajuste com justificativa visível.",
                "RN-PARC-03: Só sua parceria.",
                "RN-PARC-04: Edição só Rascunho/Ajuste.",
                "RN-PARC-05: Enviar exige ≥1 produto.",
                "RN-PARC-06: Não publica/versiona/apostila.",
            ],
        )
        self.table_slide(
            "Portal do Parceiro — RF consolidado",
            ["ID", "Requisito", "Prioridade"],
            [
                ["RF-PARC-01", "Salvar rascunho da própria parceria", "Alta"],
                ["RF-PARC-02", "Incluir produtos manual ou CSV", "Alta"],
                ["RF-PARC-03", "Exportar CSV/PDF", "Alta"],
                ["RF-PARC-04", "Enviar à MTI e travar edição", "Alta"],
                ["RF-PARC-05", "Corrigir após ajuste", "Alta"],
                ["RF-PARC-06", "Status espelhado", "Alta"],
                ["RF-PARC-07", "Visão planilha + filtros", "Alta"],
                ["RF-PARC-08", "Template CSV", "Média"],
                ["RF-PARC-09", "Fila enviados/ajustes", "Alta"],
                ["RF-PARC-10", "Justificativa em destaque", "Alta"],
                ["RF-PARC-11", "Novo rascunho pós-reprovação", "Média"],
                ["RF-PARC-12", "Histórico de status", "Média"],
                ["RF-PARC-13", "KPIs na home", "Baixa"],
                ["RF-PARC-14", "Mensagem ao analista", "Baixa"],
            ],
        )

        self.section("Portal do Cliente", "Consulta · Cotação · Orçamento · OS")
        self.story(
            "Portal do Cliente — visão geral",
            "O cliente precisa consumir o catálogo contratado com segurança de versão, sem cadastrar produtos.",
            "Cliente = PJ; pessoas = responsáveis. Hierarquia permite rateio/segregação por secretaria. Serviço exige orçamento prévio; licença pode dispensar. Cotação gera demanda (proposto).",
            [
                "RF-CLI-01: Consultar versão apostilada.",
                "RF-CLI-02: Demanda / orçamento.",
                "RF-CLI-03: Distinguir OS, orçamento prévio e consumo (serviço: orçamento antes do consumo).",
                "RF-CLI-04: Acompanhar demandas/orçamentos/OS.",
                "RF-CLI-05: Não exibir versão nova até apostila.",
                "RF-CLI-06: Cotação gera demanda — proposto.",
                "RF-CLI-07/08: Aceitar/recusar e autorizar.",
                "RF-CLI-09…14: Contratos, snapshot, itens, motivo, hierarquia, sem cadastro.",
            ],
            [
                "RN-CLI-01: Cliente cadastralmente é PJ; pessoas = responsáveis.",
                "RN-CLI-02: Visibilidade pela versão apostilada.",
                "RN-CLI-03: OS/orçamento preservam a versão (snapshot).",
                "RN-CLI-04: Rateio/segregação de saldo por secretaria/unidade no mesmo contrato.",
                "RN-ORC-01: Serviço exige orçamento prévio antes do consumo.",
                "RN-ORC-02: Licenciamento — orçamento pode ser dispensável.",
                "RN-SNAP-01: Snapshot na abertura de OS/orçamento.",
                "RN-SNAP-02: Auditoria recupera catálogo/produtos do snapshot.",
            ],
        )

        self.section("Backoffice MTI", "Análise · Publicação · Apostila")
        self.story(
            "Backoffice MTI — análise e publicação",
            "Nada que veio do parceiro entra em produção sem parecer da MTI.",
            "Fila Aguardando → analista confere valores/métricas/CSV → decide → status espelha no portal.",
            [
                "RF-MTI-01: Listas, parcerias e soluções.",
                "RF-MTI-02: Analisar (aprovar/ajuste/reprovar).",
                "RF-MTI-03: Publicar homologado.",
                "RF-MTI-04: Nova versão (congela anterior).",
                "RF-MTI-05: Apostilar no contrato.",
                "RF-MTI-06: Paralisar.",
                "RF-MTI-07: Catálogo universal.",
                "RF-MTI-08: Conferir fator com planilha CGS.",
                "RF-AN-01…05: Fila, decisão, motivo, publicar, espelho.",
            ],
            [
                "RN-MTI-01: Sem parecer, não publica.",
                "RN-MTI-02: Ajuste exige justificativa.",
                "RN-MTI-03: Publicar só após homologar.",
                "RN-MTI-04: Não edita versão homologada no lugar.",
                "RN-AN-01/02: Motivo e publicar só Homologado.",
            ],
        )

        # CATALOGO + ACEITE
        self.section("Catálogo RF / RN e aceite", "Consolidado · Critérios · Glossário · Fontes")
        self.table_slide(
            "RF por área (consolidado)",
            ["Área", "IDs", "Foco"],
            [
                ["Escopo / tipos", "RF-ESC-*, RF-TIPO-*", "Cobertura e flags"],
                ["Listas", "RF-MET, COB, CATG, GRP, MV", "Parametrização MTI"],
                ["Parceria / Solução", "RF-PAR-*, RF-SOL-*", "Governança"],
                ["Catálogo / Produto / Universal", "RF-CTL-*, RF-PROD-*, RF-UNI-*", "Cadastro, CSV, fator"],
                ["Dados de Parceria", "RF-DP-*", "Markup e 100%"],
                ["Fluxo / MTI", "RF-FLUX-*, RF-MTI-*, RF-AN-*", "Parecer e apostila"],
                ["Portal Parceiro", "RF-PARC-01…14", "Rascunho → envio"],
                ["Portal Cliente", "RF-CLI-01…14", "Consulta → OS"],
            ],
        )
        self.table_slide(
            "Regras de negócio críticas",
            ["ID", "Regra"],
            [
                ["RN-TIPO-01…06", "Tipos 1/2/3; elegibilidade USN/UST/HST; tipo decorre do objeto"],
                ["RN-P-01", "Produto referencia um catálogo (1:1 candidata — formalizar)"],
                ["RN-P-02…04", "Flags independentes; sem fator se só indiv. ou fora de métrica"],
                ["RN-P-05", "Markup em Dados de Parceria, não no Produto — protótipo"],
                ["RN-DP-01", "% parceiro + % MTI = 100 — protótipo"],
                ["RN-VER-01…03", "Versão no contrato; não sobrescreve; cliente após apostila"],
                ["RN-SNAP-01", "OS/orçamento com snapshot"],
                ["RN-SNAP-02", "Auditoria recupera catálogo/produtos do snapshot"],
                ["RN-FAT-01…03", "Fator a partir de unitário/moeda; consumo Tipo 3"],
                ["RN-ORC-01/02", "Serviço: orçamento prévio; licença: pode dispensar"],
                ["RN-CLI-01/04", "Cliente = PJ; saldo/segregação por secretaria"],
                ["RN-CSV-01", "Template CSV round-trip"],
                ["RN-FLUX / RN-PARC", "Parecer, justificativa, escopo do portal"],
                ["RN-PAR-01", "Primeiro cadastro da entidade Parceria = MTI"],
                ["RN-SOL-01", "Fabricante em texto (provisório)"],
            ],
        )
        self.table_slide(
            "Critérios de aceite",
            ["#", "Critério", "Origem"],
            [
                ["1", "Catálogo com 3 produtos (só univ., só indiv., ambos) e tags", "QA"],
                ["2", "Nova versão sem perder a anterior; auditável", "Confirmado (nº 1.0/1.1 = teste)"],
                ["3", "CSV ≥20 linhas; rejeitar inválida com mensagem", "QA / a especificar"],
                ["4", "Exportar CSV e reimportar em nova versão", "Confirmado (round-trip)"],
                ["5", "Serviço exibe complexidade; Licença oculta", "Confirmado"],
                ["6", "Fator se Universal=Sim; sem fator se só Individualizado", "Confirmado (UI = proposta)"],
                ["7", "Duas parcerias na mesma Org — permitido", "Confirmado"],
                ["8", "Ajuste com justificativa; aprovar→Homologado; publicar", "Proposto"],
                ["9", "Card: nome + tipo + tag; sem cobrança/métrica", "Parcial — Grupo candidata"],
                ["10", "Códigos SIAG/Protheus só com Preencher manualmente=Sim", "Protótipo"],
                ["11", "Parceiro: cadastro/CSV + validação MTI", "Confirmado (estados = proposto)"],
                ["12a", "Versão invisível ao cliente até apostila", "Confirmado"],
                ["12b", "Cotação gera demanda", "Proposto — separar de 12a"],
                ["13", "Distribuição ≠100% rejeitada", "Protótipo"],
                ["14", "OS/orçamento gravam versão (snapshot)", "Confirmado"],
                ["15", "Cliente = PJ; saldo/segregação por secretaria", "Confirmado"],
                ["16", "Serviço: orçamento antes do consumo", "Confirmado"],
            ],
        )
        self.table_slide(
            "Glossário e fontes",
            ["Termo / Fonte", "Significado / Uso"],
            [
                ["CGS", "Catálogo Geral de Soluções (Tipo 3)"],
                ["USN / UST / HST", "Métricas universais"],
                ["Fator", "Conversão preço → moeda Tipo 3 (fórmula oficial pendente)"],
                ["Apostila", "Libera nova versão no contrato do cliente"],
                ["Snapshot", "Versão gravada na OS/orçamento"],
                ["Status do fluxo", "Andamento portal↔MTI (proposto)"],
                ["Discovery 08/07/2026", "Regras confirmadas"],
                ["Protótipo atlas-prototipo", "Campos, métodos, classes"],
                ["Portais (apps)", "Fluxo e ações"],
                ["Minuta HTML/PPTX v0.9 · 04/08/2026", "Documento para validação MTI"],
            ],
        )

        self.prs.save(str(OUT))
        print(f"Wrote {OUT} · {OUT.stat().st_size/1024:.0f} KB · slides {self.n}")


if __name__ == "__main__":
    Deck().build()
